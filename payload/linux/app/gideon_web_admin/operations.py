from __future__ import annotations


import asyncio
import datetime as dt
import csv
import hashlib
import io
import json
import os
import shutil
import subprocess
import tempfile
import threading
import time
import sys
import uuid
import zipfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from gideon_admin import Site
from gideon_agent_orchestrator import build_agent_package
from gideon_agent_security import sha256_file, sign_file, sign_file_bytes
from gideon_agent_job_contract import build_job_envelope
from gideon_host_profiles import HostProfileStore
from gideon_domain import ArtifactDescriptor, ArtifactKind, HostTarget
from gideon_infrastructure.agent_file_transport import AgentHttpsPullFileTransport
from gideon_common.redaction import is_sensitive_key, redact_error
from gideon_core.connectivity import ConnectivityGate
from gideon_core.launch import AGENTLESS, LaunchRequest, LaunchService
from gideon_core.transport_profiles import VALID_TRANSPORT_PROFILES, normalize_transport_profile
from gideon_execution.runtime_service import FanoutStore
from gideon_execution.manifest_service import recalc_manifest, finalize_manifest, terminal_site
from gideon_site_groups import SITE_FIELDS, SiteGroupStore
from gideon_version import PRODUCT_VERSION

from .auth import SecurityStore
from .gateway import GatewayFacade
from .security_assets import SecurityAssetManager
from .ssh_probe import SshConnectionProbe
from .ssh_vault import harden_openssh_private_key_permissions
from .media_validation import MediaValidationError, validate_media_file


ALLOWED_SCRIPT_EXTENSIONS = {".ps1", ".py", ".sh", ".command", ".cmd", ".bat"}


def _accept_precreated_agentless_manifest_preflight(
    plan: Any,
    existing: dict[str, Any] | None,
    *,
    job_id: str,
    actor: str,
) -> bool:
    """Promote only the WebAdmin-owned async placeholder for this exact JobId.

    LaunchService correctly rejects an already-existing manifest for normal launches.
    RC30.95.27 intentionally creates a placeholder first so the browser can redirect
    immediately.  During the worker promotion that exact placeholder is expected.
    Only the single `Job único` preflight error is waived, and only when ownership,
    method and lifecycle state prove it is the same queued WebAdmin submission.
    Every other launch preflight error remains fail-closed.
    """
    if not isinstance(existing, dict):
        return False
    if str(existing.get("JobId") or "") != str(job_id):
        return False
    if not bool(existing.get("WebAdminLaunch")):
        return False
    if str(existing.get("CreatedBy") or "") != str(actor):
        return False
    if str(existing.get("ExecutionMethod") or "").upper() != AGENTLESS:
        return False
    state = str(existing.get("SubmissionState") or existing.get("Status") or "").upper()
    if state not in {"PREFLIGHT_QUEUED", "PREFLIGHT_RUNNING", "LAUNCHING"}:
        return False

    errors = list(getattr(plan.preflight, "errors", []) or [])
    unique = [item for item in errors if str(item).startswith("Job único:")]
    if not unique:
        return False
    plan.preflight.errors[:] = [item for item in errors if not str(item).startswith("Job único:")]
    for check in list(getattr(plan.preflight, "checks", []) or []):
        if str(check.get("Name") or "") == "Job único":
            check["Ok"] = True
            check["Warning"] = False
            check["Detail"] = "Manifest precreado por WebAdmin para este Job asíncrono"
    return True


class WebOperations:
    def __init__(self, app_root: Path, data_dir: Path, config: dict[str, Any], security: SecurityStore, gateway: GatewayFacade, security_assets: SecurityAssetManager | None = None, ssh_vault: Any | None = None, media_campaign_store: Any | None = None) -> None:
        self.app_root = Path(app_root).resolve()
        self.data_dir = Path(data_dir).resolve()
        self.config = config
        self.security = security
        self.gateway = gateway
        self.security_assets = security_assets
        self.ssh_vault = ssh_vault
        self.script_dir = self.data_dir / "scripts"
        self.job_dir = self.data_dir / "jobs"
        self.transfer_dir = self.data_dir / "transfers"
        self.transfer_job_dir = self.data_dir / "transfer-jobs"
        # The library is content-addressed and is never an executable path.
        # It remains inside GIDEON data with restrictive permissions; worker jobs
        # link/copy from it back into quarantine before sending a signed ZIP.
        self.media_asset_dir = self.data_dir / "media-assets"
        self.media_asset_blob_dir = self.media_asset_dir / "blobs"
        self.media_campaign_store = media_campaign_store
        self.script_dir.mkdir(parents=True, exist_ok=True)
        self.job_dir.mkdir(parents=True, exist_ok=True)
        self.transfer_dir.mkdir(parents=True, exist_ok=True)
        self.transfer_job_dir.mkdir(parents=True, exist_ok=True)
        self.media_asset_blob_dir.mkdir(parents=True, exist_ok=True)
        try:
            os.chmod(self.media_asset_dir, 0o750)
            os.chmod(self.media_asset_blob_dir, 0o750)
        except OSError:
            pass
        self._malware_scan_lock = threading.Lock()
        try:
            clamd_slots = max(1, min(int(self.config.get("clamd_parallel_scans", 2)), 4))
        except (TypeError, ValueError):
            clamd_slots = 2
        self._clamd_scan_slots = threading.BoundedSemaphore(clamd_slots)
        self._transfer_jobs_lock = threading.RLock()
        self._scheduled_transfer_jobs: set[str] = set()
        # Transfer preparation can include a long-running local antivirus scan.
        # Keep the HTTP request asynchronous, but never create an unbounded
        # number of threads when several operators upload at once.
        self._ensure_transfer_scheduler()
        self.profiles = HostProfileStore(self.app_root / "config" / "host_profiles.json")
        self.site_groups = SiteGroupStore(self.app_root / "config", self.app_root / "config" / "sites.csv")
        self.agentless_root = self.data_dir / "agentless"
        self.agentless_fanout_dir = self.agentless_root / "fanout_jobs"
        self.agentless_log_dir = self.agentless_root / "logs"
        self.agentless_config_dir = self.agentless_root / "configs"
        self.agentless_key_dir = self.agentless_root / "ephemeral-keys"
        for directory in (self.agentless_root, self.agentless_fanout_dir, self.agentless_log_dir, self.agentless_config_dir, self.agentless_key_dir):
            directory.mkdir(parents=True, exist_ok=True)
            try:
                os.chmod(directory, 0o700 if directory == self.agentless_key_dir else 0o750)
            except OSError:
                pass
        self._recover_transfer_jobs()

    def bind_media_campaign_store(self, store: Any | None) -> None:
        """Attach the additive Media catalogue without coupling normal transfers.

        The binding is injected by the WebAdmin composition root.  Existing
        script/file jobs keep working when Media is disabled or unavailable.
        """
        self.media_campaign_store = store

    @staticmethod
    def _windows_defender_scanner() -> str | None:
        if os.name != "nt":
            return None
        direct = shutil.which("MpCmdRun.exe")
        if direct:
            return direct
        candidates: list[Path] = []
        program_files = Path(os.environ.get("ProgramFiles") or r"C:\Program Files")
        candidates.append(program_files / "Windows Defender" / "MpCmdRun.exe")
        program_data = Path(os.environ.get("ProgramData") or r"C:\ProgramData")
        platform = program_data / "Microsoft" / "Windows Defender" / "Platform"
        try:
            candidates.extend(sorted(platform.glob("*/MpCmdRun.exe"), key=lambda p: p.parent.name, reverse=True))
        except OSError:
            pass
        return str(next((path for path in candidates if path.is_file()), "")) or None

    def _scan_with_windows_defender(self, candidate: Path) -> tuple[str, str]:
        scanner = self._windows_defender_scanner()
        if not scanner:
            return "SCAN_UNAVAILABLE", "Microsoft Defender MpCmdRun.exe no está disponible."
        started = time.monotonic()
        try:
            proc = subprocess.run(
                [scanner, "-Scan", "-ScanType", "3", "-File", str(candidate), "-DisableRemediation"],
                capture_output=True, text=True, timeout=180, shell=False,
                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            )
        except Exception as exc:
            elapsed_ms = int((time.monotonic() - started) * 1000)
            return "SCAN_ERROR", f"Engine=windows-defender; DurationMs={elapsed_ms}; {type(exc).__name__}: {exc}"
        output = ((proc.stdout or "") + "\n" + (proc.stderr or "")).strip()[:3000]
        elapsed_ms = int((time.monotonic() - started) * 1000)
        detail = f"Engine=windows-defender; DurationMs={elapsed_ms}; {output or 'análisis personalizado completado'}"
        if proc.returncode == 0:
            return "CLEAN", detail
        lowered = output.casefold()
        if any(token in lowered for token in ("threat", "infected", "malware", "virus", "amenaza")):
            return "INFECTED", detail
        return "SCAN_ERROR", detail + f"; rc={proc.returncode}"

    def _clamd_stream_command(self) -> tuple[list[str] | None, str]:
        """Return the private GIDEON clamd command, never a host-wide TCP daemon."""
        if os.name == "nt" or not bool(self.config.get("use_private_clamd", True)):
            return None, "clamd privado deshabilitado para esta plataforma/configuración"
        scanner = shutil.which("clamdscan")
        config_value = str(self.config.get("clamd_config_file") or "").strip()
        socket_value = str(self.config.get("clamd_socket") or "").strip()
        if not scanner:
            return None, "clamdscan no está instalado"
        if not config_value or not socket_value:
            return None, "socket/configuración clamd GIDEON no configurados"
        config_path = Path(os.path.expandvars(config_value)).expanduser().resolve()
        socket_path = Path(os.path.expandvars(socket_value)).expanduser()
        if not config_path.is_file():
            return None, f"configuración clamd GIDEON no disponible: {config_path}"
        if not socket_path.exists():
            return None, f"socket clamd GIDEON no disponible: {socket_path}"
        return [scanner, f"--config-file={config_path}", "--stream", "--no-summary"], ""

    def preferred_malware_engine(self) -> str:
        command, _ = self._clamd_stream_command()
        if command:
            return "clamd-stream"
        if shutil.which("clamscan"):
            return "clamscan-isolated"
        if self._windows_defender_scanner():
            return "windows-defender"
        return "unavailable"

    @staticmethod
    def _malware_engine_from_detail(detail: str, fallback: str) -> str:
        text = str(detail or "")
        # When clamd is unavailable the detail contains both the failed primary
        # engine and the successful fallback. Report the engine that produced
        # the final verdict.
        for engine in ("clamscan-isolated", "windows-defender", "clamd-stream"):
            if f"Engine={engine}" in text:
                return engine
        return fallback

    def _scan_with_private_clamd(self, candidate: Path) -> tuple[str, str] | None:
        command, unavailable_reason = self._clamd_stream_command()
        if command is None:
            return None
        try:
            timeout_seconds = max(30, min(int(self.config.get("clamd_scan_timeout_seconds", 180)), 1800))
        except (TypeError, ValueError):
            timeout_seconds = 180
        started = time.monotonic()
        try:
            # --stream uses clamd INSTREAM. The daemon never receives or opens the
            # GIDEON file path and therefore needs no access to Web/Gateway data.
            slots = getattr(self, "_clamd_scan_slots", None)
            if slots is None:
                proc = subprocess.run(
                    [*command, str(candidate)], capture_output=True, text=True,
                    timeout=timeout_seconds, shell=False,
                )
            else:
                with slots:
                    proc = subprocess.run(
                        [*command, str(candidate)], capture_output=True, text=True,
                        timeout=timeout_seconds, shell=False,
                    )
        except subprocess.TimeoutExpired:
            elapsed_ms = int((time.monotonic() - started) * 1000)
            return (
                "SCAN_TIMEOUT",
                f"Engine=clamd-stream; DurationMs={elapsed_ms}; no terminó en {timeout_seconds} segundos. "
                "El archivo permanece bloqueado y no fue enviado.",
            )
        except Exception as exc:
            elapsed_ms = int((time.monotonic() - started) * 1000)
            return "SCAN_UNAVAILABLE", f"Engine=clamd-stream; DurationMs={elapsed_ms}; {type(exc).__name__}: {exc}"
        output = ((proc.stdout or "") + "\n" + (proc.stderr or "")).strip()[:3000]
        elapsed_ms = int((time.monotonic() - started) * 1000)
        detail = f"Engine=clamd-stream; DurationMs={elapsed_ms}; {output or 'análisis completado'}"
        if proc.returncode == 0:
            return "CLEAN", detail
        if proc.returncode == 1:
            return "INFECTED", detail
        return "SCAN_UNAVAILABLE", detail + f"; rc={proc.returncode}"

    def scan_script(self, path: Path) -> tuple[str, str]:
        """Scan one GIDEON-owned file with the isolated GIDEON engine and YARA.

        Linux prefers a resident private clamd reached only through a GIDEON Unix
        socket and sends bytes with INSTREAM. This avoids reloading signatures for
        each upload and does not grant clamd access to Web/Gateway data. The older
        private-DB clamscan path remains a fail-safe fallback. Windows uses a
        one-file Microsoft Defender custom scan when ClamAV is unavailable.
        """
        required = bool(self.config.get("require_malware_scan_for_scripts", True))
        candidate = Path(path).resolve()
        allowed_roots = [self.data_dir]
        extra_root = str(self.config.get("malware_scan_allowed_root") or "").strip()
        if extra_root:
            allowed_roots.append(Path(os.path.expandvars(extra_root)).expanduser().resolve())
        if not any(candidate == root or root in candidate.parents for root in allowed_roots):
            return "SCAN_ERROR", f"Ruta fuera del sandbox GIDEON de análisis: {candidate}"

        details: list[str] = []
        clamd_result = self._scan_with_private_clamd(candidate)
        use_clamscan = clamd_result is None
        if clamd_result is not None:
            status, detail = clamd_result
            if status == "CLEAN":
                details.append(detail)
                use_clamscan = False
            elif status == "SCAN_UNAVAILABLE" and bool(self.config.get("clamd_fallback_to_clamscan", True)):
                details.append(detail)
                use_clamscan = True
            else:
                return status, detail

        scanner = shutil.which("clamscan") if use_clamscan else None
        if scanner:
            cmd = [scanner, "--no-summary"]
            database_dir = str(self.config.get("clamav_database_dir") or "").strip()
            if database_dir:
                db = Path(os.path.expandvars(database_dir)).expanduser().resolve()
                if not db.is_dir():
                    if required:
                        return "SCAN_ERROR", f"Base ClamAV GIDEON no disponible: {db}"
                else:
                    cmd.append(f"--database={db}")
            cmd.append(str(candidate))
            try:
                timeout_seconds = max(30, min(int(self.config.get("malware_scan_timeout_seconds", 300)), 1800))
            except (TypeError, ValueError):
                timeout_seconds = 300
            try:
                started = time.monotonic()
                # Fallback only: clamscan loads signatures for every run, so keep
                # it serialized to prevent concurrent uploads exhausting memory.
                with self._malware_scan_lock:
                    proc = subprocess.run(
                        cmd, capture_output=True, text=True, timeout=timeout_seconds, shell=False,
                    )
            except subprocess.TimeoutExpired:
                elapsed_ms = int((time.monotonic() - started) * 1000)
                return (
                    "SCAN_TIMEOUT",
                    f"Engine=clamscan-isolated; DurationMs={elapsed_ms}; no terminó en {timeout_seconds} segundos. "
                    "El archivo permanece bloqueado y no fue enviado.",
                )
            except Exception as exc:
                elapsed_ms = int((time.monotonic() - started) * 1000)
                return "SCAN_ERROR", f"Engine=clamscan-isolated; DurationMs={elapsed_ms}; {type(exc).__name__}: {exc}"
            output = ((proc.stdout or "") + "\n" + (proc.stderr or "")).strip()[:3000]
            elapsed_ms = int((time.monotonic() - started) * 1000)
            detail = f"Engine=clamscan-isolated; DurationMs={elapsed_ms}; {output or 'análisis completado'}"
            details.append(detail)
            if proc.returncode == 1:
                return "INFECTED", detail
            if proc.returncode not in {0}:
                if required:
                    return "SCAN_ERROR", detail + f"; rc={proc.returncode}"
        elif use_clamscan and os.name == "nt" and bool(self.config.get("windows_defender_custom_scan", True)):
            status, detail = self._scan_with_windows_defender(candidate)
            if status != "CLEAN":
                return status, detail
            details.append(detail)
        elif use_clamscan and required:
            return "SCAN_UNAVAILABLE", "No hay analizador antimalware GIDEON disponible (clamscan/Defender custom scan)."
        elif use_clamscan:
            details.append("Analizador antimalware no disponible; política no obligatoria.")

        yara_rules = str(self.config.get("yara_rules_file") or "").strip()
        if yara_rules:
            yara = shutil.which("yara")
            rules = Path(os.path.expandvars(yara_rules)).expanduser()
            if not yara or not rules.is_file():
                if required:
                    return "SCAN_ERROR", "YARA configurado pero binario/reglas no disponibles."
            else:
                proc = subprocess.run([yara, str(rules), str(candidate)], capture_output=True, text=True, timeout=30, shell=False)
                output = (proc.stdout + "\n" + proc.stderr).strip()[:3000]
                if proc.returncode == 0 and proc.stdout.strip():
                    return "REVIEW_REQUIRED", "YARA: " + proc.stdout.strip()[:2500]
                if proc.returncode not in {0, 1}:
                    return "SCAN_ERROR", output or f"YARA rc={proc.returncode}"
                details.append(output)
        return "CLEAN", " | ".join(x for x in details if x) or "Análisis estático limpio."

    def ingest_script(self, *, filename: str, content: bytes, actor: str) -> dict[str, Any]:
        name = Path(str(filename or "script")).name
        ext = Path(name).suffix.lower()
        if ext not in ALLOWED_SCRIPT_EXTENSIONS:
            raise ValueError(f"Extensión no permitida: {ext or '<sin extensión>'}")
        max_bytes = int(self.config.get("max_script_upload_bytes", 5 * 1024 * 1024))
        if not content or len(content) > max_bytes:
            raise ValueError(f"El script debe tener entre 1 byte y {max_bytes} bytes.")
        digest = hashlib.sha256(content).hexdigest()
        safe = "".join(c if c.isalnum() or c in "._-" else "_" for c in name)
        target = self.script_dir / f"{dt.datetime.now().strftime('%Y%m%d%H%M%S')}-{digest[:12]}-{safe}"
        target.write_bytes(content)
        os.chmod(target, 0o640)
        sid = self.security.add_script(
            name=name, stored_path=str(target), sha256=digest, size_bytes=len(content), extension=ext, actor=actor,
        )
        scan_status, scan_detail = self.scan_script(target)
        self.security.set_script_scan(sid, scan_status, scan_detail)
        return self.security.get_script(sid) or {}

    def _signing_material(self) -> tuple[Path | None, bytes | None, str]:
        """Resolve signing material without persisting managed private keys in plaintext."""
        if self.security_assets is not None:
            status = self.security_assets.signing_status()
            if bool(status.get("configured")):
                private_bytes = self.security_assets.signing_private_key_bytes()
                return None, private_bytes, self.security_assets.signing_key_id()

        configured = str(self.config.get("execution_signing_private_key") or os.environ.get("GIDEON_AGENT_SIGNING_PRIVATE_KEY", "")).strip()
        if not configured:
            raise RuntimeError(
                "Ejecución Web deshabilitada: configure una clave Ed25519 desde Seguridad > Claves y secretos "
                "o use una ruta externa de solo lectura."
            )
        path = Path(os.path.expandvars(configured)).expanduser().resolve()
        if not path.is_file():
            raise FileNotFoundError(f"Clave de firma no disponible: {path}")
        return path, None, str(self.config.get("signing_key_id") or "gideon-agent-ed25519-1")

    def execution_available(self) -> tuple[bool, str]:
        try:
            private_path, private_bytes, key_id = self._signing_material()
            if self.security_assets is not None and private_bytes is not None:
                test = self.security_assets.test_signing()
                return True, f"Firma disponible · {test.get('key_id') or key_id} · {test.get('fingerprint') or ''}".strip()
            return True, f"Firma disponible · {key_id}"
        except Exception as exc:
            return False, str(exc)

    @staticmethod
    def _preflight_agent_operation(row: dict[str, Any], *, script_paths: list[Path] | None = None, file_transfer: bool = False) -> None:
        """Fail early when the last heartbeat proves the selected Agent cannot run the requested operation."""
        capabilities = row.get("Capabilities") if isinstance(row.get("Capabilities"), dict) else {}
        operations = capabilities.get("Operations") if isinstance(capabilities.get("Operations"), dict) else {}
        agent_id = str(row.get("agent_id") or "")
        label = f"{row.get('site_code') or '-'} · {row.get('agent_name') or row.get('hostname') or agent_id}"
        if operations:
            if file_transfer and operations.get("FileTransfer") is False:
                raise RuntimeError(f"AGENT_CAPABILITY_MISSING: {label} no declara FILE_TRANSFER.")
            if script_paths and operations.get("ScriptExecution") is False:
                raise RuntimeError(f"AGENT_CAPABILITY_MISSING: {label} no declara SCRIPT_EXECUTION.")

        if not script_paths:
            return
        allowed = {str(x).lower() for x in (capabilities.get("AllowedExtensions") or []) if str(x).strip()}
        executors = capabilities.get("Executors") if isinstance(capabilities.get("Executors"), dict) else {}
        for path in script_paths:
            ext = path.suffix.lower()
            if allowed and ext not in allowed:
                raise RuntimeError(f"AGENT_SCRIPT_EXTENSION_UNSUPPORTED: {label} no permite {ext} ({path.name}).")
            # Executor checks are applied only when the Agent reported its runtime
            # capabilities. Older Agents remain compatible and will be checked by
            # the signed worker on the endpoint.
            if executors:
                if ext == ".ps1" and not (executors.get("powershell") or executors.get("pwsh")):
                    raise RuntimeError(f"AGENT_EXECUTOR_MISSING: {label} no tiene PowerShell/pwsh para {path.name}.")
                if ext == ".py" and not executors.get("python"):
                    raise RuntimeError(f"AGENT_EXECUTOR_MISSING: {label} no tiene Python para {path.name}.")
                if ext in {".sh", ".command"} and not (executors.get("bash") or executors.get("sh") or executors.get("zsh")):
                    raise RuntimeError(f"AGENT_EXECUTOR_MISSING: {label} no tiene shell compatible para {path.name}.")
                if ext in {".cmd", ".bat"} and str(row.get("platform") or "").lower().startswith("win") and not executors.get("cmd"):
                    # 2.3.3 did not report COMSPEC; don't block that legacy shape.
                    if operations:
                        raise RuntimeError(f"AGENT_EXECUTOR_MISSING: {label} no reporta cmd.exe para {path.name}.")

    def agentless_groups_view(self, group_id: str | None = None) -> dict[str, Any]:
        groups = self.site_groups.list_groups()
        selected = str(group_id or "").strip()
        valid_ids = {str(row.get("id") or "") for row in groups}
        if selected not in valid_ids:
            selected = self.site_groups.active_group_id
        sites = self.site_groups.read_sites(selected)
        core_cfg = self._load_core_config()
        ssh_cfg = dict(core_cfg.get("ssh") or {})
        try:
            default_port = int(ssh_cfg.get("port") or 22)
        except (TypeError, ValueError):
            default_port = 22
        return {
            "groups": groups,
            "selected_group": selected,
            "sites": sites,
            "ssh_identities": self.ssh_vault.list_identities() if self.ssh_vault is not None else [],
            "ssh_defaults": {
                "username": str(ssh_cfg.get("user") or "damautapp").strip() or "damautapp",
                "port": max(1, min(default_port, 65535)),
            },
        }

    def _load_core_config(self) -> dict[str, Any]:
        path = self.app_root / "config" / "config.json"
        value = json.loads(path.read_text(encoding="utf-8-sig"))
        if not isinstance(value, dict):
            raise RuntimeError("config/config.json no contiene un objeto JSON válido.")
        return value

    def _approved_script_paths(self, script_ids: list[int], *, agentless: bool = False) -> tuple[list[dict[str, Any]], list[Path]]:
        if not script_ids:
            raise ValueError("Seleccione al menos un script aprobado.")
        rows: list[dict[str, Any]] = []
        paths: list[Path] = []
        for sid in script_ids:
            row = self.security.get_script(int(sid))
            if row is None or str(row.get("approval_state") or "").upper() != "APPROVED":
                raise RuntimeError(f"Script {sid} no está APPROVED.")
            path = Path(str(row.get("stored_path") or ""))
            if not path.is_file() or sha256_file(path).lower() != str(row.get("sha256") or "").lower():
                raise RuntimeError(f"Integridad inválida para script {row.get('name') or sid}.")
            if agentless and path.suffix.lower() != ".ps1":
                raise RuntimeError(
                    f"AGENTLESS_SCRIPT_UNSUPPORTED: {row.get('name') or path.name}. "
                    "La ruta SSH/SFTP ejecuta PowerShell .ps1; para Python/CMD/Shell use Agent HTTPS/PULL."
                )
            rows.append(row)
            paths.append(path)
        return rows, paths

    def _agentless_job_config(self, *, job_id: str, key_path: Path, transport_profile: str, ssh_username: str, ssh_port: int) -> tuple[dict[str, Any], Path]:
        cfg = self._load_core_config()
        paths = dict(cfg.get("paths") or {})
        root = self.agentless_root
        absolute_paths = {
            "known_hosts": root / "known_hosts",
            "jobs_dir": root / "jobs",
            "logs_dir": self.agentless_log_dir,
            "results_dir": root / "results",
            "raw_dir": root / "results" / "raw",
            "fanout_jobs_dir": self.agentless_fanout_dir,
            "script_library_dir": root / "script_library",
            "batch_jobs_dir": root / "batches",
            "packages_dir": root / "packages",
            "transfer_jobs_dir": root / "transfer_jobs",
        }
        for name, path in absolute_paths.items():
            path.parent.mkdir(parents=True, exist_ok=True)
            paths[name] = str(path.resolve())
        cfg["paths"] = paths
        ssh = dict(cfg.get("ssh") or {})
        ssh["key_path"] = str(Path(key_path).resolve())
        ssh["user"] = str(ssh_username or "").strip()
        ssh["port"] = int(ssh_port)
        cfg["ssh"] = ssh
        fanout = dict(cfg.get("fanout") or {})
        fanout["transport_profile"] = normalize_transport_profile(transport_profile)
        cfg["fanout"] = fanout
        out = self.agentless_config_dir / f"{job_id}.json"
        out.write_text(json.dumps(cfg, ensure_ascii=False, indent=2), encoding="utf-8")
        try:
            os.chmod(out, 0o640)
        except OSError:
            pass
        return cfg, out

    @staticmethod
    def _hidden_process_kwargs() -> dict[str, Any]:
        if os.name != "nt":
            return {"start_new_session": True}
        startup = subprocess.STARTUPINFO()
        startup.dwFlags |= subprocess.STARTF_USESHOWWINDOW
        startup.wShowWindow = getattr(subprocess, "SW_HIDE", 0)
        return {
            "creationflags": getattr(subprocess, "CREATE_NO_WINDOW", 0),
            "startupinfo": startup,
        }

    @staticmethod
    def _agentless_prelaunch_terminal(status: str) -> bool:
        return str(status or "").upper() in {
            "COMPLETED", "COMPLETED_WITH_WARNINGS", "COMPLETED_WITH_ERRORS",
            "COMPLETED_WITH_SKIPPED", "PARTIAL", "FAILED", "LAUNCH_FAILED",
            "ACTION_REQUIRED", "CANCELLED", "CANCELLED_WITH_ERRORS", "INTERRUPTED",
        }

    def _agentless_submission_cancelled(self, job_id: str) -> bool:
        if not job_id:
            return False
        row = self.agentless_job(job_id)
        return bool(row and str(row.get("Status") or "").upper() == "CANCELLED")

    def _agentless_write_submission_state(
        self,
        job_id: str,
        *,
        status: str,
        error: str = "",
        finished: bool = False,
    ) -> None:
        row = self.agentless_job(job_id)
        if row is None:
            return
        current_status = str(row.get("Status") or "").upper()
        if self._agentless_prelaunch_terminal(current_status) and current_status != "ACTION_REQUIRED":
            return
        stamp = dt.datetime.now().astimezone().isoformat(timespec="seconds")
        row = dict(row)
        row["Status"] = str(status or "PREFLIGHT_RUNNING").upper()
        row["SubmissionState"] = row["Status"]
        if error:
            row["LaunchError"] = str(error)[:5000]
            row["SubmissionError"] = str(error)[:5000]
        if finished:
            row["FinishedAt"] = stamp
            row["SubmissionFinishedAt"] = stamp
        else:
            row["SubmissionHeartbeatAt"] = stamp
        FanoutStore(self.agentless_fanout_dir).write(row)

    def enqueue_agentless(
        self,
        *,
        group_id: str,
        site_codes: list[str],
        script_ids: list[int],
        actor: str,
        ssh_identity_id: str,
        ssh_username: str = "",
        ssh_port: int | None = None,
        transport_profile: str = "SFTP_SSH_SFTP",
        network_mode: str = "AUTO",
        continue_on_error: bool = True,
        export_format: str = "xlsx",
    ) -> dict[str, Any]:
        """Accept an Agentless submission quickly and run network preflight off-request.

        The browser receives a JobId immediately.  TCP/22, SSH authentication and SFTP
        validation continue in a GIDEON-owned background thread while /jobs/<JobId> shows
        progress.  No key material is written until authenticated preflight has passed.
        """
        transport_profile = normalize_transport_profile(transport_profile)
        if transport_profile not in VALID_TRANSPORT_PROFILES:
            raise ValueError(f"TransportProfile no soportado: {transport_profile}")
        network_mode = str(network_mode or "AUTO").strip().upper()
        if network_mode not in {"AUTO", "VPN_REQUIRED", "DIRECT_ALLOWED"}:
            raise ValueError("NetworkMode debe ser AUTO, VPN_REQUIRED o DIRECT_ALLOWED.")
        if self.ssh_vault is None:
            raise RuntimeError("SSH_VAULT_UNAVAILABLE: el WebAdmin no tiene vault SSH disponible.")
        if not ssh_identity_id:
            raise ValueError("Seleccione una identidad SSH administrada.")

        # Cheap/local validation only. The worker repeats all security checks before launch.
        script_rows, _scripts = self._approved_script_paths(script_ids, agentless=True)
        all_sites = self.site_groups.read_sites(group_id)
        wanted = {str(code).strip() for code in site_codes if str(code).strip()}
        if not wanted:
            raise ValueError("Seleccione al menos un host del inventario.")
        selected = [row for row in all_sites if str(row.get("SiteCode") or "").strip() in wanted]
        found = {str(row.get("SiteCode") or "").strip() for row in selected}
        missing = sorted(wanted - found)
        if missing:
            raise RuntimeError("Hosts no encontrados en el grupo seleccionado: " + ", ".join(missing[:20]))
        empty_ip = [str(row.get("SiteCode") or "") for row in selected if not str(row.get("ServerIP") or "").strip()]
        if empty_ip:
            raise RuntimeError("AGENTLESS_HOST_IP_MISSING: " + ", ".join(empty_ip[:20]))

        identity_row = next(
            (row for row in self.ssh_vault.list_identities() if str(row.get("id") or "") == str(ssh_identity_id)),
            None,
        )
        if identity_row is None:
            raise ValueError("SSH_IDENTITY_NOT_FOUND: seleccione una identidad SSH administrada válida.")
        profile_ok = bool(identity_row.get("connection_ssh_ok")) and bool(identity_row.get("connection_sftp_ok"))
        profile_user = str(identity_row.get("default_username") or "").strip()
        try:
            profile_port = int(identity_row.get("default_port") or 22)
        except (TypeError, ValueError):
            profile_port = 22
        if not profile_ok or not profile_user:
            raise ValueError(
                "SSH_GLOBAL_PROFILE_REQUIRED: valide primero la identidad en Claves y secretos → Probar conexión."
            )
        if len(profile_user) > 64 or any(ch.isspace() for ch in profile_user):
            raise ValueError("SSH_USERNAME_INVALID: el perfil global contiene un usuario SSH inválido.")
        if profile_port < 1 or profile_port > 65535:
            raise ValueError("SSH_PORT_INVALID: el perfil global contiene un puerto fuera de rango.")

        job_id = dt.datetime.now().strftime("%Y%m%d-%H%M%S") + "-ssh-" + uuid.uuid4().hex[:6]
        stamp = dt.datetime.now().astimezone().isoformat(timespec="seconds")
        require_sftp = transport_profile == "SFTP_SSH_SFTP"
        sites_manifest: dict[str, dict[str, Any]] = {}
        for row in selected:
            code = str(row.get("SiteCode") or "").strip()
            sites_manifest[code] = {
                "SiteCode": code,
                "SiteName": str(row.get("SiteName") or "").strip(),
                "ServerIP": str(row.get("ServerIP") or "").strip(),
                "Platform": str(row.get("Platform") or "Windows/Auto").strip(),
                "RemoteState": "PREFLIGHT_PENDING",
                "TransferStatus": "PENDING",
                "CollectionStatus": "PENDING",
            }
        placeholder = {
            "JobId": job_id,
            "ExecutionMethod": AGENTLESS,
            "TransportProfile": transport_profile,
            "NetworkPolicy": network_mode,
            "Status": "PREFLIGHT_QUEUED",
            "SubmissionState": "PREFLIGHT_QUEUED",
            "CreatedAt": stamp,
            "UpdatedAt": stamp,
            "QueuedAt": stamp,
            "CreatedBy": actor,
            "WebAdminLaunch": True,
            "TotalCount": len(selected),
            "CollectedCount": 0,
            "FailedCount": 0,
            "Scripts": [{"Id": row["id"], "Name": row["name"], "Sha256": row["sha256"]} for row in script_rows],
            "SshIdentity": {
                "id": str(identity_row.get("id") or ""),
                "label": str(identity_row.get("label") or ""),
                "fingerprint": str(identity_row.get("fingerprint") or ""),
                "username": profile_user,
                "port": profile_port,
            },
            "ProtocolPreflight": {
                "Profile": transport_profile,
                "Username": profile_user,
                "Port": profile_port,
                "SFTPRequired": bool(require_sftp),
                "TargetCount": len(selected),
                "CheckedCount": 0,
                "PassedCount": 0,
                "FailedCount": 0,
                "Sites": {},
            },
            "Sites": sites_manifest,
        }
        FanoutStore(self.agentless_fanout_dir).write(placeholder)

        def _submission_worker() -> None:
            try:
                if self._agentless_submission_cancelled(job_id):
                    return
                self._agentless_write_submission_state(job_id, status="PREFLIGHT_RUNNING")
                asyncio.run(self.dispatch_agentless(
                    group_id=group_id,
                    site_codes=list(wanted),
                    script_ids=list(script_ids),
                    actor=actor,
                    ssh_identity_id=ssh_identity_id,
                    ssh_username=ssh_username,
                    ssh_port=ssh_port,
                    transport_profile=transport_profile,
                    network_mode=network_mode,
                    continue_on_error=continue_on_error,
                    export_format=export_format,
                    job_id_override=job_id,
                ))
            except Exception as exc:
                row = self.agentless_job(job_id)
                if row is not None and str(row.get("Status") or "").upper() == "CANCELLED":
                    return
                message = redact_error(exc)
                action_tokens = (
                    "SSH_GLOBAL_PROFILE_REQUIRED", "SSH_NETWORK_UNREACHABLE",
                    "VPN_REQUIRED_BUT_NOT_ACTIVE", "SSH_SFTP_AUTH_PREFLIGHT_FAILED",
                    "SSH_SFTP_PARTIAL_PREFLIGHT_FAILED", "SSH_IDENTITY_NOT_FOUND",
                    "AGENTLESS_HOST_IP_MISSING",
                )
                status = "ACTION_REQUIRED" if any(token in message for token in action_tokens) else "FAILED"
                self._agentless_write_submission_state(job_id, status=status, error=message, finished=True)

        threading.Thread(
            target=_submission_worker,
            name=f"gideon-agentless-submit-{job_id}",
            daemon=True,
        ).start()
        return {
            "JobId": job_id,
            "ExecutionMethod": AGENTLESS,
            "TransportProfile": transport_profile,
            "NetworkPolicy": network_mode,
            "ConnectivityMode": "PREFLIGHT_PENDING",
            "Targets": len(selected),
            "Scripts": [{"Id": row["id"], "Name": row["name"], "Sha256": row["sha256"]} for row in script_rows],
            "Accepted": True,
            "Status": "PREFLIGHT_QUEUED",
        }

    async def dispatch_agentless(
        self,
        *,
        group_id: str,
        site_codes: list[str],
        script_ids: list[int],
        actor: str,
        ssh_identity_id: str,
        ssh_username: str = "",
        ssh_port: int | None = None,
        transport_profile: str = "SFTP_SSH_SFTP",
        network_mode: str = "AUTO",
        continue_on_error: bool = True,
        export_format: str = "xlsx",
        job_id_override: str = "",
    ) -> dict[str, Any]:
        """Launch the same Agentless SSH/SFTP fan-out exposed by the desktop GUI.

        The WebAdmin never injects SSH credentials into the command line.  A selected
        vault identity is decrypted to a 0600 GIDEON-owned ephemeral file for the life of
        the backend process and deleted by a watcher when the Job exits.
        """
        transport_profile = normalize_transport_profile(transport_profile)
        if transport_profile not in VALID_TRANSPORT_PROFILES:
            raise ValueError(f"TransportProfile no soportado: {transport_profile}")
        network_mode = str(network_mode or "AUTO").strip().upper()
        if network_mode not in {"AUTO", "VPN_REQUIRED", "DIRECT_ALLOWED"}:
            raise ValueError("NetworkMode debe ser AUTO, VPN_REQUIRED o DIRECT_ALLOWED.")
        if self.ssh_vault is None:
            raise RuntimeError("SSH_VAULT_UNAVAILABLE: el WebAdmin no tiene vault SSH disponible.")
        if not ssh_identity_id:
            raise ValueError("Seleccione una identidad SSH administrada.")

        script_rows, scripts = self._approved_script_paths(script_ids, agentless=True)
        all_sites = self.site_groups.read_sites(group_id)
        wanted = {str(code).strip() for code in site_codes if str(code).strip()}
        if not wanted:
            raise ValueError("Seleccione al menos un host del inventario.")
        selected = [row for row in all_sites if str(row.get("SiteCode") or "").strip() in wanted]
        found = {str(row.get("SiteCode") or "").strip() for row in selected}
        missing = sorted(wanted - found)
        if missing:
            raise RuntimeError("Hosts no encontrados en el grupo seleccionado: " + ", ".join(missing[:20]))
        empty_ip = [str(row.get("SiteCode") or "") for row in selected if not str(row.get("ServerIP") or "").strip()]
        if empty_ip:
            raise RuntimeError("AGENTLESS_HOST_IP_MISSING: " + ", ".join(empty_ip[:20]))

        core_cfg = self._load_core_config()
        connectivity_cfg = dict(core_cfg.get("connectivity") or {})
        connectivity_cfg["sample_targets"] = max(1, min(50, len(selected)))
        snapshot = await ConnectivityGate(connectivity_cfg).probe([str(row.get("ServerIP") or "") for row in selected])
        if job_id_override:
            if self._agentless_submission_cancelled(job_id_override):
                return {"JobId": job_id_override, "ExecutionMethod": AGENTLESS, "Status": "CANCELLED"}
            current = self.agentless_job(job_id_override)
            if current is not None:
                current = dict(current)
                current["ConnectivityPreflight"] = snapshot.as_dict()
                current["Status"] = "PREFLIGHT_RUNNING"
                current["SubmissionState"] = "PREFLIGHT_RUNNING"
                FanoutStore(self.agentless_fanout_dir).write(current)
        if not snapshot.Reachable:
            raise RuntimeError(
                "SSH_NETWORK_UNREACHABLE: ninguno de los hosts seleccionados responde TCP/22 desde el Control Plane. "
                f"Modo={snapshot.Mode} fallidos={','.join(snapshot.FailedTargets[:10])}"
            )
        if network_mode == "VPN_REQUIRED" and snapshot.Mode not in {"VPN", "MIXED"}:
            raise RuntimeError(
                "VPN_REQUIRED_BUT_NOT_ACTIVE: existe ruta TCP/22, pero la IP origen no pertenece a las redes VPN "
                f"configuradas. Modo detectado={snapshot.Mode}; SourceIPs={','.join(snapshot.SourceIPs) or '-'}"
            )

        # GIDEON RC30.95.24 · validated GLOBAL identity profile.
        # The site is a destination only; username/port belong to the selected identity.
        ssh_cfg = dict(core_cfg.get("ssh") or {})
        identity_row = next(
            (row for row in self.ssh_vault.list_identities() if str(row.get("id") or "") == str(ssh_identity_id)),
            None,
        )
        if identity_row is None:
            raise ValueError("SSH_IDENTITY_NOT_FOUND: seleccione una identidad SSH administrada válida.")
        profile_ok = bool(identity_row.get("connection_ssh_ok")) and bool(identity_row.get("connection_sftp_ok"))
        profile_user = str(identity_row.get("default_username") or "").strip()
        try:
            profile_port = int(identity_row.get("default_port") or 22)
        except (TypeError, ValueError):
            profile_port = 22
        if not profile_ok or not profile_user:
            raise ValueError(
                "SSH_GLOBAL_PROFILE_REQUIRED: valide primero la identidad en Claves y secretos → Probar conexión. "
                "El usuario/puerto se reutilizan para todas las sedes y cada destino se valida antes del Job."
            )
        ssh_user = profile_user
        ssh_port = profile_port
        if len(ssh_user) > 64 or any(ch.isspace() for ch in ssh_user):
            raise ValueError("SSH_USERNAME_INVALID: el perfil global contiene un usuario SSH inválido.")
        if ssh_port < 1 or ssh_port > 65535:
            raise ValueError("SSH_PORT_INVALID: el perfil global contiene un puerto fuera de rango.")
        try:
            ssh_timeout = int(ssh_cfg.get("connect_timeout_seconds") or 12)
        except (TypeError, ValueError):
            ssh_timeout = 12
        ssh_timeout = max(3, min(ssh_timeout, 30))
        require_sftp = transport_profile == "SFTP_SSH_SFTP"
        protocol_probe = SshConnectionProbe(self.app_root, self.data_dir, self.ssh_vault)
        protocol_sites: dict[str, dict[str, Any]] = {}
        concurrency = max(1, min(6, len(selected)))
        semaphore = asyncio.Semaphore(concurrency)
        progress_lock = asyncio.Lock()

        async def _persist_protocol_progress(code: str) -> None:
            if not job_id_override:
                return
            async with progress_lock:
                current = self.agentless_job(job_id_override)
                if current is None or str(current.get("Status") or "").upper() == "CANCELLED":
                    return
                checked = len(protocol_sites)
                passed = sum(1 for item in protocol_sites.values() if bool(item.get("Ok")))
                failed = max(0, checked - passed)
                current = dict(current)
                current["Status"] = "PREFLIGHT_RUNNING"
                current["SubmissionState"] = "PREFLIGHT_RUNNING"
                current["ProtocolPreflight"] = {
                    "Profile": transport_profile,
                    "Username": ssh_user,
                    "Port": ssh_port,
                    "SFTPRequired": bool(require_sftp),
                    "TargetCount": len(selected),
                    "CheckedCount": checked,
                    "PassedCount": passed,
                    "FailedCount": failed,
                    "Sites": dict(protocol_sites),
                }
                sites_now = current.get("Sites") if isinstance(current.get("Sites"), dict) else {}
                item = protocol_sites.get(code) if isinstance(protocol_sites.get(code), dict) else {}
                site_row = dict(sites_now.get(code) or {})
                site_row["RemoteState"] = "PREFLIGHT_OK" if item.get("Ok") else "PREFLIGHT_FAILED"
                site_row["PreflightCheckedAt"] = dt.datetime.now().astimezone().isoformat(timespec="seconds")
                if not item.get("Ok"):
                    site_row["LastError"] = str(item.get("Detail") or "SSH/SFTP preflight failed")[:1500]
                sites_now[code] = site_row
                current["Sites"] = sites_now
                FanoutStore(self.agentless_fanout_dir).write(current)

        async def _probe_agentless_site(row: dict[str, Any]) -> None:
            code = str(row.get("SiteCode") or "").strip()
            name = str(row.get("SiteName") or "").strip()
            host = str(row.get("ServerIP") or "").strip()
            if job_id_override and self._agentless_submission_cancelled(job_id_override):
                return
            async with semaphore:
                if job_id_override and self._agentless_submission_cancelled(job_id_override):
                    return
                try:
                    result = await asyncio.to_thread(
                        protocol_probe.probe_host,
                        identity_id=ssh_identity_id,
                        host=host,
                        site_code=code,
                        site_name=name,
                        username=ssh_user,
                        port=ssh_port,
                        timeout_seconds=ssh_timeout,
                        require_sftp=require_sftp,
                    )
                    ssh_result = result.get("ssh") if isinstance(result.get("ssh"), dict) else {}
                    sftp_result = result.get("sftp") if isinstance(result.get("sftp"), dict) else {}
                    protocol_sites[code] = {
                        "SiteCode": code,
                        "SiteName": name,
                        "Host": host,
                        "Ok": bool(result.get("ok")),
                        "SSH": {
                            "Ok": bool(ssh_result.get("ok")),
                            "Rc": int(ssh_result.get("returncode") or 0),
                        },
                        "SFTP": {
                            "Required": bool(require_sftp),
                            "Ok": bool(sftp_result.get("ok")),
                            "Rc": int(sftp_result.get("returncode") or 0),
                        },
                        "Detail": str(result.get("detail") or "")[:1500],
                    }
                except Exception as exc:
                    protocol_sites[code] = {
                        "SiteCode": code,
                        "SiteName": name,
                        "Host": host,
                        "Ok": False,
                        "SSH": {"Ok": False, "Rc": -1},
                        "SFTP": {"Required": bool(require_sftp), "Ok": False, "Rc": -1},
                        "Detail": redact_error(exc)[:1500],
                    }
                await _persist_protocol_progress(code)

        await asyncio.gather(*(_probe_agentless_site(row) for row in selected))
        protocol_passed = sum(1 for row in protocol_sites.values() if bool(row.get("Ok")))
        protocol_failed = max(0, len(selected) - protocol_passed)
        protocol_preflight = {
            "Profile": transport_profile,
            "Username": ssh_user,
            "Port": ssh_port,
            "SFTPRequired": bool(require_sftp),
            "TargetCount": len(selected),
            "CheckedCount": len(protocol_sites),
            "PassedCount": protocol_passed,
            "FailedCount": protocol_failed,
            "CheckedAt": dt.datetime.now().astimezone().isoformat(timespec="seconds"),
            "Sites": protocol_sites,
        }
        if job_id_override:
            current = self.agentless_job(job_id_override)
            if current is not None:
                if str(current.get("Status") or "").upper() == "CANCELLED":
                    return {"JobId": job_id_override, "ExecutionMethod": AGENTLESS, "Status": "CANCELLED"}
                current = dict(current)
                current["ProtocolPreflight"] = protocol_preflight
                current["Status"] = "LAUNCHING"
                current["SubmissionState"] = "LAUNCHING"
                FanoutStore(self.agentless_fanout_dir).write(current)
        if protocol_passed == 0:
            sample = " | ".join(
                f"{code}: {str(row.get('Detail') or 'FAIL')[:260]}"
                for code, row in list(protocol_sites.items())[:5]
            )
            raise RuntimeError(
                "SSH_SFTP_AUTH_PREFLIGHT_FAILED: TCP/22 es alcanzable, pero ninguna sede superó "
                f"la validación autenticada SSH{' + SFTP' if require_sftp else ''}. {sample}"
            )
        if protocol_failed and not continue_on_error:
            failed_codes = [code for code, row in protocol_sites.items() if not bool(row.get("Ok"))]
            raise RuntimeError(
                "SSH_SFTP_PARTIAL_PREFLIGHT_FAILED: continue_on_error=False y fallaron: "
                + ", ".join(failed_codes[:20])
            )

        # RC30.95.27.1: close the cancellation race between final preflight and
        # key/backend creation. A cancelled async Job must not materialize a key.
        if job_id_override and self._agentless_submission_cancelled(job_id_override):
            return {"JobId": job_id_override, "ExecutionMethod": AGENTLESS, "Status": "CANCELLED"}

        job_id = str(job_id_override or "").strip() or (dt.datetime.now().strftime("%Y%m%d-%H%M%S") + "-ssh-" + uuid.uuid4().hex[:6])
        job_root = self.agentless_root / "jobs" / job_id
        job_root.mkdir(parents=True, exist_ok=True)
        sites_file = job_root / "selected-sites.csv"
        with sites_file.open("w", encoding="utf-8-sig", newline="") as fh:
            writer = csv.DictWriter(fh, fieldnames=SITE_FIELDS)
            writer.writeheader()
            for row in selected:
                writer.writerow({field: str(row.get(field) or "").strip() for field in SITE_FIELDS})

        key_path = self.agentless_key_dir / f"{job_id}.key"
        if job_id_override and self._agentless_submission_cancelled(job_id_override):
            return {"JobId": job_id_override, "ExecutionMethod": AGENTLESS, "Status": "CANCELLED"}
        key_bytes = self.ssh_vault.private_key_bytes(ssh_identity_id)
        key_path.write_bytes(key_bytes)
        # Native Windows OpenSSH validates NTFS ACLs, not POSIX chmod bits.
        # Use the same hardened ACL contract as the interactive SSH/SFTP probe.
        harden_openssh_private_key_permissions(key_path)
        identity_row = next(
            (row for row in self.ssh_vault.list_identities() if str(row.get("id") or "") == str(ssh_identity_id)),
            {"id": ssh_identity_id},
        )

        _cfg, config_path = self._agentless_job_config(
            job_id=job_id, key_path=key_path, transport_profile=transport_profile,
            ssh_username=ssh_user, ssh_port=ssh_port,
        )
        launcher = LaunchService(
            app_root=self.app_root,
            jobs_dir=self.agentless_fanout_dir,
            logs_dir=self.agentless_log_dir,
            python_exe=sys.executable,
            version=PRODUCT_VERSION,
            config_path=config_path,
        )
        group_row = self.site_groups.get(group_id)
        request = LaunchRequest(
            job_id=job_id,
            method=AGENTLESS,
            sites_file=sites_file,
            sites=tuple(selected),
            scripts=tuple(scripts),
            export_format=export_format,
            continue_on_error=bool(continue_on_error),
            transport_profile=transport_profile,
            group_id=group_id,
            group_name=str(group_row.get("name") or group_id),
        )
        plan = launcher.build_plan(request)
        async_placeholder = self.agentless_job(job_id) if job_id_override else None
        if job_id_override:
            _accept_precreated_agentless_manifest_preflight(
                plan, async_placeholder, job_id=job_id, actor=actor
            )
        if not plan.preflight.ok:
            key_path.unlink(missing_ok=True)
            raise RuntimeError("Preflight Agentless falló: " + " | ".join(plan.preflight.errors[:8]))
        if job_id_override and self._agentless_submission_cancelled(job_id_override):
            key_path.unlink(missing_ok=True)
            return {"JobId": job_id_override, "ExecutionMethod": AGENTLESS, "Status": "CANCELLED"}

        # Promote the already-visible async placeholder into the canonical backend
        # manifest. bootstrap_manifest intentionally overwrites only after the strict
        # ownership/state gate above accepted the expected placeholder.
        launcher.bootstrap_manifest(plan)
        manifest = FanoutStore(self.agentless_fanout_dir).read(job_id)
        if isinstance(async_placeholder, dict):
            for keep_key in ("QueuedAt", "SubmissionStartedAt", "SubmissionHeartbeatAt"):
                if async_placeholder.get(keep_key) and not manifest.get(keep_key):
                    manifest[keep_key] = async_placeholder.get(keep_key)
            manifest["AsyncPlaceholderPromoted"] = True
            manifest["AsyncPlaceholderPromotedAt"] = dt.datetime.now().astimezone().isoformat(timespec="seconds")
        manifest.update({
            "CreatedBy": actor,
            "WebAdminLaunch": True,
            "NetworkPolicy": network_mode,
            "ConnectivityPreflight": snapshot.as_dict(),
            "ProtocolPreflight": protocol_preflight,
            "SshIdentity": {
                "id": str(identity_row.get("id") or ""),
                "label": str(identity_row.get("label") or ""),
                "fingerprint": str(identity_row.get("fingerprint") or ""),
                "username": ssh_user,
                "port": ssh_port,
            },
        })
        FanoutStore(self.agentless_fanout_dir).write(manifest)

        plan.log_path.parent.mkdir(parents=True, exist_ok=True)
        env = os.environ.copy()
        env.update({
            "PYTHONUTF8": "1",
            "PYTHONIOENCODING": "utf-8",
            "GIDEON_FORCE_UTF8_STDIO": "1",
            "GIDEON_BACKEND_LOG": str(plan.log_path),
            "GIDEON_SITE_GROUP_ID": group_id,
            "GIDEON_SITE_GROUP_NAME": str(group_row.get("name") or group_id),
            "GIDEON_WEB_AGENTLESS": "1",
        })
        try:
            with plan.log_path.open("a", encoding="utf-8", buffering=1) as log_file:
                log_file.write(f"DRC WebAdmin · Agentless · {job_id}\n")
                log_file.write(f"Transport={transport_profile} NetworkPolicy={network_mode} Detected={snapshot.Mode}\n")
                log_file.write(f"Command={plan.command_text()}\n\n")
                log_file.flush()
                process = subprocess.Popen(
                    list(plan.command),
                    cwd=str(self.app_root),
                    stdin=subprocess.DEVNULL,
                    stdout=log_file,
                    stderr=subprocess.STDOUT,
                    close_fds=(os.name != "nt"),
                    env=env,
                    **self._hidden_process_kwargs(),
                )
            launcher.mark_started(plan, int(process.pid))
        except Exception:
            key_path.unlink(missing_ok=True)
            raise

        def _watch_backend() -> None:
            try:
                return_code = process.wait()
                try:
                    current = FanoutStore(self.agentless_fanout_dir).read(job_id)
                except Exception:
                    current = {}
                # GIDEON RC30.95.23.2: backend exit watchdog. A backend that exits
                # while its manifest remains non-terminal must never leave 29/29 targets
                # permanently ACTIVE. Only this Job manifest and its own ephemeral key are touched.
                status = str(current.get("Status") or "").upper()
                terminal = {
                    "COMPLETED", "COMPLETED_WITH_WARNINGS", "COMPLETED_WITH_ERRORS",
                    "COMPLETED_WITH_SKIPPED", "PARTIAL", "FAILED", "LAUNCH_FAILED",
                    "ACTION_REQUIRED", "CANCELLED", "CANCELLED_WITH_ERRORS", "INTERRUPTED",
                }
                if status not in terminal:
                    stamp = dt.datetime.now().astimezone().isoformat(timespec="seconds")
                    current["Status"] = "INTERRUPTED"
                    current["LaunchReturnCode"] = int(return_code)
                    current["LaunchError"] = (
                        f"AGENTLESS_BACKEND_EXITED_NONTERMINAL: rc={return_code}; "
                        f"previous={status or '-'}; log={plan.log_path}"
                    )
                    current["FinishedAt"] = stamp
                    current["BackendExitDetectedAt"] = stamp
                    current_sites = current.get("Sites") if isinstance(current.get("Sites"), dict) else {}
                    site_terminal = {
                        "COLLECTED", "COLLECTED_WITH_ERRORS", "FAILED", "CANCELLED",
                        "CANCELLED_NOT_STARTED", "CANCEL_FAILED", "INTERRUPTED", "DISPATCH_FAILED", "IGNORED",
                    }
                    for code, raw in list(current_sites.items()):
                        if not isinstance(raw, dict):
                            continue
                        entry = dict(raw)
                        remote_state = str(entry.get("RemoteState") or "").upper()
                        if remote_state not in site_terminal:
                            entry["RemoteState"] = "INTERRUPTED"
                            entry.setdefault("LastError", current["LaunchError"])
                            entry["InterruptedAt"] = stamp
                        current_sites[code] = entry
                    current["Sites"] = current_sites
                    FanoutStore(self.agentless_fanout_dir).write(current)
            finally:
                try:
                    key_path.unlink(missing_ok=True)
                except OSError:
                    pass

        threading.Thread(target=_watch_backend, name=f"gideon-agentless-{job_id}", daemon=True).start()
        return {
            "JobId": job_id,
            "ExecutionMethod": AGENTLESS,
            "TransportProfile": transport_profile,
            "NetworkPolicy": network_mode,
            "ConnectivityMode": snapshot.Mode,
            "Targets": len(selected),
            "Scripts": [{"Id": row["id"], "Name": row["name"], "Sha256": row["sha256"]} for row in script_rows],
            "ManifestPath": str(plan.manifest_path),
            "LogPath": str(plan.log_path),
        }

    def agentless_job(self, job_id: str) -> dict[str, Any] | None:
        clean = "".join(ch for ch in str(job_id) if ch.isalnum() or ch in "._-")
        if clean != str(job_id) or not clean:
            return None
        path = self.agentless_fanout_dir / f"{clean}.json"
        if not path.is_file():
            return None
        try:
            value = json.loads(path.read_text(encoding="utf-8-sig"))
        except Exception:
            return None
        return value if isinstance(value, dict) else None

    def agentless_jobs(self, limit: int | None = None) -> list[dict[str, Any]]:
        return FanoutStore(self.agentless_fanout_dir).list(limit=limit)

    @staticmethod
    def _agentless_age_seconds(value: Any) -> float | None:
        raw = str(value or "").strip()
        if not raw:
            return None
        try:
            parsed = dt.datetime.fromisoformat(raw.replace("Z", "+00:00"))
            if parsed.tzinfo is None:
                parsed = parsed.astimezone()
            return max(0.0, (dt.datetime.now().astimezone() - parsed.astimezone()).total_seconds())
        except Exception:
            return None

    @staticmethod
    def _agentless_backend_process(pid: int, job_id: str) -> dict[str, Any]:
        result = {"pid": int(pid or 0), "alive": False, "owned": False, "detail": "PID_MISSING"}
        if int(pid or 0) <= 0:
            return result
        pid = int(pid)
        try:
            if os.name == "nt":
                ps = shutil.which("powershell.exe") or shutil.which("powershell")
                if not ps:
                    result["detail"] = "POWERSHELL_MISSING"
                    return result
                script = (
                    f'$p=Get-CimInstance Win32_Process -Filter "ProcessId = {pid}" -ErrorAction SilentlyContinue;'
                    'if($null -eq $p){exit 3};'
                    '$o=[pscustomobject]@{Name=$p.Name;CommandLine=$p.CommandLine};'
                    '$o|ConvertTo-Json -Compress'
                )
                cp = subprocess.run(
                    [ps, "-NoProfile", "-NonInteractive", "-Command", script],
                    capture_output=True, text=True, timeout=8, shell=False,
                    encoding="utf-8", errors="replace", **WebOperations._hidden_process_kwargs(),
                )
                if cp.returncode != 0 or not (cp.stdout or "").strip():
                    result["detail"] = "PROCESS_NOT_FOUND"
                    return result
                info = json.loads((cp.stdout or "{}").strip())
                cmdline = str(info.get("CommandLine") or "")
                name = str(info.get("Name") or "")
            else:
                cp = subprocess.run(
                    ["ps", "-p", str(pid), "-o", "comm=", "-o", "args="],
                    capture_output=True, text=True, timeout=5, shell=False,
                    encoding="utf-8", errors="replace",
                )
                if cp.returncode != 0 or not (cp.stdout or "").strip():
                    result["detail"] = "PROCESS_NOT_FOUND"
                    return result
                cmdline = (cp.stdout or "").strip()
                name = cmdline.split(None, 1)[0] if cmdline else ""
            low = cmdline.lower()
            owned = str(job_id) in cmdline and ("gideon_fanout" in low or "gideon-fanout" in low)
            result.update({"alive": True, "owned": bool(owned), "detail": "GIDEON_JOB_PROCESS" if owned else "PID_ALIVE_NOT_OWNED", "name": name})
            return result
        except Exception as exc:
            result["detail"] = "PROCESS_PROBE_ERROR: " + redact_error(exc)
            return result

    def agentless_runtime_state(self, job_id: str, manifest: dict[str, Any] | None = None) -> dict[str, Any]:
        row = manifest if isinstance(manifest, dict) else self.agentless_job(job_id)
        if row is None:
            return {"stale": False, "terminal": False, "age_seconds": None, "backend": {"alive": False, "owned": False, "detail": "JOB_NOT_FOUND"}}
        status = str(row.get("Status") or "").upper()
        terminal = status in {
            "COMPLETED", "COMPLETED_WITH_WARNINGS", "COMPLETED_WITH_ERRORS", "COMPLETED_WITH_SKIPPED",
            "PARTIAL", "FAILED", "LAUNCH_FAILED", "ACTION_REQUIRED", "CANCELLED", "CANCELLED_WITH_ERRORS", "INTERRUPTED",
        }
        age = self._agentless_age_seconds(row.get("UpdatedAt") or row.get("CreatedAt"))
        try:
            configured = int(dict(self.config.get("agentless") or {}).get("stale_after_seconds") or 300)
        except Exception:
            configured = 300
        threshold = max(120, min(configured, 3600))
        stale = (not terminal) and age is not None and age >= threshold
        backend = self._agentless_backend_process(int(row.get("LaunchPid") or row.get("ControllerPid") or 0), str(job_id)) if stale else {
            "pid": int(row.get("LaunchPid") or row.get("ControllerPid") or 0), "alive": None, "owned": None, "detail": "NOT_PROBED"
        }
        return {"stale": bool(stale), "terminal": bool(terminal), "age_seconds": age, "threshold_seconds": threshold, "backend": backend}

    @staticmethod
    def _agentless_retryable_sites_from_manifest(manifest: dict[str, Any]) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []
        sites = manifest.get("Sites") if isinstance(manifest.get("Sites"), dict) else {}
        for code, raw in sites.items():
            if not isinstance(raw, dict):
                continue
            dispatch = str(raw.get("DispatchStatus") or "").upper()
            retryable = bool(raw.get("RetryEligible")) or dispatch in {"FAILED", "DEFERRED"}
            if not retryable:
                continue
            detail = str(
                raw.get("ExecutionErrorDetail")
                or raw.get("DispatchError")
                or raw.get("LastPollError")
                or raw.get("LastError")
                or raw.get("Error")
                or "Reintento disponible"
            )[:1800]
            rows.append({
                "SiteCode": str(raw.get("SiteCode") or code),
                "SiteName": str(raw.get("SiteName") or ""),
                "ServerIP": str(raw.get("ServerIP") or ""),
                "DispatchStatus": dispatch or "-",
                "RemoteState": str(raw.get("RemoteState") or "-").upper(),
                "TelemetryState": str(raw.get("TelemetryState") or "-").upper(),
                "Detail": detail,
            })
        rows.sort(key=lambda row: (row["SiteCode"], row["SiteName"]))
        return rows

    def agentless_action_state(self, job_id: str, manifest: dict[str, Any] | None = None) -> dict[str, Any]:
        current = dict(manifest or self.agentless_job(job_id) or {})
        if not current:
            raise FileNotFoundError(f"Job Agentless no encontrado: {job_id}")
        status = str(current.get("Status") or "UNKNOWN").upper()
        action_code = str(current.get("ActionRequired") or "").upper()
        retryable = self._agentless_retryable_sites_from_manifest(current)
        sites = current.get("Sites") if isinstance(current.get("Sites"), dict) else {}
        active_states = {"PENDING", "PENDING_RETRY", "DISPATCHING", "DISPATCHED", "RUNNING", "RESULT_READY", "RESULT_READY_WITH_ERRORS", "COLLECTING"}
        active = []
        for code, raw in sites.items():
            if not isinstance(raw, dict):
                continue
            remote = str(raw.get("RemoteState") or raw.get("TaskState") or "").upper()
            if remote in active_states:
                active.append(str(raw.get("SiteCode") or code))
        protocol = current.get("ProtocolPreflight") if isinstance(current.get("ProtocolPreflight"), dict) else {}
        return {
            "job_id": job_id,
            "status": status,
            "action_code": action_code or ("RETRY_OR_IGNORE_SITES" if retryable else "REVIEW_REQUIRED"),
            "retryable_count": len(retryable),
            "retryable_sites": retryable,
            "active_count": len(active),
            "active_sites": active[:50],
            "preflight_target": int(protocol.get("TargetCount") or 0),
            "preflight_passed": int(protocol.get("PassedCount") or 0),
            "preflight_failed": int(protocol.get("FailedCount") or 0),
            "can_retry": status == "ACTION_REQUIRED" and bool(retryable) and not active,
            "can_omit": status == "ACTION_REQUIRED" and bool(retryable) and not active,
            "can_approve": status == "ACTION_REQUIRED" and not active,
            "can_reconcile": status == "ACTION_REQUIRED" and not retryable and not active,
        }

    def _record_resolution_history(self, manifest: dict[str, Any], *, action: str, actor: str, sites: list[str], detail: str = "") -> None:
        history = manifest.get("ResolutionHistory") if isinstance(manifest.get("ResolutionHistory"), list) else []
        history = list(history)[-19:]
        history.append({
            "Action": str(action),
            "Actor": str(actor or "administrator"),
            "Sites": list(sites),
            "Detail": str(detail or "")[:1500],
            "At": dt.datetime.now().astimezone().isoformat(timespec="seconds"),
        })
        manifest["ResolutionHistory"] = history

    def omit_agentless_sites(self, job_id: str, *, actor: str, site_codes: list[str], approve_current: bool = False) -> dict[str, Any]:
        store = FanoutStore(self.agentless_fanout_dir)
        manifest = store.read(job_id)
        if str(manifest.get("Status") or "").upper() != "ACTION_REQUIRED":
            raise RuntimeError(f"JOB_NOT_ACTION_REQUIRED: estado actual={manifest.get('Status') or '-'}")
        state = self.agentless_action_state(job_id, manifest)
        if int(state.get("active_count") or 0) > 0:
            raise RuntimeError("JOB_HAS_ACTIVE_SITES: no se puede aprobar/omitir mientras existan destinos activos.")
        retryable = {str(row.get("SiteCode") or "") for row in state.get("retryable_sites") or []}
        requested = retryable if approve_current else {str(code).strip() for code in site_codes if str(code).strip()}
        selected = sorted(retryable & requested)
        if retryable and not selected:
            raise ValueError("Seleccione al menos una sede reintentable para omitir.")
        sites = manifest.get("Sites") if isinstance(manifest.get("Sites"), dict) else {}
        stamp = dt.datetime.now().astimezone().isoformat(timespec="seconds")
        for code in selected:
            raw = sites.get(code)
            if not isinstance(raw, dict):
                continue
            entry = dict(raw)
            previous = str(entry.get("ExecutionErrorDetail") or entry.get("DispatchError") or entry.get("LastError") or entry.get("Error") or "")[:1500]
            entry.update({
                "DispatchStatus": "IGNORED",
                "RemoteState": "IGNORED",
                "RetryEligible": False,
                "IgnoredAt": stamp,
                "IgnoredBy": str(actor or "administrator"),
                "IgnoredReason": "APPROVED_CURRENT_RESULT" if approve_current else "OPERATOR_SKIPPED_UNRESOLVED",
                "IgnoredPreviousError": previous,
                "Error": None,
                "LastError": None,
            })
            sites[code] = entry
            try:
                store.clear_site_skip(job_id, code)
            except Exception:
                pass
        manifest["Sites"] = sites
        recalc_manifest(manifest)
        remaining = self._agentless_retryable_sites_from_manifest(manifest)
        manifest["UnresolvedSitesIgnored"] = not bool(remaining)
        manifest["ActionRequired"] = None if not remaining else "RETRY_OR_IGNORE_SITES"
        manifest["ResolutionDecision"] = "APPROVED_CURRENT_RESULT" if approve_current else "OMIT_SELECTED_SITES"
        manifest["ResolutionBy"] = str(actor or "administrator")
        manifest["ResolutionAt"] = stamp
        self._record_resolution_history(
            manifest,
            action=manifest["ResolutionDecision"],
            actor=actor,
            sites=selected,
            detail=("Resultado actual aprobado; pendientes aceptados como omitidos." if approve_current else "Sedes seleccionadas omitidas por decisión operativa."),
        )
        finalize_manifest(manifest, store)
        final = store.read(job_id)
        return {
            "JobId": job_id,
            "Status": str(final.get("Status") or "UNKNOWN"),
            "Omitted": len(selected),
            "RemainingRetryable": len(self._agentless_retryable_sites_from_manifest(final)),
            "Decision": manifest["ResolutionDecision"],
        }

    def reconcile_agentless_action_required(self, job_id: str, *, actor: str) -> dict[str, Any]:
        store = FanoutStore(self.agentless_fanout_dir)
        manifest = store.read(job_id)
        if str(manifest.get("Status") or "").upper() != "ACTION_REQUIRED":
            raise RuntimeError(f"JOB_NOT_ACTION_REQUIRED: estado actual={manifest.get('Status') or '-'}")
        state = self.agentless_action_state(job_id, manifest)
        if int(state.get("active_count") or 0) > 0:
            raise RuntimeError("JOB_HAS_ACTIVE_SITES: el Job todavía tiene destinos activos.")
        if int(state.get("retryable_count") or 0) > 0:
            raise RuntimeError("JOB_HAS_RETRYABLE_SITES: use Reintentar, Omitir o Aprobar resultado actual.")
        recalc_manifest(manifest)
        manifest["ActionRequired"] = None
        manifest["ResolutionDecision"] = "RECONCILE_STATE"
        manifest["ResolutionBy"] = str(actor or "administrator")
        manifest["ResolutionAt"] = dt.datetime.now().astimezone().isoformat(timespec="seconds")
        self._record_resolution_history(manifest, action="RECONCILE_STATE", actor=actor, sites=[], detail="Estado ACTION_REQUIRED recalculado sin sedes reintentables.")
        finalize_manifest(manifest, store)
        final = store.read(job_id)
        return {"JobId": job_id, "Status": str(final.get("Status") or "UNKNOWN"), "Decision": "RECONCILE_STATE"}

    def retry_agentless_unresolved(self, job_id: str, *, actor: str) -> dict[str, Any]:
        if self.ssh_vault is None:
            raise RuntimeError("SSH_VAULT_UNAVAILABLE: no se puede reintentar sin el Vault SSH.")
        store = FanoutStore(self.agentless_fanout_dir)
        manifest = store.read(job_id)
        if str(manifest.get("Status") or "").upper() != "ACTION_REQUIRED":
            raise RuntimeError(f"JOB_NOT_ACTION_REQUIRED: estado actual={manifest.get('Status') or '-'}")
        state = self.agentless_action_state(job_id, manifest)
        retryable_rows = list(state.get("retryable_sites") or [])
        if not retryable_rows:
            raise RuntimeError("NO_RETRYABLE_SITES: use Recalcular estado.")
        if int(state.get("active_count") or 0) > 0:
            raise RuntimeError("JOB_HAS_ACTIVE_SITES: no se iniciará un segundo backend mientras existan destinos activos.")
        identity = manifest.get("SshIdentity") if isinstance(manifest.get("SshIdentity"), dict) else {}
        identity_id = str(identity.get("id") or "").strip()
        if not identity_id:
            raise RuntimeError("SSH_IDENTITY_NOT_RECORDED: configure una identidad SSH antes de reintentar.")
        identity_row = next((row for row in self.ssh_vault.list_identities() if str(row.get("id") or "") == identity_id), None)
        if not isinstance(identity_row, dict):
            raise RuntimeError("SSH_IDENTITY_NOT_FOUND: la identidad usada por el Job ya no existe en el Vault.")
        username = str(identity.get("username") or identity_row.get("default_username") or "").strip()
        try:
            port = int(identity.get("port") or identity_row.get("default_port") or 22)
        except (TypeError, ValueError):
            port = 22
        if not username:
            raise RuntimeError("SSH_GLOBAL_PROFILE_REQUIRED: la identidad no tiene usuario SSH global guardado.")
        transport_profile = str(manifest.get("TransportProfile") or "SFTP_SSH_SFTP")
        sites_file = self.agentless_root / "jobs" / job_id / "selected-sites.csv"
        if not sites_file.is_file():
            raise FileNotFoundError(f"JOB_SITES_FILE_MISSING: {sites_file}")
        key_path = self.agentless_key_dir / f"{job_id}.retry-{uuid.uuid4().hex[:8]}.key"
        try:
            key_path.write_bytes(self.ssh_vault.private_key_bytes(identity_id))
            harden_openssh_private_key_permissions(key_path)
            _cfg, config_path = self._agentless_job_config(
                job_id=job_id,
                key_path=key_path,
                transport_profile=transport_profile,
                ssh_username=username,
                ssh_port=port,
            )
            export_format = str(manifest.get("ExportFormat") or "xlsx").lower()
            if export_format not in {"xlsx", "csv", "json", "txt", "none"}:
                export_format = "xlsx"
            log_path = self.agentless_log_dir / f"{dt.datetime.now().strftime('%Y%m%d-%H%M%S')}-{job_id}-retry.log"
            command = [
                sys.executable,
                str(self.app_root / "gideon_admin.py"),
                "fanout",
                "--config", str(config_path),
                "--sites-file", str(sites_file),
                "retry-failed", job_id,
                "--export", export_format,
            ]
            stamp = dt.datetime.now().astimezone().isoformat(timespec="seconds")
            manifest = store.read(job_id)
            manifest["Status"] = "RETRY_QUEUED"
            manifest["ActionRequired"] = None
            manifest["ResolutionState"] = "RETRY_QUEUED"
            manifest["ResolutionAction"] = "RETRY_FAILED"
            manifest["ResolutionRequestedBy"] = str(actor or "administrator")
            manifest["ResolutionRequestedAt"] = stamp
            manifest["ResolutionLogPath"] = str(log_path)
            self._record_resolution_history(manifest, action="RETRY_FAILED", actor=actor, sites=[str(row.get("SiteCode") or "") for row in retryable_rows], detail="Reintento solicitado desde WebAdmin.")
            store.write(manifest)
            log_path.parent.mkdir(parents=True, exist_ok=True)
            with log_path.open("a", encoding="utf-8", buffering=1) as log_file:
                log_file.write(f"GIDEON WebAdmin resolution retry · {job_id}\n")
                log_file.write(f"Actor={actor} Retryable={len(retryable_rows)} User={username} Port={port}\n")
                log_file.flush()
                process = subprocess.Popen(
                    command,
                    cwd=str(self.app_root),
                    stdin=subprocess.DEVNULL,
                    stdout=log_file,
                    stderr=subprocess.STDOUT,
                    close_fds=(os.name != "nt"),
                    env={**os.environ, "PYTHONUTF8": "1", "PYTHONIOENCODING": "utf-8", "GIDEON_FORCE_UTF8_STDIO": "1", "GIDEON_WEB_AGENTLESS": "1"},
                    **self._hidden_process_kwargs(),
                )
            latest = store.read(job_id)
            if str(latest.get("Status") or "").upper() == "RETRY_QUEUED":
                latest["Status"] = "RETRYING"
            latest["ResolutionState"] = "RETRYING"
            latest["ResolutionPid"] = int(process.pid)
            latest["UpdatedAt"] = dt.datetime.now().astimezone().isoformat(timespec="seconds")
            store.write(latest)
        except Exception:
            try:
                key_path.unlink(missing_ok=True)
            except OSError:
                pass
            try:
                current = store.read(job_id)
                current["Status"] = "ACTION_REQUIRED"
                current["ActionRequired"] = "RETRY_OR_IGNORE_SITES"
                current["ResolutionState"] = "RETRY_START_FAILED"
                current["ResolutionError"] = "No se pudo iniciar el reintento desde WebAdmin."
                store.write(current)
            except Exception:
                pass
            raise

        def _watch_resolution() -> None:
            try:
                rc = process.wait()
                try:
                    current = store.read(job_id)
                    current["ResolutionReturnCode"] = int(rc)
                    current["ResolutionFinishedAt"] = dt.datetime.now().astimezone().isoformat(timespec="seconds")
                    st = str(current.get("Status") or "").upper()
                    if st == "ACTION_REQUIRED":
                        current["ResolutionState"] = "ACTION_REQUIRED"
                    elif st.startswith("COMPLETED") or st in {"CANCELLED", "CANCELLED_WITH_ERRORS"}:
                        current["ResolutionState"] = "COMPLETED"
                    elif rc != 0 and st in {"RETRY_QUEUED", "RETRYING", "DISPATCHING", "MONITORING", "COLLECTING"}:
                        current["Status"] = "ACTION_REQUIRED"
                        current["ActionRequired"] = "RETRY_OR_IGNORE_SITES"
                        current["ResolutionState"] = "FAILED"
                        current["ResolutionError"] = f"RETRY_BACKEND_EXITED: rc={rc}; log={log_path}"
                    else:
                        current["ResolutionState"] = "COMPLETED" if rc == 0 else "FAILED"
                    store.write(current)
                except Exception:
                    pass
            finally:
                try:
                    key_path.unlink(missing_ok=True)
                except OSError:
                    pass
        threading.Thread(target=_watch_resolution, name=f"gideon-agentless-resolution-{job_id}", daemon=True).start()
        return {
            "JobId": job_id,
            "Status": "RETRYING",
            "Retryable": len(retryable_rows),
            "Pid": int(process.pid),
            "LogPath": str(log_path),
        }

    def recover_agentless(self, job_id: str, *, actor: str) -> dict[str, Any]:
        row = self.agentless_job(job_id)
        if row is None:
            raise FileNotFoundError("Job Agentless no encontrado.")
        runtime = self.agentless_runtime_state(job_id, row)
        if runtime.get("terminal"):
            return {"JobId": job_id, "Status": str(row.get("Status") or ""), "Recovered": False, "Reason": "ALREADY_TERMINAL"}
        if not runtime.get("stale"):
            raise RuntimeError("JOB_NOT_STALE: el Job todavía está dentro de la ventana normal de actividad.")
        backend = runtime.get("backend") if isinstance(runtime.get("backend"), dict) else {}
        pid = int(backend.get("pid") or 0)
        if backend.get("alive") and not backend.get("owned"):
            raise RuntimeError("STALE_JOB_PID_NOT_OWNED: el PID fue reutilizado o no pertenece a este Job; no se tocará.")
        terminated = False
        if backend.get("alive") and backend.get("owned"):
            if os.name != "nt":
                raise RuntimeError("STALE_JOB_BACKEND_STILL_RUNNING: use Cancelar Job; recuperación forzada en Unix no está habilitada.")
            taskkill = shutil.which("taskkill.exe") or str(Path(os.environ.get("SystemRoot") or r"C:\Windows") / "System32" / "taskkill.exe")
            cp = subprocess.run(
                [taskkill, "/PID", str(pid), "/T", "/F"], capture_output=True, text=True,
                timeout=15, shell=False, encoding="utf-8", errors="replace", **self._hidden_process_kwargs(),
            )
            if cp.returncode != 0:
                raise RuntimeError("STALE_JOB_TERMINATE_FAILED: " + redact_error((cp.stdout or "") + " " + (cp.stderr or "")))
            terminated = True

        stamp = dt.datetime.now().astimezone().isoformat(timespec="seconds")
        reason = f"STALE_AGENTLESS_JOB_RECOVERED: sin actualización durante {int(runtime.get('age_seconds') or 0)}s; backend={backend.get('detail') or '-'}"
        row = dict(row)
        row["Status"] = "INTERRUPTED"
        row["FinishedAt"] = stamp
        row["RecoveredAt"] = stamp
        row["RecoveredBy"] = str(actor or "administrator")
        row["RecoveryReason"] = reason
        sites = row.get("Sites") if isinstance(row.get("Sites"), dict) else {}
        terminal_sites = {"COLLECTED", "COLLECTED_WITH_ERRORS", "FAILED", "CANCELLED", "CANCELLED_NOT_STARTED", "CANCEL_FAILED", "INTERRUPTED", "DISPATCH_FAILED", "IGNORED"}
        for code, raw in list(sites.items()):
            if not isinstance(raw, dict):
                continue
            entry = dict(raw)
            state = str(entry.get("RemoteState") or "").upper()
            if state not in terminal_sites:
                entry["RemoteState"] = "INTERRUPTED"
                entry.setdefault("LastError", reason)
                entry["InterruptedAt"] = stamp
            sites[code] = entry
        row["Sites"] = sites
        store = FanoutStore(self.agentless_fanout_dir)
        store.write(row)
        store.clear_cancel_request(job_id)
        try:
            (self.agentless_key_dir / f"{job_id}.key").unlink(missing_ok=True)
        except Exception:
            pass
        return {"JobId": job_id, "Status": "INTERRUPTED", "Recovered": True, "BackendTerminated": bool(terminated), "Reason": reason}

    def _finalize_agentless_cancel(self, job_id: str, row: dict[str, Any], *, actor: str, reason: str) -> dict[str, Any]:
        stamp = dt.datetime.now().astimezone().isoformat(timespec="seconds")
        current = dict(row)
        current["Status"] = "CANCELLED"
        current["FinishedAt"] = stamp
        current["CancelledAt"] = stamp
        current["CancelledBy"] = str(actor or "administrator")
        current["CancelReason"] = str(reason or "CANCELLED_BY_WEBADMIN")
        sites = current.get("Sites") if isinstance(current.get("Sites"), dict) else {}
        terminal_sites = {
            "COLLECTED", "COLLECTED_WITH_ERRORS", "FAILED", "CANCELLED",
            "CANCELLED_NOT_STARTED", "CANCEL_FAILED", "INTERRUPTED",
            "DISPATCH_FAILED", "IGNORED",
        }
        for code, raw in list(sites.items()):
            if not isinstance(raw, dict):
                continue
            entry = dict(raw)
            remote = str(entry.get("RemoteState") or "").upper()
            if remote not in terminal_sites:
                entry["RemoteState"] = "CANCELLED_NOT_STARTED"
                entry.setdefault("LastError", current["CancelReason"])
                entry["CancelledAt"] = stamp
            if str(entry.get("TransferStatus") or "").upper() in {"", "PENDING", "QUEUED", "RUNNING"}:
                entry["TransferStatus"] = "CANCELLED"
            if str(entry.get("CollectionStatus") or "").upper() in {"", "PENDING", "QUEUED", "RUNNING"}:
                entry["CollectionStatus"] = "CANCELLED"
            sites[code] = entry
        current["Sites"] = sites
        store = FanoutStore(self.agentless_fanout_dir)
        store.write(current)
        store.clear_cancel_request(job_id)
        try:
            (self.agentless_key_dir / f"{job_id}.key").unlink(missing_ok=True)
        except Exception:
            pass
        return {"JobId": job_id, "Status": "CANCELLED", "Cancelled": True, "Reason": current["CancelReason"]}

    def cancel_agentless(self, job_id: str, *, actor: str) -> dict[str, Any]:
        row = self.agentless_job(job_id)
        if row is None:
            raise FileNotFoundError("Job Agentless no encontrado.")
        status = str(row.get("Status") or "").upper()
        terminal = {
            "COMPLETED", "COMPLETED_WITH_WARNINGS", "COMPLETED_WITH_ERRORS",
            "COMPLETED_WITH_SKIPPED", "PARTIAL", "FAILED", "LAUNCH_FAILED",
            "ACTION_REQUIRED", "CANCELLED", "CANCELLED_WITH_ERRORS", "INTERRUPTED",
        }
        if status in terminal:
            return {"JobId": job_id, "Status": status, "Cancelled": status.startswith("CANCELLED"), "Reason": "ALREADY_TERMINAL"}
        if status in {"PREFLIGHT_QUEUED", "PREFLIGHT_RUNNING", "LAUNCHING"} and not int(row.get("LaunchPid") or row.get("ControllerPid") or 0):
            return self._finalize_agentless_cancel(
                job_id, row, actor=actor, reason="CANCELLED_BEFORE_BACKEND_START"
            )
        runtime = self.agentless_runtime_state(job_id, row)
        backend = runtime.get("backend") if isinstance(runtime.get("backend"), dict) else {}
        if runtime.get("stale") and backend.get("alive") is False:
            return self._finalize_agentless_cancel(
                job_id, row, actor=actor,
                reason=f"CANCELLED_STALE_BACKEND_MISSING: {backend.get('detail') or 'PROCESS_NOT_FOUND'}",
            )
        result = FanoutStore(self.agentless_fanout_dir).request_cancel(
            job_id, requested_by=actor, source="web-admin"
        )
        if isinstance(result, dict):
            result = dict(result)
            result.setdefault("Status", "CANCEL_REQUESTED")
            result["CancelMode"] = "GRACEFUL_JOB_SCOPED"
            return result
        return {"JobId": job_id, "Status": "CANCEL_REQUESTED", "CancelMode": "GRACEFUL_JOB_SCOPED"}


    def dispatch_restart(self, *, agent_id: str, actor: str, delay_seconds: int = 60) -> dict[str, Any]:
        """Queue a signed, built-in restart action for one Agent.

        The reboot is scheduled with a delay so the worker can upload its result
        before the operating system goes offline.  No arbitrary command text is
        accepted from the browser.
        """
        aid = str(agent_id or "").strip()
        if not aid:
            raise ValueError("AgentId requerido.")
        delay = max(30, min(int(delay_seconds or 60), 600))
        private_key, private_key_bytes, signing_key_id = self._signing_material()
        all_agents = {str(x.get("agent_id") or ""): x for x in self.gateway.agents()}
        row = all_agents.get(aid)
        if row is None:
            raise RuntimeError(f"Agent no encontrado: {aid}")
        source = self.app_root / "gideon_agent_remote_restart.py"
        if not source.is_file():
            raise RuntimeError("RESTART_SUPPORT_MISSING: falta gideon_agent_remote_restart.py")
        self._preflight_agent_operation(row, script_paths=[source])
        profile_id = str(row.get("host_profile") or "").strip()
        if not profile_id:
            raise RuntimeError(f"Agent {aid} no reporta HostProfile.")
        profile = self.profiles.get(profile_id)
        job_id = "device-restart-" + dt.datetime.now().strftime("%Y%m%d-%H%M%S") + "-" + uuid.uuid4().hex[:6]
        stage = self.job_dir / job_id / "packages" / aid
        package_dir = stage / "package"
        (package_dir / "scripts").mkdir(parents=True, exist_ok=True)
        target_script = package_dir / "scripts" / source.name
        shutil.copy2(source, target_script)
        support = ("gideon_agent_worker.py", "gideon_agent_job_contract.py", "gideon_script_host.ps1")
        missing = [name for name in support if not (self.app_root / name).is_file()]
        if missing:
            raise RuntimeError("RESTART_PACKAGE_SUPPORT_MISSING: " + ", ".join(missing))
        for name in support:
            shutil.copy2(self.app_root / name, package_dir / name)
        (package_dir / "host_profile.json").write_text(json.dumps(profile, ensure_ascii=False, indent=2), encoding="utf-8")
        manifest = build_job_envelope(
            job_id=job_id, job_type="SCRIPT_EXECUTION", agent_id=aid,
            agent_name=str(row.get("agent_name") or row.get("hostname") or ""),
            hostname=str(row.get("hostname") or ""), site_code=str(row.get("site_code") or ""),
            site_name=str(row.get("site_name") or ""), host_profile=profile,
            continue_on_error=False, default_timeout_seconds=30, ttl_seconds=1800,
            scripts=[{
                "Index": 1, "Name": source.name, "Path": "scripts/" + source.name,
                "ResultName": "01_device_restart.json", "Sha256": sha256_file(target_script),
                "TimeoutSeconds": 30, "Arguments": ["--delay", str(delay)],
                "Parameters": {}, "Environment": {},
            }],
            required_capabilities=("SIGNED_JOBS", "SCRIPT_EXECUTION", "RESULT_UPLOAD"),
        )
        (package_dir / "agent_job.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        archive = stage / f"{aid}.agent.zip"
        with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as zf:
            for file in sorted(package_dir.rglob("*")):
                if file.is_file():
                    zf.write(file, file.relative_to(package_dir).as_posix())
        signature = sign_file_bytes(archive, private_key_bytes) if private_key_bytes is not None else sign_file(archive, private_key)
        response = self.gateway.client().upload_file(
            "PUT", f"/api/v1/admin/jobs/{job_id}/{aid}", archive,
            headers={
                "X-GIDEON-Site-Code": str(row.get("site_code") or ""),
                "X-GIDEON-Site-Name": str(row.get("site_name") or ""),
                "X-GIDEON-SHA256": sha256_file(archive),
                "X-GIDEON-Signature": str(signature),
                "X-GIDEON-Key-Id": signing_key_id,
                "X-GIDEON-Host-Profile": profile_id,
            }, timeout=120,
        )
        meta = {
            "JobId": job_id, "CreatedAt": dt.datetime.now().astimezone().isoformat(timespec="seconds"),
            "CreatedBy": actor, "AgentId": aid, "Action": "RESTART", "DelaySeconds": delay,
            "Gateway": response,
        }
        (self.job_dir / job_id / "web_job.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
        return meta

    def dispatch(self, *, agent_ids: list[str], script_ids: list[int], actor: str, continue_on_error: bool = True) -> dict[str, Any]:
        if not agent_ids:
            raise ValueError("Seleccione al menos un Agent.")
        if not script_ids:
            raise ValueError("Seleccione al menos un script aprobado.")
        private_key, private_key_bytes, signing_key_id = self._signing_material()
        script_rows = []
        scripts: list[Path] = []
        for sid in script_ids:
            row = self.security.get_script(int(sid))
            if row is None or str(row.get("approval_state") or "").upper() != "APPROVED":
                raise RuntimeError(f"Script {sid} no está APPROVED.")
            path = Path(str(row.get("stored_path") or ""))
            if not path.is_file() or sha256_file(path).lower() != str(row.get("sha256") or "").lower():
                raise RuntimeError(f"Integridad inválida para script {row.get('name') or sid}.")
            script_rows.append(row)
            scripts.append(path)

        all_agents = {str(x.get("agent_id") or ""): x for x in self.gateway.agents()}
        selected = []
        for aid in agent_ids:
            row = all_agents.get(str(aid))
            if row is None:
                raise RuntimeError(f"Agent no encontrado: {aid}")
            self._preflight_agent_operation(row, script_paths=scripts)
            selected.append(row)

        job_id = dt.datetime.now().strftime("%Y%m%d-%H%M%S") + "-web-" + uuid.uuid4().hex[:6]
        package_root = self.job_dir / job_id / "packages"
        package_root.mkdir(parents=True, exist_ok=True)
        cfg = {
            "script_execution": {"default_timeout_seconds": int(self.config.get("default_script_timeout_seconds", 1800))}
        }
        client = self.gateway.client()
        results = []
        for row in selected:
            aid = str(row.get("agent_id") or "")
            profile_id = str(row.get("host_profile") or "").strip()
            if not profile_id:
                raise RuntimeError(f"Agent {aid} no reporta HostProfile.")
            profile = self.profiles.get(profile_id)
            site = Site(
                code=str(row.get("site_code") or ""), name=str(row.get("site_name") or ""), ip="",
                platform=str(row.get("platform") or ""), host_profile=profile_id, agent_id=aid,
                agent_name=str(row.get("agent_name") or row.get("hostname") or ""),
                hostname=str(row.get("hostname") or ""), agent_connection="HTTPS/PULL",
                agent_remote_ip=str(row.get("last_remote_ip") or ""), source="gateway-agent",
            )
            package_info = build_agent_package(
                site=site, job_id=job_id, scripts=scripts, cfg=cfg, packages_dir=package_root,
                continue_on_error=continue_on_error, profile=profile, signing_private_key=private_key,
                signing_key_id=signing_key_id, signing_private_key_bytes=private_key_bytes,
            )
            archive = Path(package_info["Archive"])
            response = client.upload_file(
                "PUT", f"/api/v1/admin/jobs/{job_id}/{aid}", archive,
                headers={
                    "X-GIDEON-Site-Code": site.code, "X-GIDEON-Site-Name": site.name,
                    "X-GIDEON-SHA256": sha256_file(archive), "X-GIDEON-Signature": str(package_info["Signature"]),
                    "X-GIDEON-Key-Id": signing_key_id, "X-GIDEON-Host-Profile": profile_id,
                }, timeout=120,
            )
            results.append({"AgentId": aid, "SiteCode": site.code, "Hostname": site.hostname, "Response": response})
        meta = {
            "JobId": job_id, "CreatedAt": dt.datetime.now().astimezone().isoformat(timespec="seconds"),
            "CreatedBy": actor, "Agents": [str(x.get("agent_id")) for x in selected],
            "Scripts": [{"Id": r["id"], "Name": r["name"], "Sha256": r["sha256"]} for r in script_rows],
            "Results": results,
        }
        (self.job_dir / job_id / "web_job.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
        return meta

    @staticmethod
    def _excel_safe_value(value: Any) -> Any:
        """Return an XLSX-safe scalar without allowing formula injection."""
        if value is None:
            return ""
        if isinstance(value, (bool, int, float, dt.datetime, dt.date, dt.time)):
            return value
        if isinstance(value, (dict, list, tuple, set)):
            try:
                value = json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)
            except Exception:
                value = str(value)
        text = str(value)
        if len(text) > 32700:
            text = text[:32680] + "…[truncado]"
        if text.startswith(("=", "+", "-", "@")):
            text = "'" + text
        return text

    @staticmethod
    def _xlsx_format_sheet(ws: Any, *, freeze: str = "A2") -> None:
        if ws.max_row >= 1:
            for cell in ws[1]:
                cell.font = Font(bold=True)
            ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = freeze
        for idx, column in enumerate(ws.iter_cols(1, max(1, ws.max_column)), 1):
            width = 11
            for cell in list(column)[:250]:
                value = "" if cell.value is None else str(cell.value)
                width = max(width, min(48, len(value) + 2))
            ws.column_dimensions[get_column_letter(idx)].width = width

    @staticmethod
    def _flatten_result_payload(payload: Any) -> tuple[list[dict[str, Any]], dict[str, Any]]:
        """Normalize common GIDEON result envelopes to rows plus result metadata."""
        meta: dict[str, Any] = {}
        values: Any = payload
        if isinstance(payload, dict):
            for key in ("Success", "Error", "Message", "Script", "ScriptName", "Timestamp", "GeneratedAt"):
                if key in payload:
                    meta[key] = payload.get(key)
            for candidate in ("Data", "Rows", "Results", "Items", "Result"):
                if candidate in payload:
                    values = payload.get(candidate)
                    break
            else:
                # A plain dictionary is itself a useful result row.
                values = payload
        if values is None:
            return [], meta
        if not isinstance(values, list):
            values = [values]
        rows: list[dict[str, Any]] = []
        for value in values:
            if isinstance(value, dict):
                rows.append(dict(value))
            elif isinstance(value, (list, tuple)):
                rows.append({f"Column{idx + 1}": item for idx, item in enumerate(value)})
            else:
                rows.append({"Value": value})
        return rows, meta

    def _agentless_result_candidates(self, job_id: str, manifest: dict[str, Any]) -> list[Path]:
        """Discover only files belonging to an Agentless Job under GIDEON-owned roots."""
        clean = "".join(ch for ch in str(job_id) if ch.isalnum() or ch in "._-")
        if not clean or clean != str(job_id):
            raise ValueError("JobId inválido.")
        root = self.agentless_root.resolve()
        allowed_suffixes = {".json", ".csv", ".xlsx", ".zip"}
        candidates: dict[str, Path] = {}

        def add_file(path_value: Any) -> None:
            if not path_value:
                return
            try:
                path = Path(str(path_value))
                if not path.is_absolute():
                    path = root / path
                resolved = path.resolve()
                resolved.relative_to(root)
            except Exception:
                return
            if resolved.is_file() and resolved.suffix.lower() in allowed_suffixes:
                candidates[str(resolved).casefold()] = resolved

        # Prefer explicit paths recorded by the runtime manifest.
        path_keys = {
            "resultpath", "resultfile", "collectionpath", "collectedpath", "localresultpath",
            "archivepath", "exportpath", "outputpath", "reportpath", "resultarchive",
        }
        def walk(value: Any) -> None:
            if isinstance(value, dict):
                for key, item in value.items():
                    if str(key).replace("_", "").casefold() in path_keys:
                        add_file(item)
                    if isinstance(item, (dict, list, tuple)):
                        walk(item)
            elif isinstance(value, (list, tuple)):
                for item in value:
                    walk(item)
        walk(manifest)

        # Canonical per-job paths.  Search is bounded and only under GIDEON Agentless.
        roots = [
            root / "results" / clean,
            root / "results" / "raw" / clean,
            root / "jobs" / clean / "results",
            root / "jobs" / clean / "collection",
        ]
        for base in roots:
            try:
                if base.is_file():
                    add_file(base)
                elif base.is_dir():
                    for path in base.rglob("*"):
                        if len(candidates) >= 600:
                            break
                        if path.is_file() and path.suffix.lower() in allowed_suffixes:
                            add_file(path)
            except OSError:
                continue

        # Some historical runtimes write the JobId into result filenames directly.
        for base in (root / "results", root / "results" / "raw"):
            if len(candidates) >= 600 or not base.is_dir():
                continue
            try:
                for path in base.rglob(f"*{clean}*"):
                    if len(candidates) >= 600:
                        break
                    if path.is_file():
                        add_file(path)
                    elif path.is_dir():
                        for child in path.rglob("*"):
                            if len(candidates) >= 600:
                                break
                            if child.is_file() and child.suffix.lower() in allowed_suffixes:
                                add_file(child)
            except OSError:
                continue
        return sorted(candidates.values(), key=lambda x: str(x).casefold())

    @staticmethod
    def _agentless_site_for_source(source: str, sites: dict[str, dict[str, Any]]) -> tuple[str, str]:
        probe = str(source or "").casefold()
        # Match longer codes first to avoid accidental partial matches.
        for code in sorted((str(k) for k in sites.keys()), key=len, reverse=True):
            if code and code.casefold() in probe:
                row = sites.get(code) if isinstance(sites.get(code), dict) else {}
                return code, str(row.get("SiteName") or "")
        return "", ""

    def export_agentless_excel(self, job_id: str) -> bytes:
        """Build a consolidated, read-only XLSX for an Agentless SSH/SFTP Job.

        The workbook always contains operational summary and per-site status.  When
        structured JSON/CSV/XLSX/ZIP result files are available under the GIDEON
        Agentless result roots, their records are consolidated into the Data sheet.
        """
        manifest = self.agentless_job(job_id)
        if manifest is None:
            raise FileNotFoundError("Job Agentless no encontrado.")
        sites_raw = manifest.get("Sites") if isinstance(manifest.get("Sites"), dict) else {}
        sites: dict[str, dict[str, Any]] = {
            str(code): dict(row) for code, row in sites_raw.items() if isinstance(row, dict)
        }
        protocol = manifest.get("ProtocolPreflight") if isinstance(manifest.get("ProtocolPreflight"), dict) else {}
        protocol_sites = protocol.get("Sites") if isinstance(protocol.get("Sites"), dict) else {}
        candidates = self._agentless_result_candidates(job_id, manifest)

        wb = Workbook()
        summary = wb.active
        summary.title = "Resumen"
        summary.append(["Campo", "Valor"])
        summary_rows = [
            ("JobId", job_id),
            ("Estado", manifest.get("Status") or ""),
            ("Método", manifest.get("ExecutionMethod") or "AGENTLESS"),
            ("Perfil transporte", manifest.get("TransportProfile") or ""),
            ("Política de red", manifest.get("NetworkPolicy") or ""),
            ("Formato solicitado", manifest.get("ExportFormat") or "xlsx"),
            ("Creado por", manifest.get("CreatedBy") or ""),
            ("Creado", manifest.get("CreatedAt") or ""),
            ("Actualizado", manifest.get("UpdatedAt") or ""),
            ("Finalizado", manifest.get("FinishedAt") or manifest.get("EndedAt") or ""),
            ("Total sedes", manifest.get("TotalCount") if manifest.get("TotalCount") is not None else len(sites)),
            ("Recolectadas", manifest.get("CollectedCount") if manifest.get("CollectedCount") is not None else sum(1 for r in sites.values() if str(r.get("CollectionStatus") or "").upper() in {"COLLECTED", "SFTP_COLLECTED"})),
            ("Fallidas", manifest.get("FailedCount") if manifest.get("FailedCount") is not None else sum(1 for r in sites.values() if str(r.get("RemoteState") or "").upper() in {"FAILED", "INTERRUPTED", "CANCELLED", "DISPATCH_FAILED"})),
            ("Preflight comprobados", protocol.get("CheckedCount") or 0),
            ("Preflight OK", protocol.get("PassedCount") or 0),
            ("Preflight con observaciones", protocol.get("FailedCount") or 0),
            ("Archivos de resultado detectados", len(candidates)),
        ]
        identity = manifest.get("SshIdentity") if isinstance(manifest.get("SshIdentity"), dict) else {}
        if identity:
            summary_rows.extend([
                ("Identidad SSH", identity.get("label") or ""),
                ("Usuario SSH", identity.get("username") or ""),
                ("Puerto SSH", identity.get("port") or ""),
                ("Huella SSH", identity.get("fingerprint") or ""),
            ])
        for key, value in summary_rows:
            summary.append([self._excel_safe_value(key), self._excel_safe_value(value)])
        self._xlsx_format_sheet(summary)

        site_ws = wb.create_sheet("Sedes")
        site_headers = [
            "JobId", "SiteCode", "SiteName", "ServerIP", "Platform", "SSHPreflight", "SFTPPreflight",
            "RemoteState", "TransferStatus", "CollectionStatus", "DispatchStatus", "TelemetryState",
            "RetryEligible", "Error", "LastError", "ExecutionErrorDetail", "UpdatedAt",
        ]
        site_ws.append(site_headers)
        for code, raw in sites.items():
            pre = protocol_sites.get(code) if isinstance(protocol_sites.get(code), dict) else {}
            ssh = pre.get("SSH") if isinstance(pre.get("SSH"), dict) else {}
            sftp = pre.get("SFTP") if isinstance(pre.get("SFTP"), dict) else {}
            row = {
                "JobId": job_id,
                "SiteCode": raw.get("SiteCode") or code,
                "SiteName": raw.get("SiteName") or "",
                "ServerIP": raw.get("ServerIP") or "",
                "Platform": raw.get("Platform") or "",
                "SSHPreflight": "OK" if ssh.get("Ok") else ("FAIL" if pre else ""),
                "SFTPPreflight": "OK" if sftp.get("Ok") else ("FAIL" if pre and protocol.get("SFTPRequired") else ("N/A" if pre else "")),
                "RemoteState": raw.get("RemoteState") or raw.get("TaskState") or "",
                "TransferStatus": raw.get("TransferStatus") or "",
                "CollectionStatus": raw.get("CollectionStatus") or "",
                "DispatchStatus": raw.get("DispatchStatus") or "",
                "TelemetryState": raw.get("TelemetryState") or "",
                "RetryEligible": raw.get("RetryEligible") if raw.get("RetryEligible") is not None else "",
                "Error": raw.get("Error") or "",
                "LastError": raw.get("LastError") or "",
                "ExecutionErrorDetail": raw.get("ExecutionErrorDetail") or "",
                "UpdatedAt": raw.get("UpdatedAt") or raw.get("FinishedAt") or raw.get("CollectedAt") or "",
            }
            site_ws.append([self._excel_safe_value(row.get(h)) for h in site_headers])
        self._xlsx_format_sheet(site_ws)

        data_ws = wb.create_sheet("Datos")
        data_rows: list[dict[str, Any]] = []
        data_headers: list[str] = ["JobId", "SiteCode", "SiteName", "SourceFile", "SourceEntry", "Success", "Error"]
        file_rows: list[dict[str, Any]] = []
        max_rows = 150000
        max_file_bytes = 40 * 1024 * 1024

        def add_data(records: list[dict[str, Any]], *, source_file: str, source_entry: str = "", meta: dict[str, Any] | None = None, site_code: str = "", site_name: str = "") -> None:
            if len(data_rows) >= max_rows:
                return
            meta = meta or {}
            if not site_code:
                site_code, site_name_guess = self._agentless_site_for_source(source_file + " " + source_entry, sites)
                site_name = site_name or site_name_guess
            success = meta.get("Success")
            error = meta.get("Error") or meta.get("Message") or ""
            for record in records:
                if len(data_rows) >= max_rows:
                    break
                row = {
                    "JobId": job_id,
                    "SiteCode": site_code,
                    "SiteName": site_name,
                    "SourceFile": source_file,
                    "SourceEntry": source_entry,
                    "Success": success if success is not None else "",
                    "Error": error,
                }
                row.update(record)
                data_rows.append(row)
                for key in row.keys():
                    if key not in data_headers:
                        data_headers.append(str(key))

        def parse_json_bytes(blob: bytes, source_file: str, source_entry: str = "") -> None:
            try:
                payload = json.loads(blob.decode("utf-8-sig"))
            except Exception:
                return
            rows, meta = self._flatten_result_payload(payload)
            add_data(rows, source_file=source_file, source_entry=source_entry, meta=meta)

        def parse_csv_text(text: str, source_file: str, source_entry: str = "") -> None:
            try:
                reader = csv.DictReader(io.StringIO(text))
                rows = [dict(row) for _, row in zip(range(max_rows - len(data_rows)), reader)]
            except Exception:
                return
            add_data(rows, source_file=source_file, source_entry=source_entry)

        def parse_xlsx_bytes(blob: bytes, source_file: str, source_entry: str = "") -> None:
            try:
                book = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
            except Exception:
                return
            try:
                for sheet in book.worksheets:
                    iterator = sheet.iter_rows(values_only=True)
                    try:
                        first = next(iterator)
                    except StopIteration:
                        continue
                    headers = [str(v).strip() if v is not None and str(v).strip() else f"Column{idx + 1}" for idx, v in enumerate(first)]
                    rows: list[dict[str, Any]] = []
                    for values in iterator:
                        if len(data_rows) + len(rows) >= max_rows:
                            break
                        if not any(v not in (None, "") for v in values):
                            continue
                        rows.append({headers[idx]: values[idx] if idx < len(values) else "" for idx in range(len(headers))})
                    add_data(rows, source_file=source_file, source_entry=(source_entry + ":" if source_entry else "") + sheet.title)
                    if len(data_rows) >= max_rows:
                        break
            finally:
                try:
                    book.close()
                except Exception:
                    pass

        for path in candidates:
            if len(data_rows) >= max_rows:
                break
            try:
                size = int(path.stat().st_size)
                digest = sha256_file(path)
                rel = str(path.resolve().relative_to(self.agentless_root.resolve()))
            except Exception:
                continue
            file_rows.append({"RelativePath": rel, "Type": path.suffix.lower().lstrip("."), "Bytes": size, "SHA256": digest})
            if size <= 0 or size > max_file_bytes:
                continue
            try:
                if path.suffix.lower() == ".json":
                    parse_json_bytes(path.read_bytes(), rel)
                elif path.suffix.lower() == ".csv":
                    blob = path.read_bytes()
                    try:
                        text = blob.decode("utf-8-sig")
                    except UnicodeDecodeError:
                        text = blob.decode("latin-1", errors="replace")
                    parse_csv_text(text, rel)
                elif path.suffix.lower() == ".xlsx":
                    parse_xlsx_bytes(path.read_bytes(), rel)
                elif path.suffix.lower() == ".zip":
                    with zipfile.ZipFile(path) as zf:
                        for info in zf.infolist():
                            if len(data_rows) >= max_rows:
                                break
                            name = str(info.filename or "")
                            if info.is_dir() or info.file_size <= 0 or info.file_size > max_file_bytes:
                                continue
                            # Prevent zip-slip style path confusion even though we never extract.
                            member = Path(name)
                            if member.is_absolute() or ".." in member.parts:
                                continue
                            suffix = member.suffix.lower()
                            if suffix not in {".json", ".csv", ".xlsx"}:
                                continue
                            try:
                                blob = zf.read(info)
                            except Exception:
                                continue
                            if suffix == ".json":
                                parse_json_bytes(blob, rel, name)
                            elif suffix == ".csv":
                                try:
                                    text = blob.decode("utf-8-sig")
                                except UnicodeDecodeError:
                                    text = blob.decode("latin-1", errors="replace")
                                parse_csv_text(text, rel, name)
                            else:
                                parse_xlsx_bytes(blob, rel, name)
            except (OSError, ValueError, zipfile.BadZipFile):
                continue

        data_ws.append(data_headers)
        for row in data_rows:
            data_ws.append([self._excel_safe_value(row.get(header, "")) for header in data_headers])
        if not data_rows:
            data_ws.append([job_id, "", "", "", "", "", "No se encontraron filas estructuradas; revise las hojas Sedes y Archivos para el estado y trazabilidad del Job."])
        self._xlsx_format_sheet(data_ws)

        files_ws = wb.create_sheet("Archivos")
        files_headers = ["RelativePath", "Type", "Bytes", "SHA256"]
        files_ws.append(files_headers)
        for row in file_rows:
            files_ws.append([self._excel_safe_value(row.get(h, "")) for h in files_headers])
        if not file_rows:
            files_ws.append(["No se detectaron archivos estructurados vinculados al Job.", "", "", ""])
        self._xlsx_format_sheet(files_ws)

        errors_ws = wb.create_sheet("Observaciones")
        error_headers = ["SiteCode", "SiteName", "ServerIP", "RemoteState", "TransferStatus", "CollectionStatus", "Error"]
        errors_ws.append(error_headers)
        problem_count = 0
        for code, raw in sites.items():
            error = str(raw.get("Error") or raw.get("LastError") or raw.get("ExecutionErrorDetail") or "").strip()
            state = str(raw.get("RemoteState") or raw.get("TaskState") or "").upper()
            if error or state in {"FAILED", "INTERRUPTED", "CANCELLED", "DISPATCH_FAILED", "SKIPPED"}:
                problem_count += 1
                values = [
                    raw.get("SiteCode") or code, raw.get("SiteName") or "", raw.get("ServerIP") or "",
                    raw.get("RemoteState") or raw.get("TaskState") or "", raw.get("TransferStatus") or "",
                    raw.get("CollectionStatus") or "", error,
                ]
                errors_ws.append([self._excel_safe_value(v) for v in values])
        if problem_count == 0:
            errors_ws.append(["", "", "", "", "", "", "Sin observaciones registradas."])
        self._xlsx_format_sheet(errors_ws)

        # Keep the user's analysis sheets first and make the output self-describing.
        summary["A1"] = "Campo"
        summary["B1"] = "Valor"
        summary.append(["Filas consolidadas en Datos", len(data_rows)])
        summary.append(["Límite de filas del exportador", max_rows])
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    def export_excel(self, job_id: str) -> bytes:
        rows = self.gateway.job(job_id)
        if not rows:
            raise FileNotFoundError("Job no encontrado.")
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        headers: list[str] = ["JobId", "SiteCode", "SiteName", "AgentId", "Hostname", "Script", "Success", "Error"]
        data_rows: list[dict[str, Any]] = []
        with tempfile.TemporaryDirectory(prefix="gideon-web-result-") as tmp:
            root = Path(tmp)
            for item in rows:
                if not bool(item.get("ResultReady")):
                    continue
                aid = str(item.get("agent_id") or item.get("AgentId") or "")
                archive = root / f"{aid}.zip"
                self.gateway.download_result(job_id, aid, archive)
                with zipfile.ZipFile(archive) as zf:
                    names = [n for n in zf.namelist() if n.startswith("collection/scripts/") and n.lower().endswith(".json")]
                    if not names:
                        names = [n for n in zf.namelist() if n.startswith("scripts/") and n.lower().endswith(".json")]
                    for name in names:
                        try:
                            payload = json.loads(zf.read(name).decode("utf-8-sig"))
                        except Exception:
                            continue
                        values = payload.get("Data") if isinstance(payload, dict) else None
                        if values is None:
                            values = []
                        if not isinstance(values, list):
                            values = [values]
                        if not values:
                            values = [{}]
                        for value in values:
                            record = value if isinstance(value, dict) else {"Value": value}
                            prefix = {
                                "JobId": job_id,
                                "SiteCode": str(item.get("site_code") or item.get("SiteCode") or ""),
                                "SiteName": str(item.get("site_name") or item.get("SiteName") or ""),
                                "AgentId": aid,
                                "Hostname": str(item.get("Hostname") or item.get("hostname") or ""),
                                "Script": Path(name).name,
                                "Success": bool(payload.get("Success")) if isinstance(payload, dict) else True,
                                "Error": str(payload.get("Error") or "") if isinstance(payload, dict) else "",
                            }
                            merged = {**prefix, **record}
                            data_rows.append(merged)
                            for key in merged:
                                if key not in headers:
                                    headers.append(key)
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in data_rows:
            ws.append([row.get(h, "") for h in headers])
        if not data_rows:
            ws.append([job_id, "", "", "", "", "", "", "Sin resultados estructurados disponibles"])
        for i, header in enumerate(headers, 1):
            width = min(42, max(11, len(str(header)) + 2))
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "A2"
        info = wb.create_sheet("Resumen")
        info.append(["Campo", "Valor"])
        info["A1"].font = info["B1"].font = Font(bold=True)
        info.append(["JobId", job_id])
        info.append(["Targets", len(rows)])
        info.append(["Resultados listos", sum(1 for r in rows if bool(r.get("ResultReady")))])
        info.append(["Filas Data", len(data_rows)])
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    @staticmethod
    def _transfer_stamp() -> str:
        return dt.datetime.now().astimezone().isoformat(timespec="seconds")

    def _transfer_job_path(self, job_id: str) -> Path | None:
        clean = "".join(ch for ch in str(job_id or "") if ch.isalnum() or ch in "._-")
        if not clean or clean != str(job_id):
            return None
        return self.transfer_job_dir / f"{clean}.json"

    def _write_transfer_job(self, record: dict[str, Any]) -> dict[str, Any]:
        job_id = str(record.get("JobId") or "")
        path = self._transfer_job_path(job_id)
        if path is None:
            raise ValueError("JobId de transferencia inválido.")
        current = dict(record)
        current["UpdatedAt"] = self._transfer_stamp()
        temp = path.with_name(f".{path.name}.{uuid.uuid4().hex}.tmp")
        temp.write_text(json.dumps(current, ensure_ascii=False, indent=2), encoding="utf-8")
        try:
            os.chmod(temp, 0o640)
        except OSError:
            pass
        os.replace(temp, path)
        return current

    def transfer_job(self, job_id: str) -> dict[str, Any] | None:
        path = self._transfer_job_path(job_id)
        if path is None or not path.is_file():
            return None
        with self._transfer_jobs_lock:
            try:
                value = json.loads(path.read_text(encoding="utf-8-sig"))
            except Exception:
                return None
        return value if isinstance(value, dict) else None

    def transfer_jobs(self, limit: int | None = None) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []
        with self._transfer_jobs_lock:
            paths = sorted(self.transfer_job_dir.glob("transfer-*.json"), reverse=True)
            if limit is not None:
                paths = paths[:max(0, int(limit))]
            for path in paths:
                try:
                    value = json.loads(path.read_text(encoding="utf-8-sig"))
                except Exception:
                    continue
                if isinstance(value, dict):
                    rows.append(value)
        return sorted(rows, key=lambda row: str(row.get("CreatedAt") or ""), reverse=True)

    def _update_transfer_job(
        self,
        job_id: str,
        *,
        status: str | None = None,
        phase: str | None = None,
        detail: str | None = None,
        event: str | None = None,
        **changes: Any,
    ) -> dict[str, Any]:
        with self._transfer_jobs_lock:
            current = self.transfer_job(job_id)
            if current is None:
                raise RuntimeError(f"Job de transferencia no encontrado: {job_id}")
            if status is not None:
                current["Status"] = str(status).upper()
            if phase is not None:
                current["Phase"] = str(phase)
            if detail is not None:
                current["Detail"] = str(detail)[:4000]
            current.update(changes)
            if event:
                events = [dict(item) for item in (current.get("Events") or []) if isinstance(item, dict)]
                events.append({"At": self._transfer_stamp(), "Status": current.get("Status"), "Detail": str(event)[:1000]})
                current["Events"] = events[-50:]
            return self._write_transfer_job(current)

    def _transfer_worker_limits(self) -> tuple[int, int]:
        """Return conservative, configurable transfer worker and queue limits.

        The scan itself may use a private clamd socket and the Gateway can also
        be slow.  More workers do not make either dependency infinitely
        parallel, so intentionally cap this process-local pool.  The pending
        records remain persisted in ``transfer-jobs`` and survive a restart.
        """
        config = self.config if isinstance(getattr(self, "config", None), dict) else {}
        try:
            workers = int(config.get("transfer_worker_count", 2))
        except (TypeError, ValueError):
            workers = 2
        try:
            queued = int(config.get("transfer_queue_limit", 32))
        except (TypeError, ValueError):
            queued = 32
        return max(1, min(workers, 4)), max(1, min(queued, 500))

    def _ensure_transfer_scheduler(self) -> None:
        """Lazily initialize the bounded scheduler.

        Keeping this lazy makes the persistence helpers usable by recovery and
        by focused tests that construct ``WebOperations`` without its full
        service dependencies.  It also avoids relying on an unbounded executor
        work queue: a semaphore covers running *and* waiting submitted jobs.
        """
        if getattr(self, "_transfer_executor", None) is not None:
            return
        jobs_lock = getattr(self, "_transfer_jobs_lock", None)
        if jobs_lock is None:
            jobs_lock = threading.RLock()
            self._transfer_jobs_lock = jobs_lock
        with jobs_lock:
            if getattr(self, "_transfer_executor", None) is not None:
                return
            workers, queued = self._transfer_worker_limits()
            self._transfer_worker_count = workers
            self._transfer_queue_limit = queued
            self._transfer_capacity = threading.BoundedSemaphore(workers + queued)
            self._transfer_dispatch_lock = threading.RLock()
            self._transfer_executor = ThreadPoolExecutor(
                max_workers=workers,
                thread_name_prefix="gideon-transfer",
            )

    def shutdown_transfer_workers(self, *, wait: bool = False) -> None:
        """Stop the local executor during a controlled WebAdmin shutdown.

        Queued records are deliberately not discarded; the next process
        recovers them from disk before accepting new uploads.
        """
        executor = getattr(self, "_transfer_executor", None)
        if executor is not None:
            executor.shutdown(wait=bool(wait), cancel_futures=False)

    def _mark_transfer_waiting(self, job_id: str, *, detail: str, event: str) -> None:
        """Persist a queue state without turning scheduler pressure into a job error."""
        try:
            current = self.transfer_job(job_id)
            if current is None:
                return
            if str(current.get("QueueState") or "").upper() == "WAITING":
                return
            self._update_transfer_job(
                job_id,
                status="QUEUED_FOR_ANALYSIS",
                phase="WAITING_FOR_WORKER",
                detail=detail,
                event=event,
                QueueState="WAITING",
                QueuePosition=None,
            )
        except Exception:
            # A scheduler notification must never hide a persisted upload.
            pass

    def _start_transfer_job(self, job_id: str) -> bool:
        """Submit one persisted job only when bounded capacity is available.

        ``False`` means the job safely remains on disk for a future worker; it
        is not a failure and the browser can still redirect to its job status.
        """
        clean_job_id = str(job_id or "").strip()
        if not clean_job_id:
            return False
        self._ensure_transfer_scheduler()
        with self._transfer_jobs_lock:
            if clean_job_id in self._scheduled_transfer_jobs:
                return True
            capacity = self._transfer_capacity
            if not capacity.acquire(blocking=False):
                self._mark_transfer_waiting(
                    clean_job_id,
                    detail="Preparación en cola; esperando un worker de análisis disponible.",
                    event="Capacidad de análisis ocupada; el Job permanece protegido en cuarentena.",
                )
                return False
            self._scheduled_transfer_jobs.add(clean_job_id)
            try:
                current = self.transfer_job(clean_job_id)
                if current is not None:
                    self._update_transfer_job(
                        clean_job_id,
                        QueueState="SCHEDULED",
                        QueuePosition=0,
                    )
            except Exception:
                # The worker below retains the authoritative error handling.
                pass

        def _runner() -> None:
            try:
                self._process_transfer_job(clean_job_id)
            finally:
                with self._transfer_jobs_lock:
                    self._scheduled_transfer_jobs.discard(clean_job_id)
                capacity.release()
                # One completion is enough to admit the oldest persisted job.
                self._dispatch_pending_transfer_jobs()

        try:
            self._transfer_executor.submit(_runner)
        except Exception:
            with self._transfer_jobs_lock:
                self._scheduled_transfer_jobs.discard(clean_job_id)
            capacity.release()
            self._mark_transfer_waiting(
                clean_job_id,
                detail="El worker se está reiniciando; la preparación quedó guardada para recuperación.",
                event="No se pudo reservar un worker; se conserva el Job persistido.",
            )
            return False
        return True

    def _dispatch_pending_transfer_jobs(self) -> None:
        """Fill available bounded capacity from the oldest durable queue entries."""
        self._ensure_transfer_scheduler()
        with self._transfer_dispatch_lock:
            pending = [
                row for row in self.transfer_jobs(limit=500)
                if str(row.get("Status") or "").upper() == "QUEUED_FOR_ANALYSIS"
                and str(row.get("JobId") or "")
                and Path(str(row.get("StagedPath") or "")).is_file()
            ]
            pending.sort(key=lambda row: str(row.get("CreatedAt") or ""))
            position = 1
            for row in pending:
                job_id = str(row.get("JobId") or "")
                with self._transfer_jobs_lock:
                    scheduled = job_id in self._scheduled_transfer_jobs
                if scheduled:
                    continue
                try:
                    self._update_transfer_job(job_id, QueuePosition=position)
                except Exception:
                    continue
                if not self._start_transfer_job(job_id):
                    break
                position += 1

    def _recover_transfer_jobs(self) -> None:
        active = {"QUEUED_FOR_ANALYSIS", "ANALYZING", "DISPATCHING"}
        for row in self.transfer_jobs(limit=500):
            if str(row.get("Status") or "").upper() not in active:
                continue
            staged = Path(str(row.get("StagedPath") or ""))
            results = [item for item in (row.get("Results") or []) if isinstance(item, dict)]
            target_count = int(row.get("TargetCount") or 0)
            if staged.is_file():
                job_id = str(row.get("JobId") or "")
                if str(row.get("Status") or "").upper() != "QUEUED_FOR_ANALYSIS":
                    self._update_transfer_job(
                        job_id, status="QUEUED_FOR_ANALYSIS", phase="RECOVERY_QUEUED",
                        detail="Recuperando preparación interrumpida después del reinicio.",
                        event="Recuperación: el Job vuelve a la cola durable de análisis.",
                        QueueState="WAITING", QueuePosition=None,
                    )
            elif target_count and len(results) >= target_count:
                self._update_transfer_job(
                    str(row.get("JobId") or ""), status="COMPLETED", phase="COMPLETED",
                    detail=f"Transferencia encolada para {len(results)} destino(s).",
                    event="Recuperación: todos los despachos ya estaban registrados.",
                    QueuedCount=len(results), Error="",
                )
            else:
                self._update_transfer_job(
                    str(row.get("JobId") or ""), status="FAILED", phase="QUARANTINE_MISSING",
                    detail="El archivo en cuarentena no existe después del reinicio; no se envió a destinos pendientes.",
                    event="Recuperación fallida: archivo en cuarentena ausente.",
                    Error="QUARANTINE_FILE_MISSING",
                )
        self._dispatch_pending_transfer_jobs()

    def enqueue_transfer(
        self,
        *,
        filename: str,
        source_path: Path,
        destination_dir: str,
        agent_ids: list[str],
        actor: str,
        job_type: str = "FILE_TRANSFER_PREPARATION",
        known_sha256: str = "",
        known_scan_status: str = "PENDING",
        known_scan_detail: str = "",
        extra_fields: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        """Persist an upload and return before malware analysis or dispatch starts."""
        selected = list(dict.fromkeys(str(item).strip() for item in agent_ids if str(item).strip()))
        kind = str(job_type or "FILE_TRANSFER_PREPARATION").strip().upper()
        media_ingest = kind == "MEDIA_ASSET_INGEST"
        if not selected and not media_ingest:
            raise ValueError("Seleccione al menos un Agent.")
        destination = str(destination_dir or "").strip()
        if not destination and not media_ingest:
            raise ValueError("El directorio remoto es obligatorio.")
        source = Path(source_path).resolve()
        if not source.is_file() or not (source == self.transfer_dir or self.transfer_dir in source.parents):
            raise ValueError("El archivo temporal no pertenece a la cuarentena GIDEON.")
        size = source.stat().st_size
        max_bytes = int(self.config.get("max_transfer_upload_bytes", 1024 * 1024 * 1024))
        if size <= 0 or size > max_bytes:
            raise ValueError(f"El archivo debe tener entre 1 byte y {max_bytes} bytes.")
        name = Path(str(filename or "archivo")).name
        safe = "".join(c if c.isalnum() or c in "._-" else "_" for c in name) or "archivo"
        job_id = "transfer-" + dt.datetime.now().strftime("%Y%m%d-%H%M%S") + "-" + uuid.uuid4().hex[:8]
        target = self.transfer_dir / f"{job_id}-{safe}"
        os.replace(source, target)
        os.chmod(target, 0o640)
        stamp = self._transfer_stamp()
        known_digest = str(known_sha256 or "").strip().lower()
        if known_digest and (len(known_digest) != 64 or any(ch not in "0123456789abcdef" for ch in known_digest)):
            raise ValueError("El SHA-256 interno de transferencia no es válido.")
        scan_state = str(known_scan_status or "PENDING").strip().upper()
        if scan_state not in {"PENDING", "CLEAN"}:
            raise ValueError("El estado interno de análisis no es válido.")
        record = {
            "JobId": job_id,
            "JobType": kind,
            "ExecutionMethod": "Agent HTTPS/PULL",
            "Status": "QUEUED_FOR_ANALYSIS",
            "Phase": "QUARANTINED",
            "Detail": "Archivo recibido; análisis de seguridad pendiente.",
            "CreatedAt": stamp,
            "UpdatedAt": stamp,
            "CreatedBy": str(actor or ""),
            "File": name,
            "SizeBytes": size,
            "StagedPath": str(target),
            "Destination": destination,
            "AgentIds": selected,
            "TargetCount": len(selected),
            "QueuedCount": 0,
            "QueueState": "QUEUED",
            "QueuePosition": None,
            "Sha256": known_digest,
            "ScanStatus": scan_state,
            # Engine discovery may touch PATH, the Defender platform directory
            # or a Unix socket. Keep the HTTP enqueue path deterministic; the
            # worker records the effective engine as soon as analysis begins.
            "ScanEngine": "pending",
            "ScanDurationMs": None,
            "ScanStartedAt": None,
            "ScanCompletedAt": None,
            "ScanDetail": str(known_scan_detail or "")[:3000],
            "Results": [],
            "Error": "",
            "Events": [{"At": stamp, "Status": "QUEUED_FOR_ANALYSIS", "Detail": "Archivo guardado en cuarentena."}],
        }
        if extra_fields:
            # Job records are trusted internal state.  Keep extension data
            # JSON-safe and bounded before it reaches the durable queue.
            for key, value in extra_fields.items():
                if isinstance(key, str) and key and len(key) <= 80:
                    record[key] = value
        with self._transfer_jobs_lock:
            self._write_transfer_job(record)
        self._start_transfer_job(job_id)
        return record

    @staticmethod
    def _safe_media_name(value: str, digest: str) -> str:
        """Use a managed on-endpoint filename, never a marketing-provided path."""
        suffix = Path(str(value or "")).suffix.lower()
        suffix = ".mp4" if suffix != ".mp4" else suffix
        return f"{str(digest).lower()}{suffix}"

    def _media_blob_target(self, digest: str) -> Path:
        safe = str(digest or "").strip().lower()
        if len(safe) != 64 or any(ch not in "0123456789abcdef" for ch in safe):
            raise ValueError("SHA-256 de MediaAsset inválido.")
        return self.media_asset_blob_dir / f"{safe}.mp4"

    def _promote_media_blob(self, source: Path, digest: str, *, move: bool) -> tuple[Path, str]:
        """Persist a clean, content-addressed copy under the GIDEON media root."""
        candidate = Path(source).resolve()
        target = self._media_blob_target(digest)
        if target.exists():
            # An existing blob must still prove it is the exact expected
            # immutable content before a new asset refers to it.
            if sha256_file(target).lower() != str(digest).lower():
                raise RuntimeError("MEDIA_LIBRARY_INTEGRITY_CONFLICT")
            if move:
                candidate.unlink(missing_ok=True)
        elif move:
            os.replace(candidate, target)
        else:
            try:
                os.link(candidate, target)
            except OSError:
                shutil.copy2(candidate, target)
        try:
            os.chmod(target, 0o640)
        except OSError:
            pass
        return target, str(target.relative_to(self.media_asset_dir)).replace("\\", "/")

    def enqueue_media_asset(
        self, *, filename: str, source_path: Path, asset_id: str, actor: str,
    ) -> dict[str, Any]:
        """Queue a campaign upload for the normal security pipeline.

        The HTTP route only streams to quarantine and receives this durable
        record.  Antivirus, SHA-256, ffprobe and promotion run in the bounded
        background worker, so Marketing never waits for analysis.
        """
        name = Path(str(filename or "")).name
        if Path(name).suffix.lower() != ".mp4":
            raise ValueError("Solo se aceptan videos MP4 para una campaña.")
        clean_asset = str(asset_id or "").strip()
        if not clean_asset:
            raise ValueError("El identificador interno del video es obligatorio.")
        max_media = int(self.config.get("max_media_upload_bytes", self.config.get("max_transfer_upload_bytes", 1024 * 1024 * 1024)))
        source = Path(source_path).resolve()
        if not source.is_file() or source.stat().st_size > max_media:
            raise ValueError(f"El video excede el máximo de {max_media} bytes.")
        return self.enqueue_transfer(
            filename=name, source_path=source, destination_dir="", agent_ids=[], actor=actor,
            job_type="MEDIA_ASSET_INGEST", extra_fields={"MediaAssetId": clean_asset},
        )

    def enqueue_media_delivery(
        self, *, asset_path: Path, asset_id: str, sha256: str, size_bytes: int,
        agent_id: str, actor: str, delivery_id: str,
    ) -> dict[str, Any]:
        """Stage a previously approved MediaAsset for one signed Agent pull.

        A hardlink is preferred to avoid duplicating disk data; the normal
        worker re-verifies SHA-256 before it creates the signed Gateway job.
        """
        source = Path(asset_path).resolve()
        if not source.is_file() or self.media_asset_dir not in source.parents:
            raise ValueError("El video aprobado no está disponible en la biblioteca GIDEON.")
        digest = str(sha256 or "").strip().lower()
        if sha256_file(source).lower() != digest:
            raise RuntimeError("MEDIA_LIBRARY_INTEGRITY_CHANGED")
        staging = self.transfer_dir / f".media-delivery-{uuid.uuid4().hex}.part"
        try:
            os.link(source, staging)
        except OSError:
            shutil.copy2(source, staging)
        try:
            os.chmod(staging, 0o640)
        except OSError:
            pass
        return self.enqueue_transfer(
            filename=self._safe_media_name(source.name, digest), source_path=staging,
            destination_dir="Distribucion/Videos", agent_ids=[str(agent_id)], actor=actor,
            job_type="MEDIA_DISTRIBUTION", known_sha256=digest, known_scan_status="CLEAN",
            known_scan_detail="Biblioteca Media GIDEON aprobada previamente.",
            extra_fields={"MediaAssetId": str(asset_id), "MediaDeliveryId": str(delivery_id)},
        )

    def _complete_media_ingest(
        self, *, record: dict[str, Any], staged_path: Path, digest: str, size: int,
    ) -> None:
        asset_id = str(record.get("MediaAssetId") or "").strip()
        store = self.media_campaign_store
        if not asset_id or store is None:
            raise RuntimeError("MEDIA_CATALOG_UNAVAILABLE")
        self._update_transfer_job(
            str(record["JobId"]), status="ANALYZING", phase="MEDIA_FORMAT_VALIDATION",
            detail="Verificando formato MP4 compatible para pantallas GIDEON.",
            event="Antimalware aprobado; inicia la validación de formato multimedia.",
        )
        try:
            metadata = validate_media_file(
                staged_path,
                ffprobe_path=str(self.config.get("ffprobe_path") or "ffprobe"),
                timeout_seconds=self.config.get("media_validation_timeout_seconds", 20),
            ).as_dict()
        except MediaValidationError as exc:
            try:
                store.reject_media_asset(asset_id, message=str(exc), actor=str(record.get("CreatedBy") or "system"))
            except Exception:
                pass
            raise RuntimeError(f"{exc.code}: {exc}") from exc
        blob, relative = self._promote_media_blob(staged_path, digest, move=True)
        try:
            store.complete_media_asset(
                asset_id, sha256=digest, size_bytes=size, metadata=metadata, blob_relpath=relative,
                actor=str(record.get("CreatedBy") or "system"),
            )
        except Exception:
            # The blob is content-addressed and may be shared by a prior
            # approved asset.  Do not unlink it on a catalogue write error.
            raise
        self._update_transfer_job(
            str(record["JobId"]), status="COMPLETED", phase="MEDIA_READY",
            detail="Video aprobado y disponible para asignarlo a pantallas.",
            event="Formato y controles de seguridad aprobados; video listo en Biblioteca.",
            MediaAssetId=asset_id, MediaMetadata=metadata, LibraryBlob=relative,
            CompletedAt=self._transfer_stamp(), QueueState="COMPLETED", QueuePosition=0, Error="",
        )

    def _offer_transfer_media_candidate(
        self, *, record: dict[str, Any], staged_path: Path, digest: str, size: int,
    ) -> None:
        """Retain an approved compatible MP4 as an optional campaign asset.

        A normal transfer is never failed solely because it is not compatible
        with the first Media profile.  The optional campaign affordance appears
        only after the same accepted MP4 profile is proven.
        """
        if staged_path.suffix.lower() != ".mp4" or self.media_campaign_store is None:
            return
        try:
            metadata = validate_media_file(
                staged_path, ffprobe_path=str(self.config.get("ffprobe_path") or "ffprobe"),
                timeout_seconds=self.config.get("media_validation_timeout_seconds", 20),
            ).as_dict()
            _blob, relative = self._promote_media_blob(staged_path, digest, move=False)
            asset = self.media_campaign_store.import_transfer_media_asset(
                display_name=str(record.get("File") or "Video MP4"), sha256=digest, size_bytes=size,
                metadata=metadata, blob_relpath=relative, actor=str(record.get("CreatedBy") or "system"),
            )
            self._update_transfer_job(
                str(record["JobId"]), MediaCampaignCandidate={
                    "AssetId": str(asset.get("asset_id") or ""), "Status": "READY",
                    "DisplayName": str(asset.get("display_name") or "Video MP4"),
                }, event="Video compatible disponible también para una campaña.",
            )
        except MediaValidationError:
            # A generic transfer may still carry other valid business MP4s;
            # only the optional Media campaign action stays unavailable.
            self._update_transfer_job(
                str(record["JobId"]), MediaCampaignCandidate={"Status": "NOT_COMPATIBLE"},
                event="El MP4 se transferirá normalmente; no cumple el perfil de campaña GIDEON.",
            )
        except Exception:
            # Never allow an optional library promotion to block an approved
            # generic transfer.  Avoid writing internal failure details into
            # a record shown to Marketing.
            self._update_transfer_job(
                str(record["JobId"]), MediaCampaignCandidate={"Status": "UNAVAILABLE"},
                event="La opción de campaña no está disponible para este archivo; la transferencia continúa.",
            )

    def _process_transfer_job(self, job_id: str) -> None:
        staged_path: Path | None = None
        try:
            record = self.transfer_job(job_id)
            if record is None:
                return
            staged_path = Path(str(record.get("StagedPath") or "")).resolve()
            if not staged_path.is_file() or self.transfer_dir not in staged_path.parents:
                raise RuntimeError("QUARANTINE_FILE_MISSING: el archivo pendiente ya no existe.")

            self._update_transfer_job(
                job_id, status="ANALYZING", phase="HASHING",
                detail="Calculando integridad SHA-256.", event="Inicio del análisis interno.", Error="",
                QueueState="RUNNING", QueuePosition=0,
            )
            size = staged_path.stat().st_size
            digest = sha256_file(staged_path)
            record = self.transfer_job(job_id) or record
            previous_digest = str(record.get("Sha256") or "").lower()
            previous_scan = str(record.get("ScanStatus") or "").upper()
            if previous_digest and previous_digest != digest.lower():
                raise RuntimeError("QUARANTINE_INTEGRITY_CHANGED: el SHA-256 cambió durante la preparación.")
            if previous_scan == "CLEAN" and previous_digest == digest.lower():
                scan_status = "CLEAN"
                scan_detail = str(record.get("ScanDetail") or "Análisis limpio registrado antes del reinicio.")
            else:
                scan_engine = self.preferred_malware_engine()
                scan_started_at = self._transfer_stamp()
                scan_started = time.monotonic()
                self._update_transfer_job(
                    job_id, status="ANALYZING", phase="MALWARE_SCAN", Sha256=digest, SizeBytes=size,
                    ScanEngine=scan_engine, ScanStartedAt=scan_started_at,
                    detail="Analizando antimalware; el archivo continúa bloqueado en cuarentena.",
                    event=f"SHA-256 calculado; inicia {scan_engine} + YARA.",
                )
                scan_status, scan_detail = self.scan_script(staged_path)
                scan_duration_ms = int((time.monotonic() - scan_started) * 1000)
                scan_engine = self._malware_engine_from_detail(scan_detail, scan_engine)
                self._update_transfer_job(
                    job_id, status="ANALYZING", phase="MALWARE_SCAN_COMPLETE",
                    Sha256=digest, ScanStatus=scan_status, ScanDetail=scan_detail,
                    ScanEngine=scan_engine, ScanDurationMs=scan_duration_ms, ScanCompletedAt=self._transfer_stamp(),
                    detail=f"Análisis finalizado: {scan_status} en {scan_duration_ms} ms.",
                    event=f"Resultado antimalware: {scan_status} ({scan_engine}, {scan_duration_ms} ms).",
                )
            if bool(self.config.get("require_malware_scan_for_transfers", True)) and scan_status != "CLEAN":
                raise RuntimeError(f"TRANSFER_SECURITY_BLOCKED: {scan_status} · {scan_detail}")

            record = self.transfer_job(job_id) or record
            if str(record.get("JobType") or "").upper() == "MEDIA_ASSET_INGEST":
                self._complete_media_ingest(record=record, staged_path=staged_path, digest=digest, size=size)
                return
            if str(record.get("JobType") or "").upper() == "FILE_TRANSFER_PREPARATION":
                self._offer_transfer_media_candidate(record=record, staged_path=staged_path, digest=digest, size=size)
                record = self.transfer_job(job_id) or record
            results = [dict(item) for item in (record.get("Results") or []) if isinstance(item, dict)]
            completed_agents = {str(item.get("AgentId") or "") for item in results}
            pending = [str(item) for item in (record.get("AgentIds") or []) if str(item) not in completed_agents]
            self._update_transfer_job(
                job_id, status="DISPATCHING", phase="SIGNED_PACKAGE_DISPATCH",
                detail=f"Análisis limpio. Encolando paquetes firmados ({len(results)}/{int(record.get('TargetCount') or 0)}).",
                event="Controles de seguridad aprobados; inicia el despacho.",
            )

            def _on_queued(result: dict[str, Any]) -> None:
                current = self.transfer_job(job_id) or {}
                current_results = [dict(item) for item in (current.get("Results") or []) if isinstance(item, dict)]
                agent_id = str(result.get("AgentId") or "")
                if agent_id not in {str(item.get("AgentId") or "") for item in current_results}:
                    current_results.append(dict(result))
                self._update_transfer_job(
                    job_id, status="DISPATCHING", phase="SIGNED_PACKAGE_DISPATCH",
                    detail=f"Paquetes firmados encolados: {len(current_results)}/{int(current.get('TargetCount') or 0)}.",
                    event=f"Destino {agent_id or '-'} encolado en Gateway.",
                    Results=current_results, QueuedCount=len(current_results),
                )

            if pending:
                staged = {
                    "Path": staged_path, "Name": str(record.get("File") or staged_path.name),
                    "SizeBytes": size, "Sha256": digest, "ScanStatus": scan_status, "ScanDetail": scan_detail,
                }
                self.queue_transfer(
                    staged=staged, destination_dir=str(record.get("Destination") or ""),
                    agent_ids=pending, actor=str(record.get("CreatedBy") or ""), on_queued=_on_queued,
                )
            final = self.transfer_job(job_id) or record
            final_results = [item for item in (final.get("Results") or []) if isinstance(item, dict)]
            target_count = int(final.get("TargetCount") or 0)
            if len(final_results) < target_count:
                raise RuntimeError(f"TRANSFER_DISPATCH_INCOMPLETE: {len(final_results)}/{target_count} destinos encolados.")
            self._update_transfer_job(
                job_id, status="COMPLETED", phase="COMPLETED",
                detail=f"Archivo aprobado y transferencia encolada para {len(final_results)} destino(s).",
                event="Preparación completada; continúe el progreso en los Jobs de destino.",
                QueuedCount=len(final_results), Error="", CompletedAt=self._transfer_stamp(),
                QueueState="COMPLETED", QueuePosition=0,
            )
        except Exception as exc:
            message = redact_error(f"{type(exc).__name__}: {exc}")[:4000]
            try:
                self._update_transfer_job(
                    job_id, status="FAILED", phase="FAILED", detail=message,
                    event="La transferencia fue bloqueada o no pudo encolarse.", Error=message,
                    FailedAt=self._transfer_stamp(), QueueState="FAILED", QueuePosition=0,
                )
            except Exception:
                pass
        finally:
            final = self.transfer_job(job_id)
            if final is not None and str(final.get("Status") or "").upper() in {"COMPLETED", "FAILED"}:
                candidate = staged_path or Path(str(final.get("StagedPath") or ""))
                try:
                    resolved = candidate.resolve()
                    if self.transfer_dir in resolved.parents:
                        resolved.unlink(missing_ok=True)
                except OSError:
                    pass

    def stage_transfer(self, *, filename: str, source_path: Path) -> dict[str, Any]:
        """Move an already-streamed upload into quarantine and statically scan it.

        Nothing from the upload is executed on the Ubuntu host.  The staged file
        remains mode 0640 and is only read for hashing/scanning/signing.
        """
        name = Path(str(filename or "archivo")).name
        max_bytes = int(self.config.get("max_transfer_upload_bytes", 1024 * 1024 * 1024))
        source = Path(source_path).resolve()
        size = source.stat().st_size
        if size <= 0 or size > max_bytes:
            raise ValueError(f"El archivo debe tener entre 1 byte y {max_bytes} bytes.")
        digest = sha256_file(source)
        safe = "".join(c if c.isalnum() or c in "._-" else "_" for c in name)
        target = self.transfer_dir / f"{dt.datetime.now().strftime('%Y%m%d%H%M%S')}-{digest[:12]}-{safe}"
        os.replace(source, target)
        os.chmod(target, 0o640)
        status, detail = self.scan_script(target)
        if bool(self.config.get("require_malware_scan_for_transfers", True)) and status != "CLEAN":
            raise RuntimeError(f"Transferencia bloqueada por análisis de seguridad: {status} · {detail}")
        return {"Path": target, "Name": name, "SizeBytes": size, "Sha256": digest, "ScanStatus": status, "ScanDetail": detail}

    @staticmethod
    def _portable_destination_parts(destination_dir: str) -> list[str]:
        """Return a safe GIDEON-root-relative destination for every endpoint OS."""
        raw = str(destination_dir or "").strip().replace("\\", "/")
        if not raw:
            raise ValueError("La carpeta GIDEON de destino es obligatoria.")
        lowered = raw.casefold()
        known_roots = (
            "gideon://", "gideon_root/", "c:/gideon/", "/opt/gideon/",
            "/var/lib/gideon/files/", "/library/application support/gideon/",
        )
        for root in known_roots:
            if lowered.startswith(root):
                raw = raw[len(root):]
                break
        raw = raw.strip("/")
        parts = [part.strip() for part in raw.split("/") if part.strip()]
        if not parts:
            raise ValueError("Indique una carpeta dentro de la raíz GIDEON, por ejemplo Distribucion/Videos.")
        if any(part in {".", ".."} or ":" in part or "\x00" in part for part in parts):
            raise ValueError("La carpeta GIDEON contiene un segmento no permitido.")
        return parts

    @classmethod
    def _remote_join(cls, platform: str, destination_dir: str, filename: str) -> str:
        parts = cls._portable_destination_parts(destination_dir)
        platform_name = str(platform or "").strip().lower()
        if platform_name.startswith("win"):
            return "\\".join(["C:\\GIDEON", *parts, Path(filename).name])
        if platform_name in {"darwin", "mac", "macos"}:
            return "/".join(["/Library/Application Support/GIDEON", *parts, Path(filename).name])
        # Linux, Ubuntu, Debian, Raspberry Pi/Raspbian and future Unix variants.
        return "/".join(["/opt/gideon", *parts, Path(filename).name])

    def queue_transfer(
        self,
        *,
        staged: dict[str, Any],
        destination_dir: str,
        agent_ids: list[str],
        actor: str,
        on_queued: Any | None = None,
    ) -> list[dict[str, Any]]:
        if not agent_ids:
            raise ValueError("Seleccione al menos un Agent.")
        private_key, private_key_bytes, signing_key_id = self._signing_material()
        cfg = {
            "ignore_local_gateway_runtime": True,
            "execution_signing_private_key": str(private_key) if private_key is not None else "",
            "execution_signing_private_key_bytes": private_key_bytes,
            "execution_signing_public_key": str(self.security_assets.managed_public_path if self.security_assets is not None and self.security_assets.managed_public_path.is_file() else (self.config.get("execution_signing_public_key") or os.environ.get("GIDEON_AGENT_SIGNING_PUBLIC_KEY", ""))),
            "agent_control_plane": {
                "base_url": self.config.get("gateway_base_url"),
                "public_url": self.config.get("gateway_base_url"),
                "admin_api_key_env": self.config.get("gateway_admin_token_env", "GIDEON_AGENT_ADMIN_TOKEN"),
                "verify_tls": self.config.get("gateway_verify_tls", True),
                "ca_file": self.config.get("gateway_ca_file"),
                "require_https": self.config.get("gateway_require_https", True),
                "signing_key_id": signing_key_id,
                "timeout_seconds": 30,
            },
            "file_transfer": {"max_wait_seconds": int(self.config.get("transfer_max_wait_seconds", 3600))},
        }
        transport = AgentHttpsPullFileTransport(app_root=self.app_root, cfg=cfg)
        all_agents = {str(x.get("agent_id") or ""): x for x in self.gateway.agents()}
        path = Path(staged["Path"])
        results=[]
        for aid in agent_ids:
            row=all_agents.get(str(aid))
            if row is None:
                raise RuntimeError(f"Agent no encontrado: {aid}")
            self._preflight_agent_operation(row, file_transfer=True)
            remote=self._remote_join(str(row.get("platform") or "windows"), destination_dir, str(staged["Name"]))
            kind=ArtifactKind.VIDEO if path.suffix.lower() in {".mp4",".mkv",".mov",".avi",".webm"} else ArtifactKind.FILE
            artifact=ArtifactDescriptor(local_path=path, remote_path=remote, kind=kind, size_bytes=int(staged["SizeBytes"]), sha256=str(staged["Sha256"]))
            target=HostTarget(
                code=str(row.get("site_code") or ""), name=str(row.get("site_name") or ""), address="",
                platform=str(row.get("platform") or "windows"), host_profile=str(row.get("host_profile") or "windows_generic"),
                agent_id=str(row.get("agent_id") or ""), agent_name=str(row.get("agent_name") or ""), hostname=str(row.get("hostname") or ""),
            )
            queued = transport.queue_upload(target,artifact)
            results.append(queued)
            if callable(on_queued):
                on_queued(dict(queued))
        receipt={"CreatedAt":dt.datetime.now().astimezone().isoformat(timespec="seconds"),"CreatedBy":actor,"File":str(staged["Name"]),"Sha256":str(staged["Sha256"]),"Destination":destination_dir,"Results":results}
        receipt_path=self.transfer_dir/f"receipt-{dt.datetime.now().strftime('%Y%m%d%H%M%S')}-{uuid.uuid4().hex[:8]}.json"
        receipt_path.write_text(json.dumps(receipt,ensure_ascii=False,indent=2),encoding="utf-8")
        return results

    def control_plane_health(self, *, stale_after_seconds: int = 45) -> dict[str, Any]:
        """Return a read-only aggregate health model for the Web dashboard.

        Diagnostics/Coordinator remain the source of truth. WebAdmin only adds
        freshness semantics so a stale RUNNING snapshot is never presented as
        current health.
        """
        root = self.data_dir.parent
        coordinator_path = root / "coordinator" / "status.json"
        diagnostics_path = root / "diagnostics" / "health.json"
        result: dict[str, Any] = {
            "state": "STARTING", "fresh": False, "age_seconds": None,
            "generated_at_utc": "", "source": "GIDEONRemoteControlCoordinator",
            "services": [], "detail": "Esperando snapshot del Coordinator GIDEON.",
        }
        payload: dict[str, Any] | None = None
        source_path: Path | None = None
        read_errors: list[str] = []
        # Coordinator owns its private state and WebAdmin intentionally runs as
        # an unprivileged account. Prefer that snapshot when it is readable,
        # then fall back to the sanitized Diagnostics projection without ever
        # turning an expected PermissionError into an HTTP 500.
        for candidate in (coordinator_path, diagnostics_path):
            try:
                value = json.loads(candidate.read_text(encoding="utf-8-sig"))
                if not isinstance(value, dict):
                    read_errors.append(f"{candidate.name}: formato inválido")
                    continue
                payload = value
                source_path = candidate
                break
            except FileNotFoundError:
                continue
            except PermissionError:
                read_errors.append(f"{candidate.name}: acceso restringido")
            except Exception as exc:
                read_errors.append(f"{candidate.name}: {type(exc).__name__}")
        if payload is None or source_path is None:
            result["state"] = "DEGRADED"
            suffix = f" ({'; '.join(read_errors)})" if read_errors else ""
            result["detail"] = f"Snapshot GIDEON no disponible{suffix}"
            return result
        generated = str(payload.get("generated_at_utc") or "")
        age: float | None = None
        if generated:
            try:
                parsed = dt.datetime.fromisoformat(generated.replace("Z", "+00:00"))
                if parsed.tzinfo is None: parsed = parsed.replace(tzinfo=dt.timezone.utc)
                age = max(0.0, (dt.datetime.now(dt.timezone.utc) - parsed.astimezone(dt.timezone.utc)).total_seconds())
            except ValueError:
                age = None
        fresh = age is not None and age <= max(10, int(stale_after_seconds))
        raw_state = str(payload.get("state") or "FAILED").upper()
        state = raw_state if raw_state in {"RUNNING", "STARTING", "DEGRADED", "FAILED"} else "FAILED"
        if not fresh and state == "RUNNING": state = "DEGRADED"
        services = payload.get("services") if isinstance(payload.get("services"), list) else []
        result.update({
            "state": state, "fresh": fresh, "age_seconds": age,
            "generated_at_utc": generated, "source": str(payload.get("source") or source_path.name),
            "services": [dict(x) for x in services if isinstance(x, dict)],
            "detail": ("Snapshot vigente" if fresh else "Snapshot obsoleto o sin timestamp válido"),
        })
        return result

    def diagnostics(self) -> list[dict[str, Any]]:
        checks: list[dict[str, Any]]=[]
        def add(name: str, ok: bool, detail: str, level: str | None = None) -> None:
            checks.append({"Name":name,"Ok":bool(ok),"Detail":str(detail),"Level":level or ("OK" if ok else "ERROR")})
        service_health = self.data_dir.parent / "diagnostics" / "health.json"
        try:
            payload = json.loads(service_health.read_text(encoding="utf-8-sig"))
            rows = payload.get("services") if isinstance(payload, dict) else []
            if not isinstance(rows, list):
                rows = []
            for row in rows:
                if not isinstance(row, dict):
                    continue
                state = str(row.get("state") or "FAILED").upper()
                component = str(row.get("component") or "unknown")
                level = "OK" if state == "RUNNING" else "WARNING" if state in {"STARTING", "DEGRADED"} else "ERROR"
                detail = f"Estado={state} · {row.get('detail') or '-'} · fuente={row.get('source') or '-'}"
                add(f"Servicio GIDEON · {component}", state == "RUNNING", detail, level)
        except Exception as exc:
            add("Servicios GIDEON", False, f"Snapshot de Diagnostics no disponible: {type(exc).__name__}: {exc}", "WARNING")
        try:
            health=self.gateway.health()
            add("Gateway HTTPS", bool(health.get("Ok", True)), f"Versión {health.get('Version') or '?'} · SchemaReady={health.get('SchemaReady')}")
            caps=health.get("Capabilities") if isinstance(health.get("Capabilities"),dict) else {}
            add("API Web Admin", bool(caps.get("WebAdminApi")), str(caps.get("JobsListEndpoint") or "Capability ausente"))
        except Exception as exc:
            add("Gateway HTTPS",False,redact_error(f"{type(exc).__name__}: {exc}"))
        ok,key_detail=self.execution_available(); add("Firma Ed25519 de Jobs",ok,key_detail)
        clam=shutil.which("clamscan"); clamd=shutil.which("clamdscan"); db=str(self.config.get("clamav_database_dir") or "").strip(); socket=str(self.config.get("clamd_socket") or "").strip(); engine=self.preferred_malware_engine(); daemon_ready=engine == "clamd-stream"; add("ClamAV",bool(clam or daemon_ready),f"Motor={engine} · clamdscan={clamd or 'no'} · socket={socket or 'no'} · fallback={clam or 'no'} · DB GIDEON={db or 'no'}","OK" if daemon_ready else "WARNING" if clam else "ERROR")
        yara=shutil.which("yara"); rules=Path(str(self.config.get("yara_rules_file") or "")); add("YARA",bool(yara and rules.is_file()),f"Binario={yara or 'no'} · Reglas={rules}","OK" if yara and rules.is_file() else "WARNING")
        add("Directorio de datos",os.access(self.data_dir,os.W_OK),str(self.data_dir))
        add("Código de aplicación solo lectura",not os.access(self.app_root,os.W_OK),str(self.app_root),"OK" if not os.access(self.app_root,os.W_OK) else "WARNING")
        add("HTTPS obligatorio Gateway",bool(self.config.get("gateway_require_https",True)),str(self.config.get("gateway_base_url") or ""))
        add("Cookies seguras",bool(self.config.get("secure_cookies",True)),"Secure + HttpOnly + SameSite=Strict" if self.config.get("secure_cookies",True) else "Secure cookie deshabilitada")
        add("MFA requerido",bool(self.config.get("require_mfa",True)),"TOTP obligatorio" if self.config.get("require_mfa",True) else "MFA opcional","OK" if self.config.get("require_mfa",True) else "WARNING")
        add("Secretos MFA cifrados",bool(os.environ.get("GIDEON_WEB_TOTP_ENCRYPTION_KEY")),"Fernet/AES128 + HMAC-SHA256" if os.environ.get("GIDEON_WEB_TOTP_ENCRYPTION_KEY") else "Clave de cifrado TOTP no configurada", "OK" if os.environ.get("GIDEON_WEB_TOTP_ENCRYPTION_KEY") else "WARNING")
        return checks

    def safe_settings(self) -> list[tuple[str, str]]:
        rows=[]
        for key,value in sorted(self.config.items()):
            if is_sensitive_key(key):
                shown="Configurado" if value else "No configurado"
            elif isinstance(value,(dict,list)):
                shown=json.dumps(value,ensure_ascii=False)
            else:
                shown=str(value)
            rows.append((key,shown))
        return rows
