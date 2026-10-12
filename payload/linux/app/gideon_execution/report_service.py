from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any, Optional

from gideon_admin import ensure_dir
from gideon_common.hashing import sha256_file
from gideon_execution.runtime_service import safe_component

def parse_script_result(path: Path) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    payload = json.loads(path.read_text(encoding="utf-8-sig"))
    data = payload.get("Data", [])
    if isinstance(data, dict):
        data = [data]
    if data is None:
        data = []
    rows: list[dict[str, Any]] = []
    for item in data:
        if isinstance(item, dict):
            rows.append(item)
        else:
            rows.append({"Value": item})
    return payload, rows

def format_job(
    manifest: dict[str, Any],
    export_format: Optional[str] = None,
) -> list[Path]:
    export_format = export_format or manifest.get("ExportFormat", "xlsx")
    if export_format == "none":
        return []

    output_base = Path(manifest["OutputBase"])
    ensure_dir(output_base.parent)
    summary: list[dict[str, Any]] = []
    script_summary: list[dict[str, Any]] = []
    data_rows: list[dict[str, Any]] = []
    manifest_scripts = list(manifest.get("Scripts") or [])

    for code, entry in manifest["Sites"].items():
        row = {
            "SiteCode": entry.get("SiteCode"),
            "SiteName": entry.get("SiteName"),
            "AgentId": entry.get("AgentId"),
            "AgentName": entry.get("AgentName"),
            "Hostname": entry.get("Hostname"),
            "ServerIP": entry.get("ServerIP"),
            "TransferStatus": entry.get("TransferStatus"),
            "StartStatus": entry.get("StartStatus"),
            "DispatchStatus": entry.get("DispatchStatus"),
            "RemoteState": entry.get("RemoteState"),
            "CollectionStatus": entry.get("CollectionStatus"),
            "TaskState": entry.get("TaskState"),
            "LastTaskResult": entry.get("LastTaskResult"),
            "ScriptCount": entry.get("ScriptCount") or manifest.get("ScriptCount"),
            "ScriptFailures": entry.get("ScriptFailures", 0),
            "ScriptSkipped": entry.get("ScriptSkipped", 0),
            "CollectionVerified": entry.get("CollectionVerified"),
            "CollectionWarnings": " | ".join(entry.get("CollectionWarnings") or []),
            "CurrentScriptIndex": entry.get("CurrentScriptIndex"),
            "CurrentScriptName": entry.get("CurrentScriptName"),
            "PackageVerified": entry.get("PackageVerified"),
            "PackageBytes": entry.get("PackageBytes"),
            "RawDir": entry.get("RawDir"),
            "RawFile": entry.get("RawFile"),
            "Error": entry.get("Error") or entry.get("DispatchError"),
        }
        progress = entry.get("Progress")
        if isinstance(progress, dict):
            row.update({
                "ProgressType": progress.get("Type"),
                "Drive": progress.get("Drive"),
                "Directories": progress.get("Directories"),
                "BackupFolders": progress.get("BackupFolders"),
                "AccessErrors": progress.get("AccessErrors"),
                "ElapsedSec": progress.get("ElapsedSec"),
                "CurrentPath": progress.get("CurrentPath"),
            })
        summary.append(row)

        raw_dir = entry.get("RawDir")
        if not raw_dir:
            continue
        local_scripts_dir = Path(str(raw_dir)) / "scripts"

        for script in manifest_scripts:
            result_name = script.get("ResultName")
            if not result_name:
                # Manifest global created before packages; reconstruct deterministic name.
                idx = int(script.get("Index") or 0)
                stem = safe_component(str(script.get("Name") or f"script_{idx}"))
                result_name = f"{idx:02d}_{stem}.json"
            path = local_scripts_dir / str(result_name)
            script_row = {
                "SiteCode": entry.get("SiteCode"),
                "SiteName": entry.get("SiteName"),
                "AgentId": entry.get("AgentId"),
                "AgentName": entry.get("AgentName"),
                "Hostname": entry.get("Hostname"),
                "ServerIP": entry.get("ServerIP"),
                "ScriptIndex": script.get("Index"),
                "ScriptName": script.get("Name"),
                "ResultFile": str(path),
                "Success": None,
                "Skipped": None,
                "ResultSha256": None,
                "ElapsedSec": None,
                "Directories": None,
                "BackupFolders": None,
                "AccessErrors": None,
                "Error": None,
            }
            if path.exists():
                try:
                    payload, rows = parse_script_result(path)
                    script_row.update({
                        "Success": payload.get("Success"),
                        "Skipped": payload.get("Skipped", False),
                        "ResultSha256": sha256_file(path),
                        "ElapsedSec": payload.get("ElapsedSec"),
                        "Directories": (
                            (payload.get("Progress") or {}).get("Directories")
                            if isinstance(payload.get("Progress"), dict)
                            else None
                        ),
                        "BackupFolders": (
                            (payload.get("Progress") or {}).get("BackupFolders")
                            if isinstance(payload.get("Progress"), dict)
                            else None
                        ),
                        "AccessErrors": (
                            (payload.get("Progress") or {}).get("AccessErrors")
                            if isinstance(payload.get("Progress"), dict)
                            else None
                        ),
                        "StartedAt": payload.get("StartedAt"),
                        "FinishedAt": payload.get("FinishedAt"),
                        "Error": payload.get("Error"),
                    })
                    for item in rows:
                        data_rows.append({
                            "SiteCode": entry.get("SiteCode"),
                            "SiteName": entry.get("SiteName"),
                            "AgentId": entry.get("AgentId"),
                            "AgentName": entry.get("AgentName"),
                            "Hostname": entry.get("Hostname"),
                            "ServerIP": entry.get("ServerIP"),
                            "ScriptIndex": script.get("Index"),
                            "ScriptName": script.get("Name"),
                            **item,
                        })
                except Exception as exc:
                    script_row["Error"] = f"LOCAL_PARSE_ERROR: {exc}"
            else:
                script_row["Error"] = "RESULT_FILE_NOT_FOUND"
            script_summary.append(script_row)

    created: list[Path] = []

    def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
        if not rows:
            path.write_text("", encoding="utf-8-sig")
            return
        fields: list[str] = []
        seen: set[str] = set()
        for r in rows:
            for key in r.keys():
                if key not in seen:
                    seen.add(key)
                    fields.append(key)
        with path.open("w", encoding="utf-8-sig", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=fields)
            writer.writeheader()
            writer.writerows(rows)

    if export_format == "json":
        path = output_base.with_suffix(".json")
        path.write_text(
            json.dumps(
                {
                    "Job": manifest,
                    "Summary": summary,
                    "ScriptSummary": script_summary,
                    "Data": data_rows,
                },
                indent=2,
                ensure_ascii=False,
                default=str,
            ),
            encoding="utf-8",
        )
        created.append(path)

    elif export_format == "csv":
        summary_path = output_base.with_name(output_base.name + "-summary").with_suffix(".csv")
        script_path = output_base.with_name(output_base.name + "-scripts").with_suffix(".csv")
        data_path = output_base.with_name(output_base.name + "-data").with_suffix(".csv")
        write_csv(summary_path, summary)
        write_csv(script_path, script_summary)
        write_csv(data_path, data_rows)
        created.extend([summary_path, script_path, data_path])

    elif export_format == "txt":
        path = output_base.with_suffix(".txt")
        chunks = [
            f"GIDEON JOB {manifest.get('JobId')}",
            "TRANSPORT: SFTP -> SSH -> SFTP",
            f"SCRIPTS: {manifest.get('ScriptCount', 0)}",
            "",
        ]
        for entry in manifest["Sites"].values():
            chunks.append(
                f"===== {entry.get('SiteCode')}_{entry.get('SiteName')} ====="
            )
            chunks.append(
                f"State={entry.get('RemoteState')} Transfer={entry.get('TransferStatus')} "
                f"Start={entry.get('StartStatus')} Collect={entry.get('CollectionStatus')}"
            )
            for result_file in entry.get("ScriptResultFiles") or []:
                p = Path(str(result_file))
                chunks.append(f"--- {p.name} ---")
                if p.exists():
                    chunks.append(p.read_text(encoding="utf-8-sig", errors="replace"))
                else:
                    chunks.append("NO RESULT FILE")
        path.write_text("\n\n".join(chunks), encoding="utf-8")
        created.append(path)

    elif export_format == "xlsx":
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
        except ImportError as exc:
            raise RuntimeError(
                "Falta openpyxl. Ejecute: py -m pip install -r requirements.txt"
            ) from exc

        path = output_base.with_suffix(".xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Summary"

        def write_sheet(sheet, rows):
            if not rows:
                sheet.append(["Sin datos"])
                return
            fields: list[str] = []
            seen: set[str] = set()
            for r in rows:
                for key in r.keys():
                    if key not in seen:
                        seen.add(key)
                        fields.append(key)
            sheet.append(fields)
            for r in rows:
                sheet.append([
                    r.get(field)
                    if not isinstance(r.get(field), (dict, list))
                    else json.dumps(r.get(field), ensure_ascii=False, default=str)
                    for field in fields
                ])
            for cell in sheet[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="D9EAF7")
                cell.alignment = Alignment(horizontal="center")
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for idx, field in enumerate(fields, 1):
                values = [str(field)]
                for rr in sheet.iter_rows(min_row=2, min_col=idx, max_col=idx):
                    values.append(str(rr[0].value or ""))
                width = min(max(max(len(v) for v in values) + 2, 10), 60)
                sheet.column_dimensions[get_column_letter(idx)].width = width

        write_sheet(ws, summary)
        scripts_ws = wb.create_sheet("Scripts")
        write_sheet(scripts_ws, script_summary)
        data_ws = wb.create_sheet("Data")
        write_sheet(data_ws, data_rows)
        wb.save(path)
        created.append(path)

    else:
        raise ValueError(f"Formato no soportado: {export_format}")

    return created
