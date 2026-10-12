from __future__ import annotations

import json
from datetime import UTC, datetime
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    KeepTogether,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[2]
OUTPUT_DIR = ROOT / "documentation"
DOCX_PATH = OUTPUT_DIR / "GIDEON_PRE-RC31_Informe_Tecnico_Instalacion_y_Capacidad.docx"
PDF_PATH = OUTPUT_DIR / "GIDEON_PRE-RC31_Informe_Tecnico_Instalacion_y_Capacidad.pdf"
VALIDATION_PATH = ROOT / "metadata" / "oneclick-platform-validation.json"
QUALITY_PATH = ROOT / "metadata" / "code-quality-repair-report.json"

NAVY = "17365D"
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EEF5FB"
GREEN = "2E7D32"
LIGHT_GREEN = "E2F0D9"
AMBER = "9C6500"
LIGHT_AMBER = "FFF2CC"
RED = "A61B1B"
LIGHT_GRAY = "F2F2F2"
MID_GRAY = "666666"
WHITE = "FFFFFF"

REPORT_DATE = "31 de agosto de 2026"
REPORT_VERSION = "GIDEON Master PRE-RC31"

SOURCES = [
    (
        "[1] Microsoft — Requisitos del sistema para Hyper-V",
        "https://learn.microsoft.com/es-es/windows-server/virtualization/hyper-v/host-hardware-requirements",
    ),
    (
        "[2] Docker — Install Docker Desktop on Windows",
        "https://docs.docker.com/desktop/setup/install/windows-install/",
    ),
    (
        "[3] Microsoft — Habilitar Escritorio remoto",
        "https://learn.microsoft.com/en-us/windows-server/remote/remote-desktop-services/remotepc/remote-desktop-allow-access",
    ),
    (
        "[4] Microsoft — Licencias RDS CAL",
        "https://learn.microsoft.com/es-es/windows-server/remote/remote-desktop-services/rds-client-access-license",
    ),
    (
        "[5] Microsoft — Requisitos de ancho de banda RDP",
        "https://learn.microsoft.com/en-us/azure/virtual-desktop/rdp-bandwidth",
    ),
    (
        "[6] Microsoft — Requisitos de hardware de Windows Server",
        "https://learn.microsoft.com/en-us/windows-server/get-started/hardware-requirements",
    ),
    (
        "[7] Apache Guacamole — Configuración y protocolos",
        "https://guacamole.apache.org/doc/gug/configuring-guacamole.html",
    ),
]


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text_color(cell, rgb: str, bold: bool = False) -> None:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor.from_string(rgb)
            run.bold = bold


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def add_hyperlink(paragraph, text: str, url: str, color: str = "0563C1") -> None:
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Página ")
    run.font.size = Pt(8)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def configure_docx_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, color in (
        ("Title", 26, NAVY),
        ("Heading 1", 16, NAVY),
        ("Heading 2", 12, BLUE),
        ("Heading 3", 10.5, BLUE),
    ):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(10 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)

    for style_name in ("List Bullet", "List Number"):
        doc.styles[style_name].font.name = "Aptos"
        doc.styles[style_name].font.size = Pt(9.5)


def add_docx_table(doc: Document, headers: list[str], rows: list[list[str]], widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    for index, heading in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = heading
        set_cell_shading(cell, NAVY)
        set_cell_text_color(cell, WHITE, bold=True)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            cell.width = widths[index]
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value)
            set_cell_margins(cells[index])
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                set_cell_shading(cells[index], PALE_BLUE)
            if widths:
                cells[index].width = widths[index]
            for p in cells[index].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_docx_callout(doc: Document, title: str, body: str, fill: str, title_color: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=150, start=160, bottom=150, end=160)
    p = cell.paragraphs[0]
    r = p.add_run(title + "\n")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(title_color)
    r.font.size = Pt(11)
    r = p.add_run(body)
    r.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_docx_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def build_docx(validation: dict, quality: dict) -> None:
    doc = Document()
    configure_docx_styles(doc)
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    header = section.header.paragraphs[0]
    header.text = f"{REPORT_VERSION}  |  Informe técnico"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MID_GRAY)
    footer = section.footer.paragraphs[0]
    add_page_number(footer)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(38)
    r = p.add_run("GIDEON REMOTE CONTROL")
    r.font.name = "Aptos Display"
    r.font.size = Pt(15)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(BLUE)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Informe técnico de instalación,\nrequisitos y capacidad")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"{REPORT_VERSION}\n{REPORT_DATE}")
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(MID_GRAY)
    doc.add_paragraph()
    add_docx_callout(
        doc,
        "DICTAMEN EJECUTIVO",
        "GO para una prueba de instalación controlada. La aprobación de producción queda condicionada a ejecutar el OneClick completo en un host Microsoft Hyper-V real y a realizar una prueba de carga con la cantidad prevista de agentes y sesiones remotas.",
        LIGHT_AMBER,
        AMBER,
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Documento de ingeniería · No sustituye una prueba de capacidad en el entorno objetivo")
    r.italic = True
    r.font.color.rgb = RGBColor.from_string(MID_GRAY)
    doc.add_page_break()

    doc.add_heading("1. Alcance y criterio de evaluación", level=1)
    doc.add_paragraph(
        "Este informe consolida el estado técnico del bundle GIDEON Master PRE-RC31, el comportamiento adaptativo del instalador OneClick, los requisitos de hardware y sistema operativo, los límites configurados por servicio y una estimación conservadora de capacidad. Se diferencia entre:"
    )
    add_docx_bullets(
        doc,
        [
            "Límite configurado: restricción impuesta actualmente por el código o la configuración.",
            "Capacidad práctica inicial: carga prudente para comenzar una prueba de aceptación.",
            "Capacidad ampliada: posible únicamente después de redimensionar, ajustar parámetros y validar con carga real.",
        ],
    )

    doc.add_heading("2. Estado de preparación", level=1)
    result = validation.get("result", "NO DISPONIBLE")
    quality_result = quality.get("result", "NO DISPONIBLE")
    add_docx_table(
        doc,
        ["Área", "Resultado", "Interpretación"],
        [
            ["OneClick / plataforma", result, "Windows Home + Docker validado; ejecución Hyper-V real pendiente."],
            ["Calidad integral", quality_result, "Validaciones de sintaxis, estructura, hashes e importación aprobadas."],
            ["Imagen Hyper-V", "PREVALIDACIÓN PASS", "UEFI, cloud-init, Docker offline y TCP 4822 validados con QEMU; aún no con Hyper-V."],
            ["Producción", "CONDICIONAL", "Requiere prueba Hyper-V real, reinicio, enrolamiento, acceso remoto y prueba de carga."],
        ],
        [Cm(3.5), Cm(5.2), Cm(8.5)],
    )

    tests = quality.get("tests", {})
    ps = tests.get("powershell_syntax", {})
    py = tests.get("python_ast", {})
    js = tests.get("json_parse", {})
    sh = tests.get("shell_syntax", {})
    add_docx_table(
        doc,
        ["Control", "Cobertura", "Estado"],
        [
            ["PowerShell", f"{ps.get('files', 551)} archivos", ps.get("status", "PASS")],
            ["Python AST", f"{py.get('files', 3243)} archivos", py.get("status", "PASS")],
            ["JSON", f"{js.get('files', 1390)} archivos", js.get("status", "PASS")],
            ["Shell", f"{sh.get('files', 45)} archivos", sh.get("status", "PASS")],
            ["Manifiesto / hashes", "16,644 archivos elegibles", "PASS · 0 errores"],
            ["YARA", "Código del bundle", tests.get("yara_code_scan", "PASS")],
            ["Importación transaccional", f"{tests.get('validated_import_transaction_rounds', 2)} rondas", tests.get("validated_import_transaction", "PASS")],
        ],
        [Cm(5.2), Cm(6.4), Cm(5.6)],
    )

    doc.add_heading("3. Comportamiento del OneClick por sistema operativo", level=1)
    add_docx_table(
        doc,
        ["Plataforma", "Backend seleccionado", "Automatización y observaciones"],
        [
            ["Windows Home / Core", "Docker Desktop + WSL2", "Preflight, imagen offline, contenedor guacd loopback e idempotencia. La aceptación de licencia de Docker requiere intervención del operador."],
            ["Windows Pro / Enterprise / Education", "Hyper-V Generación 2", "Habilita Hyper-V si falta, registra continuación, reinicia, valida la candidata y solo la promueve después de PASS."],
            ["Windows Server", "Hyper-V Generación 2", "No depende de Docker Desktop. Recomendado: Windows Server 2022; Server 2025 debe pasar smoke test específico antes de adoptarse."],
            ["Ubuntu", "Instalación Linux systemd", "Ubuntu 24.04 LTS recomendado; Python 3.11+; paquetes offline primero y repositorio firmado como fallback."],
            ["Linux Debian/RHEL compatible", "Instalación Linux systemd", "Debian no Ubuntu, RHEL, Rocky, Alma, CentOS o Fedora cuando sus repositorios proporcionan Python 3.11+."],
        ],
        [Cm(4.0), Cm(4.2), Cm(9.0)],
    )
    doc.add_paragraph(
        "Hyper-V requiere CPU x64 con SLAT/EPT/NPT, extensiones de modo monitor, VT-x o AMD-V, DEP/NX y virtualización habilitada en UEFI/BIOS [1]. Si GIDEON se instala dentro de otra máquina virtual, el hipervisor superior debe exponer virtualización anidada."
    )
    doc.add_paragraph(
        "Docker Desktop requiere actualmente 8 GB de RAM, WSL2 actualizado y virtualización por hardware; no está soportado en Windows Server [2]. Por ello, la ruta Server del OneClick usa la micro-VM Hyper-V."
    )

    doc.add_heading("4. Dimensionamiento del servidor central", level=1)
    add_docx_table(
        doc,
        ["Escenario", "CPU", "RAM", "SSD/NVMe libre", "Capacidad inicial prudente"],
        [
            ["Laboratorio mínimo", "4 vCPU/hilos", "8 GB", "40–60 GB", "25–50 agentes; 1–2 RDP/VNC."],
            ["Producción pequeña", "8 núcleos", "16 GB", "100 GB", "100–200 agentes; 2–4 RDP/VNC sin redimensionar guacd."],
            ["Producción con crecimiento", "12–16 núcleos", "32 GB", "250 GB", "400–800 agentes y 20–30 escritorios solo después de ajustes y pruebas."],
            ["Alta disponibilidad", "16+ núcleos", "32–64 GB", "500 GB+", "Exige múltiples gateways, base de datos externa y arquitectura HA."],
        ],
        [Cm(3.4), Cm(2.5), Cm(2.2), Cm(3.0), Cm(6.3)],
    )
    add_docx_bullets(
        doc,
        [
            "Red recomendada: Ethernet 1 Gbps como mínimo, IP estable, DNS funcional y sincronización NTP.",
            "En host físico: RAM ECC y almacenamiento NVMe con respaldo son preferibles.",
            "Bundle actual: 4.89 GiB, 17,240 archivos; payload Windows 1.49 GiB, datos 0.50 GiB y caché offline 2.88 GiB.",
            "El preflight del instalador rechaza menos de 6 GB libres, pero ese valor es un piso técnico; no es una recomendación operativa.",
            "Planificar espacio adicional para copia transaccional, VHDX, respaldos, logs, actualizaciones, SQLite y transferencias de hasta 512 MB.",
        ],
    )

    doc.add_heading("5. Servicios instalados y consumo base", level=1)
    doc.add_paragraph("El manifiesto instala siete servicios automáticos bajo LocalSystem:")
    add_docx_table(
        doc,
        ["Servicio", "Función principal"],
        [
            ["GIDEONRemoteControl", "Controlador/orquestación principal."],
            ["GIDEONRemoteControlBackup", "Respaldos y recuperación; en Windows el host nativo inicia Python solo durante el respaldo diario."],
            ["GIDEONRemoteControlDiagnostics", "Diagnóstico, salud y exportación de evidencia."],
            ["GIDEONRemoteControlGateway", "API HTTPS, agentes, jobs y túneles reversos."],
            ["GIDEONRemoteControlUpdater", "Aplicación transaccional de actualizaciones; activación por evento .ready, sin polling Python residente."],
            ["GIDEONRemoteControlWebAdmin", "Portal administrativo, sesiones y operaciones."],
            ["GIDEONAgent", "Agente local del propio servidor."],
        ],
        [Cm(6.2), Cm(11.0)],
    )
    add_docx_table(
        doc,
        ["Medición en reposo", "Consumo observado"],
        [
            ["Árboles de procesos GIDEON", "≈222.8 MB de working set."],
            ["Contenedor guacd", "≈13.73 MiB, CPU 0 %, 1 PID."],
            ["Docker Desktop + vmmem", "≈1.41 GB de working set."],
            ["Incremento total Windows Home", "≈1.65 GB en reposo; las sesiones y trabajos añaden consumo."],
        ],
        [Cm(7.0), Cm(10.2)],
    )

    doc.add_heading("6. Límites configurados por componente", level=1)
    add_docx_table(
        doc,
        ["Componente", "Límite actual", "Lectura operativa"],
        [
            ["Gateway", "256 conexiones; backlog 512", "Planificar 150–200 agentes online con margen. Las conexiones HTTP y túneles comparten recursos."],
            ["Agentes / jobs", "20 jobs activos en la flota", "El resto queda QUEUED. Un agente ejecuta un trabajo a la vez; admite cola offline."],
            ["Heartbeat y polling", "Heartbeat ≈30 s; polling 5 s", "Tráfico base bajo; inventario, scripts y archivos producen los picos."],
            ["WebAdmin", "256 requests; backlog 512", "No equivale a 256 administradores humanos. Uvicorn de un proceso y SQLite WAL."],
            ["Sesiones web", "5 por usuario; TTL/idle 12 h", "MFA obligatorio y controles de intentos fallidos."],
            ["RDP/VNC web", "10 globales; 3/usuario; 1/destino; 4 h", "Máximo lógico. Iniciar pruebas con 2–4 sesiones por el tamaño actual de guacd."],
            ["Túneles reversos", "64 reservas; 4 streams/agente", "La política de sesión limita a 1 escritorio por destino de forma predeterminada."],
            ["SSH agentless", "5 workers", "Windows OpenSSH/SFTP/PowerShell. Linux/macOS deben usar preferentemente Agent HTTPS/PULL."],
            ["Fanout", "6 dispatch + 6 poll", "VPN adaptativa entre 2 y 6 workers; cancelación con 10 workers."],
            ["Transferencias", "512 MB por upload", "4 workers normales, 2 grandes, 1 crítico; espera máxima 1 hora."],
            ["Scripts", "5 MB; 30 min por defecto", "Doble aprobación y análisis antimalware activados."],
            ["Enrolamiento", "8 usos iniciales", "Debe emitirse y rotarse conforme a la política de incorporación."],
        ],
        [Cm(3.6), Cm(4.5), Cm(9.1)],
    )

    doc.add_heading("7. RDP, VNC, SSH y agentes", level=1)
    doc.add_heading("7.1 RDP y VNC desde navegador", level=2)
    doc.add_paragraph(
        "La integración WebAdmin actual acepta únicamente protocolos RDP y VNC. Apache Guacamole upstream también puede soportar SSH [7], pero el adaptador GIDEON actual rechaza ese protocolo para sesiones interactivas de navegador."
    )
    add_docx_bullets(
        doc,
        [
            "Docker: guacd limitado a 1 vCPU y 768 MB.",
            "Hyper-V: micro-VM con memoria dinámica y límites adaptativos compact/standard/performance; guacd conserva límites explícitos de RAM/CPU.",
            "Capacidad inicial recomendada: 2–4 sesiones simultáneas de oficina.",
            "Para sostener 5–10: asignar 2–4 vCPU y 2–4 GB a guacd y ejecutar carga representativa.",
            "Política RDP endurecida: portapapeles, impresión, unidades, entrada de audio, wallpaper, temas y suavizado deshabilitados de forma predeterminada.",
        ],
    )
    doc.add_paragraph(
        "Windows Home puede actuar como cliente, pero no como host RDP. Los destinos deben usar Windows Pro, Enterprise, Education o Server [3]. Para múltiples usuarios simultáneos en Windows Server se requieren las RDS CAL correspondientes [4]. Linux necesita xrdp y macOS usa VNC/Apple Remote Management."
    )

    doc.add_heading("7.2 Ancho de banda RDP", level=2)
    add_docx_table(
        doc,
        ["Actividad 1080p", "Referencia Microsoft", "Presupuesto GIDEON recomendado"],
        [
            ["Reposo", "≈0.3 Kbps", "Despreciable."],
            ["Word / Excel", "≈100–500 Kbps", "0.5–1 Mbps por sesión."],
            ["Web / gráficos", "≈0.9–6.5 Mbps", "2–5 Mbps por sesión normal."],
            ["Video", "≈2.5–9.5 Mbps", "Hasta 10 Mbps por sesión de alta actividad."],
        ],
        [Cm(4.3), Cm(5.3), Cm(7.6)],
    )
    doc.add_paragraph(
        "Microsoft indica que la resolución, el movimiento, el códec y las redirecciones determinan el consumo, y recomienda medir conexiones reales [5]. Para diez sesiones de uso normal conviene reservar 20–50 Mbps además del tráfico de agentes y transferencias."
    )

    doc.add_heading("7.3 SSH y administración agentless", level=2)
    add_docx_bullets(
        doc,
        [
            "SSH se usa para sondeo, ejecución, PowerShell y SFTP; no hay terminal SSH interactiva en el navegador en esta versión.",
            "Máximo actual: 5 conexiones SSH agentless simultáneas; 2 workers para exploración de backups.",
            "Strict host key checking activo, ruta VPN requerida y destinos públicos bloqueados.",
            "El modo agentless está diseñado principalmente para Windows con OpenSSH, PowerShell y tareas programadas.",
        ],
    )

    doc.add_heading("7.4 Agentes administrados", level=2)
    add_docx_table(
        doc,
        ["Plataforma destino", "Gestión", "Acceso remoto"],
        [
            ["Windows Pro/Enterprise/Education/Server", "Agent HTTPS/PULL o Windows agentless por SSH", "RDP nativo y túnel reverso."],
            ["Windows Home", "Agent HTTPS/PULL", "Sin host RDP nativo; GIDEON no instala servidores RDP/VNC de terceros."],
            ["Linux x64/ARM64/ARM32", "Agent HTTPS/PULL", "RDP si xrdp está instalado; túnel reverso saliente."],
            ["Raspberry Pi", "Agent HTTPS/PULL", "Según escritorio/xrdp disponible."],
            ["macOS", "Agent HTTPS/PULL", "VNC/Apple Remote Management."],
        ],
        [Cm(5.5), Cm(5.4), Cm(6.3)],
    )

    doc.add_heading("8. Red, puertos y controles de seguridad", level=1)
    add_docx_table(
        doc,
        ["Puerto", "Uso", "Exposición recomendada"],
        [
            ["TCP 28443", "WebAdmin HTTPS", "Solo red administrativa, VPN o ZTNA. No publicar directamente en Internet."],
            ["TCP 29443", "Gateway HTTPS, enrolamiento, heartbeat, jobs y túneles", "Publicar de forma controlada; preferible reverse proxy/WAF/LB con TLS/mTLS cuando haya agentes externos."],
            ["TCP 4822", "guacd", "Loopback únicamente; no exponer a la LAN o Internet."],
            ["TCP 22", "SSH/SFTP hacia destinos agentless", "Solo redes privadas/VPN y claves de host verificadas."],
            ["TCP 3389", "RDP hacia destinos", "Privado/VPN o túnel reverso; nunca exposición pública directa."],
            ["TCP 5900", "VNC hacia destinos", "Privado/VPN o túnel reverso; nunca exposición pública directa."],
        ],
        [Cm(2.5), Cm(6.0), Cm(8.7)],
    )
    add_docx_bullets(
        doc,
        [
            "MFA obligatorio para WebAdmin.",
            "Aprobación dual para scripts.",
            "Análisis antimalware de scripts y transferencias.",
            "Tickets remotos de un solo uso, duración corta y vinculados al usuario/destino.",
            "Destinos remotos limitados a CIDR privados y acceso por VPN; objetivos públicos bloqueados.",
            "Contenedor guacd read-only, sin capacidades Linux, no-new-privileges y publicación loopback.",
        ],
    )

    doc.add_heading("9. Evaluación del equipo actual", level=1)
    add_docx_table(
        doc,
        ["Elemento", "Detectado", "Evaluación"],
        [
            ["Sistema operativo", "Windows 10 Home Single Language 22H2", "Apto solo para laboratorio Docker/WSL2; no dispone del rol Microsoft Hyper-V ni puede ser host RDP."],
            ["Procesador", "Intel Core i5-6300HQ; 4 lógicos", "Cumple el mínimo funcional, no el recomendado para producción."],
            ["RAM", "15.9 GB", "Suficiente para laboratorio y carga pequeña."],
            ["Disco C:", "25.8 GB libres", "Probablemente alcanza un smoke test, pero queda debajo del margen recomendado de 40–60 GB."],
            ["Disco D:", "344.2 GB libres", "Adecuado para datos; Docker/WSL suele mantener su disco virtual en C: salvo reconfiguración."],
            ["Backend", "Docker Desktop + WSL2", "Funcional e idempotencia validada."],
            ["Uso productivo", "No recomendado", "Windows 10 finalizó soporte general el 14 de octubre de 2025."],
        ],
        [Cm(4.0), Cm(5.2), Cm(8.0)],
    )
    add_docx_callout(
        doc,
        "RECOMENDACIÓN PARA LA ACEPTACIÓN",
        "Usar una máquina limpia Windows Server 2022 o Windows 11 Pro con 8 vCPU, 16 GB RAM y 100 GB NVMe libres. El equipo actual puede verificar nuevamente la rama Docker, pero no puede cerrar el gate Hyper-V.",
        LIGHT_GREEN,
        GREEN,
    )

    doc.add_heading("10. Procedimiento de prueba de instalación", level=1)
    steps = [
        "Preparar un host limpio, aplicar actualizaciones del sistema, verificar UEFI/BIOS, virtualización, DNS, NTP, espacio libre y conectividad.",
        "Copiar y extraer el bundle localmente; conservar el manifiesto y los hashes sin modificar.",
        "Ejecutar installers\\WINDOWS\\INSTALL_GIDEON_MASTER_WINDOWS.cmd como administrador.",
        "Permitir el reinicio si el OneClick habilita Hyper-V. Verificar que la continuación automática retome la instalación.",
        "Confirmar INSTALL=PASS y ausencia de un ZIP de diagnóstico de fallo.",
        "Verificar los siete servicios GIDEON, la micro-VM o contenedor guacd y los puertos 28443/29443/4822 según su alcance.",
        "Acceder a WebAdmin, cambiar la credencial temporal, configurar MFA y conservar los secretos en un almacén seguro.",
        "Enrolar al menos un agente Windows y uno Linux; comprobar DNS, TCP, TLS, enrolamiento, heartbeat e inventario.",
        "Ejecutar un script aprobado, una transferencia, un trabajo en cola, una cancelación y una recuperación después de pérdida de conectividad.",
        "Probar RDP contra Windows Pro/Server y VNC o xrdp según plataforma.",
        "Reiniciar el servidor y confirmar arranque automático, persistencia de configuración y recuperación de agentes/túneles.",
        "Ejecutar carga inicial con 20 jobs activos y 2–4 sesiones remotas; registrar CPU, RAM, disco, latencia y errores.",
    ]
    for step in steps:
        doc.add_paragraph(step, style="List Number")

    doc.add_heading("11. Criterios de aceptación y gates pendientes", level=1)
    add_docx_table(
        doc,
        ["Gate", "Criterio PASS"],
        [
            ["Instalación", "OneClick finaliza con INSTALL=PASS; rollback y diagnóstico disponibles si falla."],
            ["Hyper-V", "Candidata arranca como Generación 2 con Secure Boot, cloud-init, Docker offline y guacd TCP 4822."],
            ["Servicios", "Siete servicios automáticos en estado esperado; reinicio limpio y sin procesos huérfanos."],
            ["Gateway", "Agente completa DNS/TCP/TLS, enrolamiento, heartbeat, polling y reporte."],
            ["Operaciones", "Jobs, aprobación, cancelación, cola offline, scripts y transferencias funcionan."],
            ["Acceso remoto", "RDP/VNC abre, respeta VPN/CIDR, timeout, límites y cierre de ticket/sesión."],
            ["Carga", "Sin saturación sostenida, timeouts ni crecimiento anómalo de memoria/disco en la carga objetivo."],
            ["Seguridad", "MFA, doble aprobación, claves de host, TLS, ACL, firewall y secretos verificados."],
        ],
        [Cm(4.0), Cm(13.2)],
    )
    add_docx_callout(
        doc,
        "DICTAMEN FINAL",
        "Listo para prueba controlada. No declarar producción hasta completar el gate Hyper-V real y una prueba de carga representativa. Antes de escalar por encima de aproximadamente 200 agentes o 10 sesiones, ajustar arquitectura, límites y persistencia.",
        LIGHT_AMBER,
        AMBER,
    )

    doc.add_heading("12. Mejoras recomendadas antes de producción", level=1)
    add_docx_bullets(
        doc,
        [
            "Elevar el preflight de disco desde el piso de 6 GB a un cálculo dinámico por backend y volumen.",
            "Validar con carga real los perfiles adaptativos compact/standard/performance y ajustar sus umbrales únicamente con métricas.",
            "Añadir una prueba de carga automatizada que cubra Gateway, 20 jobs, transferencias y sesiones RDP/VNC.",
            "Definir métricas y alertas de CPU, RAM, conexiones, cola de jobs, latencia, SQLite WAL y espacio libre.",
            "Para más de 200 agentes sostenidos, validar aumento de gateway_max_connections y estudiar base de datos externa/HA.",
            "Documentar retención de logs, respaldos, recuperación y objetivo RPO/RTO del entorno productivo.",
        ],
    )

    doc.add_heading("13. Evidencia local", level=1)
    for label, path in (
        ("Validación OneClick", VALIDATION_PATH),
        ("Informe de calidad y reparaciones", QUALITY_PATH),
        ("Instalador Windows", ROOT / "installers" / "WINDOWS" / "INSTALL_GIDEON_MASTER_WINDOWS.cmd"),
        ("Configuración maestra", ROOT / "master.config.json"),
    ):
        p = doc.add_paragraph()
        p.add_run(label + ": ").bold = True
        p.add_run(str(path))

    doc.add_heading("14. Fuentes oficiales", level=1)
    for label, url in SOURCES:
        p = doc.add_paragraph()
        add_hyperlink(p, label, url)
        p.add_run("\n" + url).font.size = Pt(8)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Fin del informe")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_PATH)


def register_pdf_fonts() -> tuple[str, str]:
    regular = Path(r"C:\Windows\Fonts\arial.ttf")
    bold = Path(r"C:\Windows\Fonts\arialbd.ttf")
    if regular.exists() and bold.exists():
        pdfmetrics.registerFont(TTFont("GIDEONArial", str(regular)))
        pdfmetrics.registerFont(TTFont("GIDEONArialBold", str(bold)))
        return "GIDEONArial", "GIDEONArialBold"
    return "Helvetica", "Helvetica-Bold"


def build_pdf(validation: dict, quality: dict) -> None:
    regular_font, bold_font = register_pdf_fonts()
    page_width, page_height = A4
    doc = BaseDocTemplate(
        str(PDF_PATH),
        pagesize=A4,
        leftMargin=17 * mm,
        rightMargin=17 * mm,
        topMargin=20 * mm,
        bottomMargin=17 * mm,
        title="GIDEON PRE-RC31 — Informe técnico de instalación y capacidad",
        author="GIDEON Remote Control",
        subject="Requisitos, límites, capacidades y prueba de aceptación",
    )

    def page_header_footer(canvas, document):
        canvas.saveState()
        canvas.setFont(regular_font, 7.5)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawString(17 * mm, page_height - 11 * mm, f"{REPORT_VERSION}  |  Informe técnico")
        canvas.drawRightString(page_width - 17 * mm, 9 * mm, f"Página {document.page}")
        canvas.setStrokeColor(colors.HexColor("#D9E2F3"))
        canvas.line(17 * mm, page_height - 13 * mm, page_width - 17 * mm, page_height - 13 * mm)
        canvas.restoreState()

    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="normal")
    doc.addPageTemplates([PageTemplate(id="gideon", frames=[frame], onPage=page_header_footer)])

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="GIDEONTitle", parent=styles["Title"], fontName=bold_font, fontSize=24, leading=28, textColor=colors.HexColor("#17365D"), alignment=TA_CENTER, spaceAfter=9 * mm))
    styles.add(ParagraphStyle(name="GIDEONSub", parent=styles["Normal"], fontName=regular_font, fontSize=11, leading=15, textColor=colors.HexColor("#666666"), alignment=TA_CENTER, spaceAfter=6 * mm))
    styles.add(ParagraphStyle(name="GIDEONH1", parent=styles["Heading1"], fontName=bold_font, fontSize=15, leading=18, textColor=colors.HexColor("#17365D"), spaceBefore=5 * mm, spaceAfter=2.5 * mm))
    styles.add(ParagraphStyle(name="GIDEONH2", parent=styles["Heading2"], fontName=bold_font, fontSize=11.5, leading=14, textColor=colors.HexColor("#1F4E78"), spaceBefore=3.5 * mm, spaceAfter=1.5 * mm))
    styles.add(ParagraphStyle(name="GIDEONBody", parent=styles["BodyText"], fontName=regular_font, fontSize=8.7, leading=12, textColor=colors.HexColor("#222222"), spaceAfter=2.2 * mm))
    styles.add(ParagraphStyle(name="GIDEONBullet", parent=styles["BodyText"], fontName=regular_font, fontSize=8.5, leading=11.5, leftIndent=5 * mm, firstLineIndent=-3 * mm, bulletIndent=1.5 * mm, spaceAfter=1.2 * mm))
    styles.add(ParagraphStyle(name="GIDEONTable", parent=styles["BodyText"], fontName=regular_font, fontSize=7.2, leading=9.2, textColor=colors.HexColor("#222222")))
    styles.add(ParagraphStyle(name="GIDEONTableHead", parent=styles["BodyText"], fontName=bold_font, fontSize=7.2, leading=9.2, textColor=colors.white, alignment=TA_LEFT))
    styles.add(ParagraphStyle(name="GIDEONCalloutTitle", parent=styles["BodyText"], fontName=bold_font, fontSize=10, leading=12, textColor=colors.HexColor("#9C6500"), spaceAfter=1.5 * mm))
    styles.add(ParagraphStyle(name="GIDEONSource", parent=styles["BodyText"], fontName=regular_font, fontSize=7.8, leading=10.5, textColor=colors.HexColor("#0563C1"), spaceAfter=2 * mm))

    story = []

    def body(text: str):
        story.append(Paragraph(text, styles["GIDEONBody"]))

    def h1(text: str):
        story.append(Paragraph(text, styles["GIDEONH1"]))

    def h2(text: str):
        story.append(Paragraph(text, styles["GIDEONH2"]))

    def bullets(items: Iterable[str]):
        for item in items:
            story.append(Paragraph("• " + item, styles["GIDEONBullet"]))

    def callout(title: str, text: str, fill="#FFF2CC", color="#9C6500"):
        title_style = ParagraphStyle("temp", parent=styles["GIDEONCalloutTitle"], textColor=colors.HexColor(color))
        data = [[Paragraph(title, title_style), Paragraph(text, styles["GIDEONBody"])]]
        t = Table(data, colWidths=[42 * mm, 123 * mm])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(fill)),
            ("BOX", (0, 0), (-1, -1), 0.7, colors.HexColor(color)),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        story.extend([t, Spacer(1, 3 * mm)])

    def table(headers: list[str], rows: list[list[str]], widths_mm: list[float]):
        data = [[Paragraph(str(h), styles["GIDEONTableHead"]) for h in headers]]
        for row in rows:
            data.append([Paragraph(str(v).replace("&", "&amp;"), styles["GIDEONTable"]) for v in row])
        t = Table(data, colWidths=[w * mm for w in widths_mm], repeatRows=1, hAlign="LEFT")
        commands = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17365D")),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#A6A6A6")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]
        for index in range(2, len(data), 2):
            commands.append(("BACKGROUND", (0, index), (-1, index), colors.HexColor("#EEF5FB")))
        t.setStyle(TableStyle(commands))
        story.extend([t, Spacer(1, 3 * mm)])

    story.append(Spacer(1, 24 * mm))
    story.append(Paragraph("GIDEON REMOTE CONTROL", ParagraphStyle("brand", parent=styles["GIDEONSub"], fontName=bold_font, fontSize=14, textColor=colors.HexColor("#1F4E78"))))
    story.append(Paragraph("Informe técnico de instalación,<br/>requisitos y capacidad", styles["GIDEONTitle"]))
    story.append(Paragraph(f"{REPORT_VERSION}<br/>{REPORT_DATE}", styles["GIDEONSub"]))
    callout("DICTAMEN EJECUTIVO", "GO para una prueba de instalación controlada. La aprobación de producción queda condicionada a ejecutar el OneClick completo en un host Microsoft Hyper-V real y a realizar una prueba de carga con la cantidad prevista de agentes y sesiones remotas.")
    story.append(Spacer(1, 20 * mm))
    story.append(Paragraph("Documento de ingeniería · No sustituye una prueba de capacidad en el entorno objetivo", styles["GIDEONSub"]))
    story.append(PageBreak())

    h1("1. Alcance y criterio de evaluación")
    body("Este informe consolida el estado técnico del bundle GIDEON Master PRE-RC31, el comportamiento adaptativo del instalador OneClick, los requisitos de hardware y sistema operativo, los límites configurados por servicio y una estimación conservadora de capacidad. Se diferencia entre:")
    bullets([
        "Límite configurado: restricción impuesta actualmente por el código o la configuración.",
        "Capacidad práctica inicial: carga prudente para comenzar una prueba de aceptación.",
        "Capacidad ampliada: posible únicamente después de redimensionar, ajustar parámetros y validar con carga real.",
    ])

    h1("2. Estado de preparación")
    table(["Área", "Resultado", "Interpretación"], [
        ["OneClick / plataforma", validation.get("result", "NO DISPONIBLE"), "Windows Home + Docker validado; ejecución Hyper-V real pendiente."],
        ["Calidad integral", quality.get("result", "NO DISPONIBLE"), "Sintaxis, estructura, hashes e importación aprobadas."],
        ["Imagen Hyper-V", "PREVALIDACIÓN PASS", "UEFI, cloud-init, Docker offline y TCP 4822 validados con QEMU; aún no con Hyper-V."],
        ["Producción", "CONDICIONAL", "Requiere prueba Hyper-V real, reinicio, enrolamiento, acceso remoto y carga."],
    ], [34, 52, 79])
    tests = quality.get("tests", {})
    table(["Control", "Cobertura", "Estado"], [
        ["PowerShell", f"{tests.get('powershell_syntax', {}).get('files', 551)} archivos", tests.get("powershell_syntax", {}).get("status", "PASS")],
        ["Python AST", f"{tests.get('python_ast', {}).get('files', 3243)} archivos", tests.get("python_ast", {}).get("status", "PASS")],
        ["JSON", f"{tests.get('json_parse', {}).get('files', 1390)} archivos", tests.get("json_parse", {}).get("status", "PASS")],
        ["Shell", f"{tests.get('shell_syntax', {}).get('files', 45)} archivos", tests.get("shell_syntax", {}).get("status", "PASS")],
        ["Manifiesto / hashes", "16,644 archivos elegibles", "PASS · 0 errores"],
        ["YARA", "Código del bundle", tests.get("yara_code_scan", "PASS")],
        ["Importación", f"{tests.get('validated_import_transaction_rounds', 2)} rondas", tests.get("validated_import_transaction", "PASS")],
    ], [54, 65, 46])

    h1("3. Comportamiento del OneClick por sistema operativo")
    table(["Plataforma", "Backend", "Automatización y observaciones"], [
        ["Windows Home / Core", "Docker Desktop + WSL2", "Preflight, imagen offline, guacd loopback e idempotencia. La licencia de Docker requiere aceptación del operador."],
        ["Windows Pro / Enterprise / Education", "Hyper-V Gen. 2", "Habilita componentes, registra continuación, reinicia, valida y solo promueve después de PASS."],
        ["Windows Server", "Hyper-V Gen. 2", "No depende de Docker Desktop. Windows Server 2022 recomendado."],
        ["Ubuntu", "Linux systemd", "Ubuntu 24.04 LTS recomendado; Python 3.11+; offline primero y repositorio firmado como fallback."],
        ["Debian/RHEL compatible", "Linux systemd", "Debian no Ubuntu, RHEL, Rocky, Alma, CentOS o Fedora con Python 3.11+."],
    ], [39, 41, 85])
    body("Hyper-V requiere CPU x64 con SLAT/EPT/NPT, extensiones de modo monitor, VT-x o AMD-V, DEP/NX y virtualización habilitada en UEFI/BIOS [1]. Si GIDEON se instala dentro de otra VM, el hipervisor superior debe exponer virtualización anidada.")
    body("Docker Desktop requiere actualmente 8 GB de RAM, WSL2 actualizado y virtualización por hardware; no está soportado en Windows Server [2]. La ruta Server del OneClick usa la micro-VM Hyper-V.")

    h1("4. Dimensionamiento del servidor central")
    table(["Escenario", "CPU", "RAM", "Disco libre", "Capacidad prudente"], [
        ["Laboratorio mínimo", "4 vCPU", "8 GB", "40–60 GB", "25–50 agentes; 1–2 RDP/VNC."],
        ["Producción pequeña", "8 núcleos", "16 GB", "100 GB", "100–200 agentes; 2–4 RDP/VNC sin redimensionar."],
        ["Con crecimiento", "12–16", "32 GB", "250 GB", "400–800 agentes y 20–30 escritorios tras ajustes y pruebas."],
        ["Alta disponibilidad", "16+", "32–64 GB", "500 GB+", "Múltiples gateways, BD externa y diseño HA."],
    ], [32, 23, 19, 28, 63])
    bullets([
        "Ethernet 1 Gbps mínimo, IP estable, DNS funcional y NTP.",
        "Bundle: 4.89 GiB y 17,240 archivos; payload Windows 1.49 GiB, datos 0.50 GiB y caché offline 2.88 GiB.",
        "El rechazo por menos de 6 GB es un piso técnico, no una recomendación operativa.",
        "Reservar espacio para copia transaccional, VHDX, respaldos, logs, SQLite y transferencias.",
    ])

    h1("5. Servicios instalados y consumo base")
    table(["Servicio", "Función principal"], [
        ["GIDEONRemoteControl", "Controlador/orquestación principal."],
        ["GIDEONRemoteControlBackup", "Respaldo diario; Python solo durante la ejecución en Windows."],
        ["GIDEONRemoteControlDiagnostics", "Diagnóstico y evidencia."],
        ["GIDEONRemoteControlGateway", "API HTTPS, agentes, jobs y túneles."],
        ["GIDEONRemoteControlUpdater", "Actualizaciones por evento .ready, sin polling Python residente."],
        ["GIDEONRemoteControlWebAdmin", "Portal administrativo y operaciones."],
        ["GIDEONAgent", "Agente local del servidor."],
    ], [62, 103])
    table(["Medición en reposo", "Consumo observado"], [
        ["Procesos GIDEON", "≈222.8 MB."],
        ["guacd", "≈13.73 MiB; CPU 0 %."],
        ["Docker Desktop + vmmem", "≈1.41 GB."],
        ["Incremento Windows Home", "≈1.65 GB; las sesiones y jobs añaden consumo."],
    ], [70, 95])

    h1("6. Límites configurados por componente")
    table(["Componente", "Límite actual", "Lectura operativa"], [
        ["Gateway", "256; backlog 512", "Planificar 150–200 agentes online con margen."],
        ["Agentes / jobs", "20 activos", "Resto QUEUED; un job por agente; cola offline."],
        ["Heartbeat / polling", "≈30 s / 5 s", "Tráfico base bajo; archivos y jobs producen picos."],
        ["WebAdmin", "256 requests; backlog 512", "Uvicorn de un proceso y SQLite WAL."],
        ["Sesiones web", "5/usuario; 12 h", "MFA y control de intentos fallidos."],
        ["RDP/VNC", "10; 3/usuario; 1/destino; 4 h", "Iniciar con 2–4 por tamaño actual de guacd."],
        ["Túneles", "64 reservas; 4/agente", "Política remota limita 1 escritorio/destino."],
        ["SSH", "5 workers", "Agentless Windows OpenSSH/SFTP/PowerShell."],
        ["Fanout", "6 dispatch + 6 poll", "VPN adaptativa 2–6; cancelación 10."],
        ["Transferencias", "512 MB", "Workers 4/2/1; espera máxima 1 hora."],
        ["Scripts", "5 MB; 30 min", "Doble aprobación y antimalware."],
        ["Enrolamiento", "8 usos", "Emitir y rotar según política."],
    ], [36, 49, 80])

    h1("7. RDP, VNC, SSH y agentes")
    h2("7.1 RDP y VNC desde navegador")
    body("WebAdmin acepta únicamente RDP y VNC. Apache Guacamole upstream puede soportar SSH [7], pero el adaptador GIDEON actual rechaza SSH para sesiones interactivas de navegador.")
    bullets([
        "Docker: guacd limitado a 1 vCPU y 768 MB.",
        "Hyper-V: micro-VM con memoria dinámica y perfiles adaptativos; guacd conserva límites explícitos.",
        "Comenzar con 2–4 sesiones de oficina; para 5–10 asignar 2–4 vCPU y 2–4 GB y probar carga.",
        "Portapapeles, impresión, unidades, audio de entrada, wallpaper y temas deshabilitados por defecto.",
    ])
    body("Windows Home no puede ser host RDP [3]. Los destinos deben usar Windows Pro, Enterprise, Education o Server. Múltiples usuarios sobre Windows Server requieren RDS CAL [4]. Linux necesita xrdp y macOS usa VNC/Apple Remote Management.")
    h2("7.2 Ancho de banda RDP")
    table(["Actividad 1080p", "Microsoft", "Presupuesto GIDEON"], [
        ["Reposo", "≈0.3 Kbps", "Despreciable."],
        ["Word / Excel", "≈100–500 Kbps", "0.5–1 Mbps/sesión."],
        ["Web / gráficos", "≈0.9–6.5 Mbps", "2–5 Mbps/sesión."],
        ["Video", "≈2.5–9.5 Mbps", "Hasta 10 Mbps/sesión."],
    ], [48, 54, 63])
    body("Microsoft recomienda medir conexiones reales [5]. Para diez sesiones normales conviene reservar 20–50 Mbps, además del tráfico de agentes y transferencias.")
    h2("7.3 SSH agentless")
    bullets([
        "Sondeo, ejecución, PowerShell y SFTP; no hay terminal SSH web en esta versión.",
        "Cinco conexiones simultáneas; dos workers para exploración de respaldos.",
        "Host key checking estricto, VPN requerida y destinos públicos bloqueados.",
        "Diseñado principalmente para Windows; Linux/macOS usan preferentemente Agent HTTPS/PULL.",
    ])
    h2("7.4 Plataformas de Agent")
    table(["Destino", "Gestión", "Acceso remoto"], [
        ["Windows Pro/Enterprise/Education/Server", "Agent o SSH agentless", "RDP + túnel reverso."],
        ["Windows Home", "Agent HTTPS/PULL", "Sin host RDP nativo; no instala terceros."],
        ["Linux x64/ARM", "Agent HTTPS/PULL", "xrdp + túnel reverso."],
        ["Raspberry Pi", "Agent HTTPS/PULL", "Según xrdp/escritorio."],
        ["macOS", "Agent HTTPS/PULL", "VNC/Apple Remote Management."],
    ], [55, 52, 58])

    h1("8. Red, puertos y seguridad")
    table(["Puerto", "Uso", "Exposición recomendada"], [
        ["28443/TCP", "WebAdmin HTTPS", "Solo VPN/ZTNA/red administrativa."],
        ["29443/TCP", "Gateway, agentes y túneles", "Publicación controlada; reverse proxy/WAF/LB con TLS/mTLS."],
        ["4822/TCP", "guacd", "Loopback únicamente."],
        ["22/TCP", "SSH/SFTP", "Privado/VPN y claves verificadas."],
        ["3389/TCP", "RDP", "Privado/VPN o túnel; nunca público directo."],
        ["5900/TCP", "VNC", "Privado/VPN o túnel; nunca público directo."],
    ], [27, 56, 82])
    bullets([
        "MFA obligatorio y aprobación dual para scripts.",
        "Análisis antimalware de scripts y transferencias.",
        "Tickets remotos de un uso, corta duración y vinculados a usuario/destino.",
        "CIDR privados y VPN; objetivos públicos bloqueados.",
        "guacd read-only, cap-drop ALL, no-new-privileges y loopback.",
    ])

    h1("9. Evaluación del equipo actual")
    table(["Elemento", "Detectado", "Evaluación"], [
        ["SO", "Windows 10 Home 22H2", "Laboratorio Docker/WSL2; sin rol Hyper-V ni host RDP."],
        ["CPU", "i5-6300HQ; 4 lógicos", "Mínimo funcional; no recomendado producción."],
        ["RAM", "15.9 GB", "Suficiente para laboratorio."],
        ["C:", "25.8 GB libres", "Probable smoke test; menor que 40–60 GB recomendado."],
        ["D:", "344.2 GB libres", "Adecuado para datos; Docker/WSL suele permanecer en C:."],
        ["Backend", "Docker Desktop + WSL2", "Funcional e idempotencia validada."],
        ["Producción", "No recomendado", "Windows 10 terminó soporte general el 14/10/2025."],
    ], [35, 50, 80])
    callout("RECOMENDACIÓN", "Usar una máquina limpia Windows Server 2022 o Windows 11 Pro con 8 vCPU, 16 GB RAM y 100 GB NVMe libres. El equipo actual no puede cerrar el gate Hyper-V.", "#E2F0D9", "#2E7D32")

    h1("10. Procedimiento de prueba")
    steps = [
        "Preparar host limpio, actualizar SO y verificar UEFI/BIOS, virtualización, DNS, NTP, disco y red.",
        "Copiar y extraer el bundle localmente; preservar manifiesto y hashes.",
        "Ejecutar installers\\WINDOWS\\INSTALL_GIDEON_MASTER_WINDOWS.cmd como administrador.",
        "Permitir reinicio si se habilita Hyper-V y confirmar continuación automática.",
        "Confirmar INSTALL=PASS y revisar diagnósticos si existieran.",
        "Verificar siete servicios, micro-VM/contenedor y puertos según alcance.",
        "Entrar a WebAdmin, cambiar credencial temporal y configurar MFA.",
        "Enrolar un agente Windows y uno Linux; comprobar TLS, heartbeat e inventario.",
        "Probar script, transferencia, cola, cancelación y recuperación de conectividad.",
        "Probar RDP y VNC/xrdp según plataforma.",
        "Reiniciar el servidor y comprobar recuperación automática.",
        "Cargar 20 jobs y 2–4 sesiones; medir CPU, RAM, disco, latencia y errores.",
    ]
    for index, step in enumerate(steps, 1):
        story.append(Paragraph(f"{index}. {step}", styles["GIDEONBullet"]))

    h1("11. Gates de aceptación")
    table(["Gate", "Criterio PASS"], [
        ["Instalación", "INSTALL=PASS; rollback y diagnóstico disponibles."],
        ["Hyper-V", "Gen. 2, Secure Boot, cloud-init, Docker offline y TCP 4822."],
        ["Servicios", "Siete servicios y reinicio limpio, sin huérfanos."],
        ["Gateway", "DNS/TCP/TLS, enrolamiento, heartbeat, polling y reporte."],
        ["Operaciones", "Jobs, aprobación, cancelación, cola offline y archivos."],
        ["Remoto", "RDP/VNC respeta VPN/CIDR, límites, timeout y cierre."],
        ["Carga", "Sin saturación sostenida, timeouts ni crecimiento anómalo."],
        ["Seguridad", "MFA, claves, TLS, ACL, firewall y secretos verificados."],
    ], [39, 126])
    callout("DICTAMEN FINAL", "Listo para prueba controlada. No declarar producción hasta completar Hyper-V real y una carga representativa. Antes de superar aproximadamente 200 agentes o 10 sesiones, ajustar arquitectura, límites y persistencia.")

    h1("12. Mejoras recomendadas antes de producción")
    bullets([
        "Calcular el mínimo de disco dinámicamente por backend y volumen, en lugar del piso fijo de 6 GB.",
        "Validar con carga real los perfiles adaptativos compact/standard/performance.",
        "Automatizar una prueba de carga de Gateway, 20 jobs, archivos y sesiones remotas.",
        "Alertar CPU, RAM, conexiones, cola, latencia, SQLite WAL y disco libre.",
        "Para más de 200 agentes, validar gateway_max_connections y evaluar base de datos externa/HA.",
        "Definir retención, respaldos, recuperación, RPO y RTO.",
    ])

    h1("13. Evidencia local")
    for label, path in (
        ("Validación OneClick", VALIDATION_PATH),
        ("Calidad y reparaciones", QUALITY_PATH),
        ("Instalador Windows", ROOT / "installers" / "WINDOWS" / "INSTALL_GIDEON_MASTER_WINDOWS.cmd"),
        ("Configuración maestra", ROOT / "master.config.json"),
    ):
        body(f"<b>{label}:</b> {str(path)}")

    h1("14. Fuentes oficiales")
    for label, url in SOURCES:
        safe_label = label.replace("&", "&amp;")
        story.append(Paragraph(f'<link href="{url}" color="#0563C1"><u>{safe_label}</u></link><br/><font size="7">{url}</font>', styles["GIDEONSource"]))

    doc.build(story)


def main() -> None:
    validation = load_json(VALIDATION_PATH)
    quality = load_json(QUALITY_PATH)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    build_docx(validation, quality)
    build_pdf(validation, quality)
    print(json.dumps({
        "docx": str(DOCX_PATH),
        "pdf": str(PDF_PATH),
        "generated_utc": datetime.now(UTC).isoformat(timespec="seconds").replace("+00:00", "Z"),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
